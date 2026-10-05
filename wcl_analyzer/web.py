"""Веб-интерфейс: python -m wcl_analyzer ui

Все данные пользователя хранятся у него в браузере: ключ Warcraft Logs API, эталоны,
настройки, история разборов и Excel-файлы. Сервер только считает:
получает запрос (ключ — в заголовках X-WCL-Id / X-WCL-Secret), загружает данные из WCL,
возвращает результат и Excel и ничего о пользователе не сохраняет — ни на диске, ни в cookie.

Режимы:
- локальный (по умолчанию) — программа на своём компьютере; кэш публичных данных WCL лежит
  на этом же компьютере; телефон в той же Wi-Fi сети подключается по ссылке-ключу из QR-кода;
- публичный (--public или WCL_PUBLIC=1) — сервис в интернете для любых пользователей;
  кэш публичных данных WCL только в оперативной памяти, на диск сервер ничего не пишет.
"""
from __future__ import annotations

import json
import math
import os
import re
import secrets
import socket
import tempfile
import threading
import time
import traceback
import uuid
import webbrowser
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path
from urllib.parse import parse_qs, urlparse

from . import settings
from .compare import CompareResult, _fmt_t
from .config import cache_path
from .metrics import CATEGORY_RU
from .names_ru import spec_ru

STATIC = Path(__file__).parent / "web" / "index.html"
JOBS: dict[str, dict] = {}
JOBS_LOCK = threading.Lock()
NEED_KEY = "Нужен ключ Warcraft Logs API: нажмите «Подключить ключ API» вверху страницы."
SERVER = {"port": 8765, "lan": False, "public": False}
MAX_JOBS = int(os.environ.get("WCL_MAX_JOBS", "3"))
MAX_QUEUE = int(os.environ.get("WCL_MAX_QUEUE", "20"))
JOB_SLOTS = threading.BoundedSemaphore(MAX_JOBS)
JOB_TTL_S = 3600            # результат и Excel живут в памяти час, пока браузер их не заберёт
DEFAULT_MAX_AGE_DAYS = 3


# ------------------------------------------------------------ клиент API
_CLIENTS: dict[str, object] = {}
_CLIENTS_LOCK = threading.Lock()
_CACHE = {"obj": None}


def _shared_cache():
    """Публичный сервис — кэш только в памяти; локальная программа — файл на этом компьютере."""
    from .api import Cache, MemoryCache
    with _CLIENTS_LOCK:
        if _CACHE["obj"] is None:
            _CACHE["obj"] = (MemoryCache(float(os.environ.get("WCL_MEMCACHE_MB", "150")))
                             if SERVER["public"] else Cache(cache_path()))
        return _CACHE["obj"]


def env_key() -> tuple[str, str] | None:
    """Ключ из .env — только для локальной программы (совместимость и командная строка)."""
    if SERVER["public"]:
        return None
    from .config import load_env
    load_env()
    cid, sec = os.environ.get("WCL_CLIENT_ID"), os.environ.get("WCL_CLIENT_SECRET")
    return (cid, sec) if cid and sec else None


def default_client_factory(creds: tuple[str, str] | None):
    """Клиент WCL для ключа из запроса. Клиенты живут в памяти, чтобы не запрашивать токен
    каждый раз; ключ по хэшу не восстановить, на диск он не попадает."""
    import hashlib
    from .api import WCLClient
    creds = creds or env_key()
    if not creds:
        raise SystemExit(NEED_KEY)
    k = hashlib.sha256(f"{creds[0]}:{creds[1]}".encode()).hexdigest()
    with _CLIENTS_LOCK:
        if k not in _CLIENTS:
            _CLIENTS[k] = WCLClient(creds[0], creds[1], None, verbose=False)
            while len(_CLIENTS) > 200:   # публичный сервер: не копим клиентов всех когда-либо приходивших ключей
                _CLIENTS.pop(next(iter(_CLIENTS)))
        else:
            _CLIENTS[k] = _CLIENTS.pop(k)   # недавно использованный — в конец очереди
        client = _CLIENTS[k]
    client.cache = _shared_cache()
    return client


CLIENT_FACTORY = default_client_factory  # тесты подменяют на фейковый клиент


def check_key(client_id: str, secret: str) -> None:
    """Проверяет ключ получением токена. Ключ нигде не сохраняется."""
    import requests
    from .config import TOKEN_URL
    if not client_id or not secret:
        raise ValueError("Заполните Client ID и Client Secret.")
    try:
        r = requests.post(TOKEN_URL, data={"grant_type": "client_credentials"},
                          auth=(client_id, secret), timeout=20)
    except requests.RequestException as e:
        raise ValueError(f"Не удалось связаться с Warcraft Logs: {e}") from e
    if r.status_code != 200:
        raise ValueError("Warcraft Logs не принял ключ: проверьте Client ID и Client Secret.")


CHECK_KEY = check_key  # тесты подменяют


# ------------------------------------------------------------- задачи
MAX_DONE = int(os.environ.get("WCL_MAX_DONE", "60"))  # готовых результатов в памяти не больше (публичный сервер)


def _cleanup() -> None:
    now = time.time()
    with JOBS_LOCK:
        done = [(j, job) for j, job in JOBS.items() if job["state"] != "running"]
        for jid, job in done:
            if now - job.get("finished", job["created"]) > JOB_TTL_S:
                JOBS.pop(jid, None)
        done = sorted(((j, job) for j, job in JOBS.items() if job["state"] != "running"),
                      key=lambda x: x[1].get("finished", x[1]["created"]))
        for jid, _job in done[:max(0, len(done) - MAX_DONE)]:  # сверх лимита — самые старые
            JOBS.pop(jid, None)


def start_job(creds, params: dict) -> str:
    _cleanup()
    with JOBS_LOCK:
        if sum(1 for j in JOBS.values() if j["state"] == "running") >= MAX_JOBS + MAX_QUEUE:
            raise ValueError("Сервер сейчас загружен — попробуйте через пару минут.")
        jid = uuid.uuid4().hex  # 128 бит: знает только браузер, запустивший разбор
        job = {"id": jid, "state": "running", "progress": 0.02, "log": [], "created": time.time()}
        JOBS[jid] = job
    threading.Thread(target=_run_job, args=(job, params, creds), daemon=True).start()
    return jid


def _run_job(job: dict, params: dict, creds) -> None:
    def log(msg: str) -> None:
        job["log"].append(msg)
        m = re.search(r"\[(\d+)/(\d+)\]", msg)
        if m:
            job["progress"] = 0.15 + 0.75 * int(m.group(1)) / int(m.group(2))

    if not JOB_SLOTS.acquire(blocking=False):
        log("Сервер занят другими разборами — ваш начнётся, как только освободится место…")
        JOB_SLOTS.acquire()
    try:
        mode = params.get("mode")
        if mode == "fight":
            _run_fight(job, params, creds, log)
        elif mode == "progress":
            _run_top_progress(job, params, creds, log)
        elif mode == "raid":
            _run_raid(job, params, creds, log)
        elif mode == "raidrot":
            _run_raid_rotation(job, params, creds, log)
        elif mode == "saves":
            _run_saves(job, params, creds, log)
        elif mode == "allbosses":
            _run_player_all(job, params, creds, log)
        elif mode == "refresh":
            _run_refresh(job, params, creds, log)
        elif mode == "rosterplan":
            _run_roster_plan(job, params, creds, log)
        else:
            _run_player(job, params, creds, log)
        job["progress"], job["state"] = 1.0, "done"
    except SystemExit:
        job["state"], job["error"], job["need_key"] = "error", NEED_KEY, True
    except Exception as e:  # noqa: BLE001
        job["state"], job["error"] = "error", _friendly(e)
        job["trace"] = traceback.format_exc()
    finally:
        job["finished"] = time.time()   # час хранения — от конца разбора, а не от начала
        JOB_SLOTS.release()


def app_mode_is_desktop() -> bool:
    from .platform_support import app_mode
    return app_mode() in ("exe", "python")


def _client(creds, job: dict, log, wait: bool = False):
    """Клиент WCL для задачи. Кончился часовой лимит API — разбор останавливается с понятной ошибкой
    (уже скачанное остаётся в кэше). wait=True — ждать сброса лимита (галочка «ждать сброса лимита»).
    Настройки ожидания — свои у каждого разбора: общий клиент ключа не меняется."""
    client = CLIENT_FACTORY(creds)

    def on_wait(seconds: float) -> None:
        job["wait_until"] = time.time() + seconds
        log(f"Закончился часовой лимит запросов Warcraft Logs. Жду сброса: {max(1, round(seconds / 60))} мин. "
            "Не закрывайте страницу: разбор продолжится сам.")

    return _limited(client, None if wait else 0, on_wait)


def _limited(client, max_wait_s, on_wait=None):
    if hasattr(client, "with_limits"):
        return client.with_limits(max_wait_s, on_wait)
    for k, v in (("max_wait_s", max_wait_s), ("on_wait", on_wait)):  # тестовые клиенты
        try:
            setattr(client, k, v)
        except AttributeError:
            pass
    return client


def _excel_bytes(job: dict, name: str, writer, data, key: str = "xlsx") -> None:
    """Excel собирается во временной папке и сразу удаляется; файл живёт только в памяти задачи.
    Внутри «Разобрать бой» (_SubJob) листы запоминаются, а файл собирается один на весь разбор."""
    if isinstance(job, _SubJob):
        job.writers.append((key, name, writer, data))
        return
    fname = re.sub(r"[^\w\-]+", "_", f"{name}_{time.strftime('%Y%m%d_%H%M%S')}") + ".xlsx"
    try:
        with tempfile.TemporaryDirectory() as tmp:
            path = writer(data, Path(tmp) / fname)
            job[key] = Path(path).read_bytes()
    except Exception as e:  # noqa: BLE001 — сбой Excel не должен терять уже готовый разбор
        job.setdefault("log", []).append(f"Excel не собран: {e}. Результат на странице — полный.")
        job["trace_excel"] = traceback.format_exc()
        return
    job[key + "_name"] = fname


def _max_age_s(params: dict) -> float:
    try:
        days = float(params.get("max_age_days") or DEFAULT_MAX_AGE_DAYS)
    except (TypeError, ValueError):
        days = DEFAULT_MAX_AGE_DAYS
    return max(1.0, min(30.0, days)) * 86400


MYTHIC = 5
# Размер эталона: топ-1 или топ-3 (по умолчанию); топ-10 — только для старых разборов из истории
REF_SIZES = {"top1": 1, "top3": 3, "top10": 10}


def _run_player(job: dict, params: dict, creds, log) -> None:
    """Разбор игрока: эталон — топ той же сложности, что и ваш бой; если бой не эпохальный —
    дополнительно эпохальный топ (переключатель в шапке результата)."""
    from .collect import collect_reference, load_my_log

    client = None
    ref_meta: dict = {}
    if params.get("demo"):
        from .demo import demo_logs
        log("Генерирую демо-логи: 25 игроков топа и ваш бой…")
        tops, me = demo_logs()
        label = "топ-25 (демо)"
        job["progress"] = 0.8
    else:
        client = _client(creds, job, log)
        log("Загружаю ваш лог…")
        me = load_my_log(client, params["url"], params.get("fight"),
                         actor_id=int(params["actor"]) if params.get("actor") not in (None, "") else None)
        log(f"{me.name}: {me.spec} {me.cls}, {me.encounter_name} ({me.difficulty_name}), "
            f"{_fmt_t(me.duration)}, {me.dps:,.0f} DPS")
        job["progress"] = 0.12
        if params.get("against"):
            log("Загружаю лог для сравнения…")
            other = load_my_log(client, params["against"])
            tops, label = [other], f"Игрок {other.name}"
        else:
            tops, label = _player_ref(client, me, me.difficulty, params, log, ref_meta)
    main = _player_result(job, params, client, me, tops, label, ref_meta, me.difficulty, log, alt=False)

    want_mythic = params.get("mythic") is not False and str(params.get("mythic")).lower() not in ("0", "false")
    if (client is not None and not params.get("against") and want_mythic
            and me.difficulty and int(me.difficulty) != MYTHIC):
        log("Эталон 2 — эпохальный топ: сравнение с самой высокой сложностью…")
        meta2: dict = {}
        try:
            tops2, label2 = _player_ref(client, me, MYTHIC, params, log, meta2)
            main["alt"] = _player_result(job, params, client, me, tops2, label2, meta2, MYTHIC, log, alt=True)
        except Exception as e:  # noqa: BLE001 — второй эталон не обязателен
            log(f"Эпохальный эталон не собран: {e}")
    job["result"] = main
    log("Готово.")


class _SubJob(dict):
    """Часть разбора «Разобрать бой» (рейд, ротация, игрок…): своя шкала прогресса внутри общей.
    Пишет лог и ожидание лимита в общую задачу; Excel не собирает — только запоминает листы."""

    def __init__(self, parent: dict, lo: float, hi: float):
        super().__init__(id=parent["id"], progress=0.0, log=parent["log"], state="running")
        self.parent, self.lo, self.hi, self.writers = parent, lo, hi, []

    def __setitem__(self, k, v):
        super().__setitem__(k, v)
        if k == "progress":
            self.parent["progress"] = max(self.parent["progress"], self.lo + (self.hi - self.lo) * min(1.0, float(v)))
        elif k == "wait_until":
            self.parent[k] = v


FIGHT_PARTS = {"raid": "Урон, сейвы, смерти", "rot": "Ротация рейда", "player": "Игрок",
               "saves": "Сейвы по боссам", "pall": "Ротация по боссам"}


def _run_fight(job: dict, params: dict, creds, log) -> None:
    """«Разобрать бой» — одна кнопка вместо трёх. Один бой: разбор рейда (урон, сейвы, смерти, механики),
    затем ротация — всех DPS (Игрок = «Все игроки») или одного игрока. Все боссы отчёта (fight="all"):
    план сейвов на каждого босса + ротация по боссам против топ-1. Сначала считается дешёвая часть —
    браузер показывает её сразу (job["partial"]), ротация досчитывается следом. Кончился лимит WCL на
    ротации — готовые вкладки остаются, во вкладке ротации — причина и «Продолжить». Excel — один на всё."""
    allb = params.get("fight") == "all"
    ref = params.get("ref") or "top3"
    actor = params.get("actor")
    demo = bool(params.get("demo"))
    parts: dict = {}
    errors: dict = {}
    writers: list = []
    first = ("saves" if allb else "raid")
    second = None if ref == "none" else ("pall" if allb else ("rot" if actor in (None, "", "all") or demo else "player"))
    order = ([] if allb else [first]) + ([second] if second else [])
    titles: dict = dict(FIGHT_PARTS)
    info_box: dict = {}   # все боссы: шапка результата — отчёт целиком (вкладки — боссы)
    kinds: dict = {}
    queued: list = []

    def combined(pending=None) -> dict:
        base = parts.get(first) or {}
        return {"mode": "fight", "all_bosses": allb, "info": info_box or base.get("info") or {}, "parts": dict(parts),
                "order": list(order), "errors": dict(errors), "pending": pending, "titles": dict(titles),
                "kinds": dict(kinds), "queued": [k for k in queued if k not in parts and k not in errors and k != pending],
                "focus": actor, "ref": ref, "excel": f"/api/report/{job['id']}", "source_url": params.get("url"),
                "params": {k: params.get(k) for k in ("url", "fight", "actor", "ref", "pick", "mythic", "demo")}}

    def publish(pending) -> None:
        job["partial"] = combined(pending=pending)
        job["partial_v"] = job.get("partial_v", 0) + 1

    def run_part(key, fn, sub_params, lo, hi) -> None:
        sj = _SubJob(job, lo, hi)
        log(f"— {FIGHT_PARTS[key]} —")
        try:
            fn(sj, sub_params, creds, log)
        except SystemExit:
            raise
        except Exception as e:  # noqa: BLE001 — одна часть не должна ронять уже готовые
            if key == first:
                raise
            errors[key] = _friendly(e)
            log(f"{FIGHT_PARTS[key]}: не досчитано — {errors[key]}")
            return
        res = sj.get("result") or {}
        res["excel"] = f"/api/report/{job['id']}"
        if (res.get("alt") or {}).get("excel"):
            res["alt"]["excel"] = f"/api/report/{job['id']}/alt"
        for pl in res.get("players") or []:  # подробности игрока внутри ротации — тот же общий Excel
            if isinstance(pl.get("detail"), dict):
                pl["detail"]["excel"] = res["excel"]
        for k, name, writer, data in sj.writers:
            if k == "xlsx":
                writers.append((name, writer, data, key))
            else:  # второй (эпохальный) эталон игрока — отдельный файл, как и раньше
                _excel_bytes(job, name, writer, data, key=k)
        parts[key] = res

    if allb:
        _fight_all_raid(job, params, creds, log, parts, errors, writers, order, titles, kinds, queued, second,
                        publish, info_box, 0.0, 0.6 if second else 0.95)
    else:
        run_part("raid", _run_raid, {**params, "mode": "raid"}, 0.0, 0.3 if second else 0.95)
    if second:
        publish(second)
        if second == "pall":
            run_part("pall", _run_player_all, {**params, "mode": "allbosses", "actor": actor or "all"}, 0.6, 0.95)
        elif second == "rot":
            run_part("rot", _run_raid_rotation, {**params, "mode": "raidrot", "ref": ref}, 0.3, 0.95)
        else:
            run_part("player", _run_player, {**params, "ref": ref}, 0.3, 0.95)
    if writers:  # один Excel на весь разбор: листы рейда, затем ротации
        short = {"raid": "Рейд", "rot": "Ротация", "player": "Игрок", "saves": "Сейвы", "pall": "Ротация"}

        def write_all(_data, path):
            wb = _new_wb()
            for _name, writer, data, key in writers:
                before = set(wb.sheetnames)
                writer(data, None, wb=wb)
                for ws in wb.worksheets:  # одноимённый лист второй части: «Выжимка1» → «Игрок — Выжимка»
                    if ws.title in before:
                        continue
                    m = re.match(r"(.+?)(\d+)$", ws.title)
                    if key in kinds:  # бой каждого босса — листы с его именем: «Sszorak — Выжимка»
                        base = m.group(1) if m and m.group(1) in before else ws.title
                        ws.title = _sheet_prefix(titles[key], key) + f" — {base}"[:31 - len(_sheet_prefix(titles[key], key))]
                    elif m and m.group(1) in before:
                        ws.title = f"{short.get(key, key)} — {m.group(1)}"[:31]
            wb.save(path)
            return path
        name = writers[0][0] + ("_и_ротация" if any(w[3] in ("rot", "player", "pall") for w in writers) else "")
        _excel_bytes(job, name, write_all, None)
    job.pop("partial", None)
    job["result"] = combined()
    log("Готово.")


def _fight_all_raid(job, params, creds, log, parts, errors, writers, order, titles, kinds, queued, second,
                    publish, info_box, lo, hi) -> None:
    """«Все боссы отчёта»: на каждого босса — полный разбор боя, как у одиночного лога (главное по бою,
    кому что поправить, урон, сейвы, смерти, механики, план с двумя заметками MRT), вкладка на босса.
    Готовые боссы
    видны сразу; кончился лимит WCL — остальные боссы с «Продолжить» (скачанное повторно лимит не тратит)."""
    from .config import DIFFICULTY_NAMES, SITE_URL
    from .excel_raid import write_raid_workbook
    from .logs import parse_report_url
    from .raid import run_raid
    from .raid_saves import pick_fights

    avoidable = set(_avoidable_list())
    if params.get("demo"):
        from .raid_demo import DEMO_AVOIDABLE, DEMO_URL, FakeRaidClient
        client, url = FakeRaidClient(), DEMO_URL
        avoidable |= {str(x) for x in DEMO_AVOIDABLE}
        log("Демо-рейд: один босс — на настоящем отчёте вкладка будет на каждого босса вечера.")
    else:
        client, url = _client(creds, job, log, wait=bool(params.get("wait"))), params["url"]
    code, _, _ = parse_report_url(url)
    report = client.report(code)
    chosen = pick_fights(report, difficulties=None)
    if not chosen:
        raise LookupError("В отчёте нет боёв с боссами")
    multi = len({int(c["fight"].get("difficulty") or 0) for c in chosen}) > 1
    keys = [f"b{i}" for i in range(len(chosen))]
    for k, c in zip(keys, chosen):
        f = c["fight"]
        titles[k] = f.get("name", "") + (f" ({DIFFICULTY_NAMES.get(int(f.get('difficulty') or 0), '')[:4]}.)" if multi else "")
        kinds[k] = "raid"
    order[0:0] = keys   # вкладки боссов, затем ротация
    queued[:] = keys + ([second] if second else [])
    zone = (report.get("zone") or {}).get("name", "")
    info_box.update({"code": code, "title": report.get("title", ""), "zone": zone, "url": f"{SITE_URL}/reports/{code}",
                     "boss": zone or report.get("title", ""), "bosses": len(chosen),
                     "difficulty": ", ".join(dict.fromkeys(DIFFICULTY_NAMES.get(int(c["fight"].get("difficulty") or 0), "")
                                                           for c in chosen)), "demo": code.startswith("DEMO")})
    log(f"Боссов в отчёте: {len(chosen)}. На каждого — полный разбор боя (последний килл, без киллов — лучший пулл).")
    n = len(chosen)
    for i, (k, c) in enumerate(zip(keys, chosen)):
        f = c["fight"]
        diff = DIFFICULTY_NAMES.get(int(f.get("difficulty") or 0), "")
        publish(k)
        log(f"[{i + 1}/{n}] {f.get('name')} ({diff}): " + ("килл" if f.get("kill") else "лучший пулл"))
        base = lo + (hi - lo) * i / n
        try:
            R = run_raid(client, url, int(f["id"]), log=lambda m: log("    " + m), avoidable=avoidable,
                         talent_data=[] if params.get("demo") else None, save_talents=not SERVER["public"],
                         mythic=params.get("mythic") is not False,
                         progress=lambda x, b=base: job.__setitem__(
                             "progress", max(job["progress"], min(0.97, b + (hi - lo) * min(1.0, x) / n))))
        except SystemExit:
            raise
        except Exception as e:  # noqa: BLE001 — один босс не должен ронять остальные
            msg = _friendly(e)
            if "лимит" in msg.lower():
                for k2 in keys[i:]:
                    errors[k2] = "Не хватило часового лимита WCL — нажмите «Продолжить» после сброса."
                log("Закончился часовой лимит WCL — остальные боссы после сброса.")
                break
            errors[k] = msg
            log(f"    пропущен: {msg}")
            continue
        parts[k] = {**R, "excel": f"/api/report/{job['id']}", "source_url": url}
        writers.append((f"Все_боссы_{zone or code}", write_raid_workbook, R, k))
    if not any(k in parts for k in keys) and not any("лимит" in e.lower() for e in errors.values()):
        raise LookupError("Ни один бой не удалось разобрать: " + "; ".join(dict.fromkeys(errors.values())))
    # кончился лимит уже на первом боссе — не ошибка: вкладки боссов с «Продолжить»


def _prune_cache() -> None:
    try:
        r = _shared_cache().prune()
        if r["removed"] or r["before_mb"] - r["after_mb"] > 1:
            print(f"Кэш WCL: удалено старых записей — {r['removed']}, размер {r['before_mb']} → {r['after_mb']} МБ")
    except Exception as e:  # noqa: BLE001 — уборка кэша необязательна
        print(f"Кэш WCL не очищен: {e}")


def _cache_file() -> Path:
    from .config import cache_path
    return Path(cache_path()).resolve()


def _sheet_prefix(title: str, key: str) -> str:
    """Начало имени листа для боя босса: номер вкладки + имя с сокращением сложности. Номер делает имя
    уникальным, даже если один босс есть на двух сложностях (обрезка до 31 символа их не склеит)."""
    m = re.match(r"b(\d+)$", key or "")
    num = f"{int(m.group(1)) + 1}." if m else ""
    diff = re.search(r"\((\w{2,5})\.?\)\s*$", title or "")
    name = re.sub(r"\s*\([^)]*\)\s*$", "", title or "").strip()
    tag = f" {diff.group(1)[:3]}" if diff else ""
    return f"{num}{name[:12 - len(tag)]}{tag}".strip()


def _new_wb():
    from openpyxl import Workbook
    wb = Workbook()
    wb.remove(wb.active)
    return wb


ALL_BOSSES_POINTS = 25.0   # оценка очков WCL на одного босса до первого замера (свой лог + лог топ-1)


def _run_player_all(job: dict, params: dict, creds, log) -> None:
    """Разбор на всех боссах отчёта: выбранный игрок или все DPS (actor="all"). Эталон — топ-1 (ради
    экономии запросов), без эпохального второго эталона. Если часового лимита WCL не хватает — разбор
    останавливается, а оставшиеся пары «босс — игрок» возвращаются списком «pending»: браузер покажет
    кнопку «Продолжить» (уже скачанное берётся из кэша и лимит не тратит)."""
    from .api import WCLError
    from .collect import load_my_log
    from .config import DIFFICULTY_NAMES
    from .excel_report import write_player_all_workbook
    from .logs import parse_report_url

    from .collect import shared_cache
    wait = bool(params.get("wait"))
    client = _client(creds, job, log, wait=wait)  # без галочки — не ждём сброса: останавливаемся и предлагаем продолжить
    shared = shared_cache(client)  # общие данные боя — один раз на всех игроков этого боя
    ranks_cache: dict = {}

    def ranks_of(fid: int) -> dict:
        """Рейтинги WCL игроков боя (процентиль и с учётом экипировки) — один запрос на бой, только киллы."""
        if fid not in ranks_cache:
            from .collect import player_percentiles
            try:
                ranks_cache[fid] = player_percentiles(client.report_rankings(code, fid)) \
                    if fights[fid].get("kill") and hasattr(client, "report_rankings") else {}
            except Exception:  # noqa: BLE001 — рейтинги необязательны
                ranks_cache[fid] = {}
        return ranks_cache[fid]
    code, _, _ = parse_report_url(params["url"])
    report = client.report(code)
    everyone = str(params.get("actor")) == "all"
    fights = {int(f["id"]): f for f in _boss_fights(report)}
    if not fights:
        raise LookupError("В отчёте нет боёв с боссами")
    if params.get("units"):  # «Продолжить»: только оставшиеся пары
        units = [(int(fid), int(aid) if aid not in (None, "") else None) for fid, aid in params["units"]
                 if int(fid) in fights]
    elif everyone:
        units = [(fid, aid) for fid in fights for aid in _dps_ids(client, code, fid)]
    else:
        aid = int(params["actor"]) if params.get("actor") not in (None, "") else None
        units = [(fid, aid) for fid in fights]
    p1 = {**params, "ref": "top1", "mythic": False}
    log(f"Боссов: {len(fights)}{f', разборов «босс — игрок»: {len(units)}' if everyone else ''}. "
        "Эталон — топ-1 на каждом; общие данные боя качаются один раз на всех игроков (экономия запросов WCL).")
    if wait:
        log("Если лимит кончится — жду сброса и продолжаю сам.")
    rows, skipped, pending = [], [], []
    per_unit: list[float] = []
    for i, (fid, aid) in enumerate(units):
        f = fights[fid]
        label = f"{f.get('name')} ({DIFFICULTY_NAMES.get(int(f.get('difficulty') or 0), '')})"
        left = client.points_left() if hasattr(client, "points_left") else None
        need = (sum(per_unit) / len(per_unit)) if per_unit else ALL_BOSSES_POINTS
        if not wait and left is not None and left < need * 1.2:
            pending = units[i:]
            log(f"Лимита WCL не хватит на следующий разбор: осталось {left:.0f} очков, на один уходит ≈ {need:.0f}.")
            break
        log(f"[{i + 1}/{len(units)}] {label}{f', игрок #{aid}' if everyone else ''}…")
        try:
            # без событий нанесённого урона: экономия лимита (нет только оценки прироста от кулдаунов)
            me = load_my_log(client, params["url"], fid, actor_id=aid, shared=shared, damage_events=False)
            tops, ref_label = _player_ref(client, me, me.difficulty, p1, lambda m: log("    " + m), {})
            res = _player_result(job, p1, client, me, tops, ref_label, {}, me.difficulty, lambda m: log("    " + m), alt=False,
                                 excel=False)
            res.pop("ref", None)
            rk = ranks_of(fid).get(me.name) or ranks_of(fid).get(me.actor_id) or {}
            rows.append({"fight_id": fid, "actor": me.actor_id, "player": me.name, "cls": me.cls, "spec": me.spec,
                         "rank": rk.get("rank"), "ilvl_pct": rk.get("ilvl"),
                         "boss": me.encounter_name, "difficulty": me.difficulty_name, "kill": me.kill,
                         "duration": _fmt_t(me.duration), "dps": round(me.dps), "ref_dps": res["info"].get("ref_dps"),
                         "ref_label": ref_label, "actions": (res.get("brief") or {}).get("actions", [])[:3], "detail": res})
        except WCLError as e:
            if "лимит" in str(e).lower():
                pending = units[i:]
                log(f"    Закончился часовой лимит WCL — остановился на «{f.get('name')}».")
                break
            skipped.append({"boss": label, "reason": str(e)})
            log(f"    пропущен: {e}")
        except Exception as e:  # noqa: BLE001 — один разбор не должен ронять остальные
            skipped.append({"boss": label, "reason": str(e)})
            log(f"    пропущен: {e}")
        after = client.points_left() if hasattr(client, "points_left") else None
        if left is not None and after is not None and left - after > 0:
            per_unit.append(left - after)
        job["progress"] = max(job["progress"], min(0.95, (i + 1) / len(units)))
    reset_in = None
    if pending:
        try:
            reset_in = float(client.rate_limit().get("pointsResetIn") or 0)
        except Exception:  # noqa: BLE001
            pass
    first = next((r["detail"]["info"] for r in rows), {})
    R = {"mode": "playerall", "everyone": everyone,
         "info": {"name": "Все игроки" if everyone else first.get("name", ""),
                  "cls": "" if everyone else first.get("cls", ""), "spec": "" if everyone else first.get("spec", ""),
                  "title": report.get("title", ""), "code": code, "url": f"https://www.warcraftlogs.com/reports/{code}",
                  "demo": False},
         "bosses": rows, "skipped": skipped,
         "pending": [{"fight_id": fid, "actor": aid, "boss": fights[fid].get("name", ""),
                      "difficulty": DIFFICULTY_NAMES.get(int(fights[fid].get("difficulty") or 0), "")} for fid, aid in pending],
         "reset_in": reset_in, "params": {"url": params["url"], "actor": params.get("actor")}}
    _excel_bytes(job, f"Все_боссы_{R['info']['name'] or code}", write_player_all_workbook,
                 {**R, "prev": params.get("prev") or []})
    job["result"] = {**R, "excel": f"/api/report/{job['id']}"}
    log("Готово." if not pending else f"Готово частично: осталось разборов — {len(pending)}. "
        "Нажмите «Продолжить», когда лимит восстановится.")


def _dps_ids(client, code: str, fid: int) -> list[int]:
    """DPS-игроки боя (танки и лекари — не по рейтингу урона, их сравнение с топом по DPS не имеет смысла)."""
    det = client.player_details(code, fid) or {}
    return [int(p["id"]) for p in det.get("dps") or [] if "id" in p]


def _boss_fights(report: dict) -> list[dict]:
    """По одному бою на босса и сложность: последний килл, а без киллов — лучший пулл."""
    groups: dict = {}
    for f in sorted(report.get("fights") or [], key=lambda x: float(x.get("startTime") or 0)):
        if int(f.get("encounterID") or 0):
            groups.setdefault((int(f["encounterID"]), int(f.get("difficulty") or 0)), []).append(f)
    out = []
    for fs in groups.values():
        kills = [f for f in fs if f.get("kill")]
        out.append(kills[-1] if kills else min(fs, key=lambda f: (f.get("fightPercentage") is None,
                                                                   f.get("fightPercentage") or 100.0)))
    return out


def _player_ref(client, me, difficulty: int, params: dict, log, meta: dict):
    from .collect import collect_reference
    from .config import DIFFICULTY_NAMES
    n = REF_SIZES.get(params.get("ref") or "top3", 3)
    log(f"Собираю эталон: топ-{n} {spec_ru(me.cls, me.spec)} на этом боссе, сложность — "
        f"{DIFFICULTY_NAMES.get(int(difficulty or 0), difficulty)}. Первый раз это занимает несколько минут, дальше быстрее.")
    if params.get("refresh"):
        log("Обновляю рейтинг лучших логов…")
    tops, label = collect_reference(client, me.encounter_id, me.cls, me.spec, difficulty,
                                    top_n=n, duration=me.duration, log=log,
                                    force=bool(params.get("refresh")), meta=meta,
                                    max_age_s=_max_age_s(params), save=not SERVER["public"])
    return tops, label


def _player_result(job, params, client, me, tops, label, ref_meta, top_diff, log, alt: bool, excel: bool = True) -> dict:
    from .compare import compare
    from .config import DIFFICULTY_NAMES
    from .excel_report import write_compare_workbook
    from .metrics import load_overrides
    from .reference import build_reference
    from .talents import compare_talents, demo_tree_data, load_tree_data
    if client is not None:
        from .raid_rotation import report_points
        report_points(client, log)
    log("Считаю эталон и сравниваю…" + (" (эпохальный топ)" if alt else ""))
    ref = build_reference(tops, load_overrides(_spell_meta_path()), label=label, me=me)
    r = compare(me, ref)
    try:  # по всем скачанным логам топа, а не только по отобранным с похожим билдом
        data = demo_tree_data() if params.get("demo") else load_tree_data(log, save=not SERVER["public"])
        r.talents = compare_talents(me, tops, data=data, log=log)
        if r.talents and r.talents.get("error"):
            log("Таланты: " + r.talents["error"])
    except Exception as e:  # noqa: BLE001 — сравнение талантов не обязательно
        log(f"Сравнение талантов недоступно: {e}")
        r.talents = None
    diff_name = DIFFICULTY_NAMES.get(int(top_diff or 0), "")
    suffix = "_эпохальный_топ" if alt else ""
    if excel:  # в «Все боссы» — общий Excel на всех игроков, отдельные книги на каждого не нужны
        _excel_bytes(job, f"{me.name}_{me.encounter_name}{suffix}", write_compare_workbook, r,
                     key="xlsx_alt" if alt else "xlsx")
    result = to_json(r, job["id"])
    if alt:
        result["excel"] = f"/api/report/{job['id']}/alt"
    result["info"]["ref_difficulty"] = diff_name
    if ref_meta:
        result["ref"] = {k: ref_meta.get(k) for k in REF_FIELDS}
        result["info"]["ref_collected_at"] = ref_meta["collected_at"]
        result["info"]["ref_stale"] = time.time() - ref_meta["collected_at"] > _max_age_s(params)
    result["params"] = {k: params.get(k) for k in ("url", "fight", "actor", "ref")}
    result["talents"] = r.talents
    return result


def _avoidable_list() -> list[str]:
    """Список избегаемых механик из spell_meta.json: "_avoidable": [id или название, ...]."""
    import json
    try:
        data = json.loads(Path(_spell_meta_path()).read_text(encoding="utf-8"))
        return [str(x) for x in data.get("_avoidable", [])]
    except (OSError, ValueError, AttributeError):
        return []


def _spell_meta_path() -> str:
    """spell_meta.json рядом с программой или в каталоге данных (для .exe и Android)."""
    from .config import data_dir
    for p in (Path("spell_meta.json"), data_dir() / "spell_meta.json"):
        if p.exists():
            return str(p)
    return "spell_meta.json"


REF_FIELDS = ("encounter_id", "boss", "cls", "spec", "difficulty", "top_n", "duration", "label", "n_logs",
              "collected_at")


def _run_refresh(job: dict, params: dict, creds, log) -> None:
    """Обновляет эталоны, присланные браузером: свежий рейтинг и бои новых игроков топа."""
    from .collect import collect_reference
    client = _client(creds, job, log)
    refs = []
    for r in (params.get("refs") or [])[:30]:
        try:
            refs.append({"encounter_id": int(r["encounter_id"]), "cls": str(r["cls"])[:40],
                         "spec": str(r["spec"])[:40], "difficulty": int(r["difficulty"]),
                         "top_n": max(1, min(50, int(r["top_n"]))),
                         "duration": float(r["duration"]) if r.get("duration") else None,
                         "boss": str(r.get("boss", ""))[:80]})
        except (KeyError, TypeError, ValueError):
            continue
    out = []
    for i, r in enumerate(refs):
        log(f"[{i}/{len(refs)}] {r['boss']}: {spec_ru(r['cls'], r['spec'])}, топ-{r['top_n']}…")
        meta: dict = {}
        collect_reference(client, r["encounter_id"], r["cls"], r["spec"], r["difficulty"], top_n=r["top_n"],
                          duration=r["duration"], force=True, log=lambda m: log("  " + m.strip()),
                          meta=meta, max_age_s=_max_age_s(params), save=not SERVER["public"])
        out.append({k: meta.get(k) for k in REF_FIELDS})
    job["result"] = {"mode": "refresh", "count": len(out), "refs": out}
    log(f"Готово: обновлено эталонов — {len(out)}.")


def _run_raid(job: dict, params: dict, creds, log) -> None:
    from .excel_raid import write_raid_workbook
    from .raid import run_raid

    def step(msg: str) -> None:
        log(msg)
        job["progress"] = min(0.5, job["progress"] + 0.07)

    avoidable = set(_avoidable_list())
    if params.get("demo"):
        from .raid_demo import DEMO_AVOIDABLE, DEMO_URL, FakeRaidClient
        step("Генерирую демо-рейд: 20 игроков, 6 пуллов…")
        client, url = FakeRaidClient(), DEMO_URL
        avoidable |= {str(x) for x in DEMO_AVOIDABLE}
    else:
        client, url = _client(creds, job, log), params["url"]
    R = run_raid(client, url, params.get("fight"), log=step, avoidable=avoidable,
                 talent_data=[] if params.get("demo") else None, save_talents=not SERVER["public"],
                 mythic=params.get("mythic") is not False,
                 progress=lambda x: job.__setitem__("progress", max(job["progress"], min(0.97, x))))
    _excel_bytes(job, f"Рейд_{R['info']['boss']}_пулл{R['summary']['pull_n']}", write_raid_workbook, R)
    job["result"] = {**R, "excel": f"/api/report/{job['id']}", "source_url": url}
    log("Готово.")


def _run_saves(job: dict, params: dict, creds, log) -> None:
    """План рейдовых сейвов на всех боссов отчёта (по одному бою на босса)."""
    from .excel_raid import write_saves_workbook
    from .raid_saves import run_raid_saves

    if params.get("demo"):
        from .raid_demo import DEMO_URL, FakeRaidClient
        log("Демо-рейд: один босс — на настоящем отчёте план будет на каждого босса вечера.")
        client, url = FakeRaidClient(), DEMO_URL
    else:
        client, url = _client(creds, job, log, wait=bool(params.get("wait"))), params["url"]
    R = run_raid_saves(client, url, log=log, talent_data=[] if params.get("demo") else None,
                       save_talents=not SERVER["public"], avoidable=set(_avoidable_list()),
                       progress=lambda x: job.__setitem__("progress", max(job["progress"], min(0.97, x))))
    _excel_bytes(job, f"Сейвы_{R['info']['zone'] or R['info']['code']}", write_saves_workbook, R)
    job["result"] = {**R, "excel": f"/api/report/{job['id']}", "source_url": url}
    log("Готово.")


def _run_roster_plan(job: dict, params: dict, creds, log) -> None:
    """План сейвов по составу из игры (WoWUtils Group Export) и длительности боя — по лучшим киллам топа."""
    from .excel_raid import write_roster_plan_workbook
    from .roster_plan import run_roster_plan
    roster = str(params.get("roster") or "")[:200_000]
    try:
        minutes = max(1.0, min(20.0, float(str(params.get("minutes") or "5").replace(",", "."))))
    except ValueError:
        raise LookupError("Длительность боя — число минут, например 5 или 6,5") from None
    if params.get("encounter") in (None, "", "all"):
        raise LookupError("Выберите босса — план составляется на одного босса")
    client = _client(creds, job, log, wait=bool(params.get("wait")))
    R = run_roster_plan(client, int(params["encounter"]), int(params.get("difficulty") or 5), minutes, roster, log=log,
                        progress=lambda x: job.__setitem__("progress", max(job["progress"], min(0.97, x))))
    _excel_bytes(job, f"План_по_составу_{R['info']['boss']}", write_roster_plan_workbook, R)
    job["result"] = {**R, "excel": f"/api/report/{job['id']}"}
    log("Готово.")


FIRST_KILL_PAGES = 3   # героик и обычный: самый ранний килл ищем среди 3 страниц рейтинга WCL (300 киллов)


def _basis_text(diff) -> str:
    return ("рейтинг WCL по прогрессу (первые киллы в мире)…" if int(diff or 0) == 5 else
            f"самый ранний по дате среди {FIRST_KILL_PAGES * 100} киллов рейтинга WCL (рейтинга прогресса для этой сложности нет)…")


def _progress_cands(client, enc: int, diff: int) -> list[dict]:
    """Чей бой показывать. Эпохальная — рейтинг WCL по прогрессу (первые киллы в мире). Для героической и
    обычной сложности такого рейтинга нет — берём самый ранний килл по дате среди FIRST_KILL_PAGES страниц
    рейтинга киллов (basis="first_kill"); дат в ответе нет — самый быстрый килл (basis="fastest")."""
    from .config import SITE_URL

    def cand(rk, place, basis):
        rep, g, srv = rk.get("report") or {}, rk.get("guild") or {}, rk.get("server") or {}
        return {"url": f"{SITE_URL}/reports/{rep['code']}", "fight": int(rep.get("fightID") or rep.get("fightId") or 0),
                "guild": g.get("name") or rk.get("name") or rep["code"], "rank": place, "basis": basis,
                "server": srv.get("name"), "region": srv.get("region"), "start": _kill_time(rk)}

    ranks = [] if int(diff or 0) != 5 else [r for r in client.fight_rankings(enc, diff, "progress") if (r.get("report") or {}).get("code")]
    if ranks:
        return [cand(rk, i + 1, "progress") for i, rk in enumerate(ranks[:10])]
    pool = []
    for page in range(1, FIRST_KILL_PAGES + 1):
        try:
            got = client.fight_rankings(enc, diff, "speed", page=page)
        except Exception:  # noqa: BLE001 — следующая страница не обязательна
            break
        pool += [r for r in got if (r.get("report") or {}).get("code")]
        if len(got) < 100:
            break
    if not pool:
        return []
    if all(_kill_time(r) for r in pool):
        pool.sort(key=_kill_time)
        return [cand(rk, i + 1, "first_kill") for i, rk in enumerate(pool[:10])]
    return [cand(rk, i + 1, "fastest") for i, rk in enumerate(pool[:10])]


def _kill_time(rk: dict) -> float | None:
    t = rk.get("startTime") or (rk.get("report") or {}).get("startTime")
    try:
        return float(t) if t else None
    except (TypeError, ValueError):
        return None


CLOSED_HINTS = ("закрыт", "не найден", "permission", "private", "not found", "does not exist", "403", "unauthorized",
                "нет доступа")
PROGRESS_TRIES = 6   # сколько первых мест пробуем, если логи закрыты


def _progress_one(job, client, cands: list[dict], params: dict, log, lo: float, hi: float,
                  difficulty: int | None = None) -> dict:
    """Разбор боя лучшей гильдии из cands без сравнения. Закрытый, удалённый или битый лог — следующее место
    (до PROGRESS_TRIES). Какие места пропущены и почему — в info.top_progress.skipped, а если не открылся
    ни один — понятная ошибка со списком мест и причин."""
    from .api import WCLError
    from .raid import run_raid
    tried = []
    for c in cands[:PROGRESS_TRIES]:
        log(f"Место {c['rank']}: {c['guild']} — разбираю их бой…")
        try:
            R = run_raid(client, c["url"], c["fight"], log=lambda m: log("    " + m), compare_top=False, mythic=False,
                         talent_data=[] if params.get("demo") else None, save_talents=not SERVER["public"],
                         avoidable=set(_avoidable_list()),
                         progress=lambda x: job.__setitem__("progress", max(job["progress"], min(hi, lo + (hi - lo) * x))))
        except Exception as e:  # noqa: BLE001
            if isinstance(e, WCLError) and "лимит" in str(e).lower():
                raise
            msg = str(e)
            reason = "лог закрыт или удалён" if any(h in msg.lower() for h in CLOSED_HINTS) else _friendly(e)
            tried.append({"rank": c["rank"], "guild": c["guild"], "reason": reason})
            log(f"    недоступен: {reason}")
            continue
        R["info"]["top_progress"] = {**{k: c.get(k) for k in ("rank", "guild", "server", "region", "start", "basis")},
                                     "skipped": tried}
        return {**R, "source_url": c["url"]}
    places = "; ".join(f"{t['rank']}-е место — {t['guild']}: {t['reason']}" for t in tried)
    hint = (" На эпохальной сложности гильдии во время гонки за первый килл часто прячут логи — попробуйте "
            "героическую сложность или зайдите позже, когда логи откроют.") if int(difficulty or 0) == 5 else \
        " Попробуйте позже — гильдии иногда открывают логи не сразу."
    raise LookupError(f"Логи первых гильдий недоступны — проверено мест: {len(tried)}. {places}.{hint}")


def _run_top_progress(job: dict, params: dict, creds, log) -> None:
    """«Топ прогресса»: бой гильдии, первой убившей этого босса на этой сложности (рейтинг WCL по прогрессу),
    разобранный как обычный бой рейда — урон, сейвы и кто какие кулдауны жал, смерти, состав. Без сравнения
    с вашим боем и с другими киллами: просто посмотреть, как это сделали лучшие.
    «Все боссы отчёта» — так же для каждого босса отчёта, вкладка на босса."""
    from .config import DIFFICULTY_NAMES
    from .excel_raid import write_raid_workbook
    from .logs import parse_report_url

    if params.get("demo"):
        from .raid_demo import DEMO_URL, FakeRaidClient
        client = FakeRaidClient()
        R = _progress_one(job, client, [{"url": DEMO_URL, "fight": None, "guild": "Демо-гильдия", "rank": 1}], params, log, 0.1, 0.95)
        _excel_bytes(job, f"Топ_прогресса_{R['info']['boss']}", write_raid_workbook, R)
        job["result"] = {**R, "excel": f"/api/report/{job['id']}"}
        log("Готово.")
        return
    client = _client(creds, job, log, wait=bool(params.get("wait")))
    if params.get("zone"):  # без лога: рейд, сложность и босс выбраны в списке
        zone = next((z for z in client.raid_zones() if str(z["id"]) == str(params["zone"])), None)
        if not zone:
            raise LookupError("Рейд не найден в Warcraft Logs — обновите страницу и выберите снова")
        diff = int(params.get("difficulty") or 5)
        encs = zone["encounters"] if params.get("encounter") in (None, "", "all") else \
            [e for e in zone["encounters"] if str(e["id"]) == str(params["encounter"])]
        if not encs:
            raise LookupError("Босс не найден в этом рейде")
        bosses = [(int(e["id"]), diff, e["name"]) for e in encs]
        if len(bosses) > 1:
            return _run_top_progress_all(job, params, client, bosses, zone["name"], log)
        enc, _, name = bosses[0]
        log(f"{name} ({DIFFICULTY_NAMES.get(diff, '')}): ищу первый килл — " + _basis_text(diff))
        cands = _progress_cands(client, enc, diff)
        if not cands:
            raise LookupError(f"{name}: на этой сложности пока нет киллов в рейтингах Warcraft Logs")
        job["progress"] = 0.1
        R = _progress_one(job, client, cands, params, log, 0.1, 0.95, difficulty=diff)
        _excel_bytes(job, f"Топ_прогресса_{R['info']['boss']}_{R['info']['top_progress']['guild']}", write_raid_workbook, R)
        job["result"] = {**R, "excel": f"/api/report/{job['id']}"}
        log("Готово.")
        return
    code, url_fight, _ = parse_report_url(params["url"])
    report = client.report(code)
    fights = [f for f in report.get("fights") or [] if int(f.get("encounterID") or 0)]
    if not fights:
        raise LookupError("В отчёте нет боёв с боссами — не понятно, чей топ прогресса искать")
    if params.get("fight") == "all":
        seen, bosses = set(), []
        for f in sorted(fights, key=lambda x: float(x.get("startTime") or 0)):
            k = (int(f["encounterID"]), int(f.get("difficulty") or 0))
            if k not in seen:
                seen.add(k)
                bosses.append((k[0], k[1], f.get("name", "")))
        title = (report.get("zone") or {}).get("name") or report.get("title") or "Все боссы"
        return _run_top_progress_all(job, params, client, bosses, title, log)
    fid = params.get("fight") if params.get("fight") not in (None, "") else url_fight
    f = next((f for f in fights if str(f["id"]) == str(fid)), None) or fights[-1]
    enc, diff = int(f["encounterID"]), int(f.get("difficulty") or 0)
    log(f"{f.get('name')} ({DIFFICULTY_NAMES.get(diff, '')}): ищу первый килл — " + _basis_text(diff))
    cands = _progress_cands(client, enc, diff)
    if not cands:
        raise LookupError("У этого босса на этой сложности пока нет киллов в рейтингах Warcraft Logs")
    job["progress"] = 0.1
    R = _progress_one(job, client, cands, params, log, 0.1, 0.95, difficulty=diff)
    _excel_bytes(job, f"Топ_прогресса_{R['info']['boss']}_{R['info']['top_progress']['guild']}", write_raid_workbook, R)
    job["result"] = {**R, "excel": f"/api/report/{job['id']}"}
    log("Готово.")


def _run_top_progress_all(job: dict, params: dict, client, bosses: list[tuple], report_title: str, log) -> None:
    """Топ прогресса на каждого босса отчёта (по сложностям из отчёта): вкладка на босса, готовые боссы
    видны сразу, пока считаются следующие. Кончился лимит WCL — остальные боссы с «Продолжить»
    (уже разобранное лимит повторно не тратит). Excel — один, листы подписаны именем босса."""
    from .config import DIFFICULTY_NAMES
    from .excel_raid import write_raid_workbook
    multi = len({b[1] for b in bosses}) > 1
    order = [f"b{i}" for i in range(len(bosses))]
    titles = {f"b{i}": name + (f" ({DIFFICULTY_NAMES.get(d, '')[:4]}.)" if multi else "") for i, (_, d, name) in enumerate(bosses)}
    parts, errors, writers = {}, {}, []
    log(f"Боссов: {len(bosses)}. На каждого — бой гильдии с первым киллом в мире.")

    def combined(pending=None):
        return {"mode": "fight", "variant": "progress_all", "all_bosses": True,
                "info": {"boss": report_title, "title": report_title, "zone": report_title}, "parts": dict(parts),
                "order": order, "errors": dict(errors), "pending": pending, "titles": titles,
                "kinds": {k: "raid" for k in order}, "queued": [k for k in order if k not in parts and k not in errors and k != pending],
                "excel": f"/api/report/{job['id']}",
                "params": {"mode": "progress", "url": params.get("url"), "fight": "all", "zone": params.get("zone"),
                           "encounter": params.get("encounter"), "difficulty": params.get("difficulty")}}
    for i, (enc, diff, name) in enumerate(bosses):
        key, lo = order[i], i / len(bosses)
        job["partial"], job["partial_v"] = combined(pending=key), job.get("partial_v", 0) + 1
        log(f"[{i + 1}/{len(bosses)}] {name} ({DIFFICULTY_NAMES.get(diff, '')})")
        try:
            cands = _progress_cands(client, enc, diff)
            if not cands:
                raise LookupError("на этой сложности пока нет киллов в рейтингах Warcraft Logs")
            R = _progress_one(job, client, cands, params, lambda m: log("  " + m), lo, lo + 1 / len(bosses), difficulty=diff)
        except Exception as e:  # noqa: BLE001
            msg = _friendly(e)
            if "лимит" in msg.lower():
                for k in order[i:]:
                    errors[k] = "Не хватило часового лимита WCL — нажмите «Продолжить» после сброса."
                log("Закончился часовой лимит WCL — остальные боссы после сброса.")
                break
            errors[key] = msg
            log(f"  {name}: {msg}")
            continue
        R["excel"] = f"/api/report/{job['id']}"
        parts[key] = R
        writers.append((titles[key], write_raid_workbook, R))
    if not parts and not any("лимит" in e.lower() for e in errors.values()):
        raise LookupError("Ни для одного босса не удалось разобрать топ прогресса: " + "; ".join(dict.fromkeys(errors.values())))

    def write_all(_data, path):
        wb = _new_wb()
        for wi, (prefix, writer, data) in enumerate(writers):
            before = set(wb.sheetnames)
            writer(data, None, wb=wb)
            for ws in wb.worksheets:  # листы каждого босса — с его именем: «Sszorak — Выжимка»
                if ws.title not in before:
                    base = re.sub(r"\d+$", "", ws.title) if re.sub(r"\d+$", "", ws.title) in before else ws.title
                    ws.title = _sheet_prefix(prefix, f"b{wi}") + f" — {base}"[:31 - len(_sheet_prefix(prefix, f"b{wi}"))]
        wb.save(path)
        return path
    _excel_bytes(job, f"Топ_прогресса_{report_title}", write_all, None)
    job.pop("partial", None)
    job["result"] = combined()
    log("Готово.")


def _report_code(url: str) -> str:
    from .logs import parse_report_url
    return parse_report_url(url)[0]


def _run_raid_rotation(job: dict, params: dict, creds, log) -> None:
    """Каждый DPS боя против топа своего спека; эталон собирается один раз на спек."""
    from .excel_report import write_raid_rotation_workbook
    from .metrics import load_overrides
    from .raid_rotation import run_demo, run_raid_rotation

    def progress(x: float) -> None:
        job["progress"] = max(job["progress"], min(0.95, x))

    n = REF_SIZES.get(params.get("ref") or "top3", 3)
    log(f"Эталон для каждого спека: топ-{n} (размер — в списке «Сравнить с»).")
    if params.get("demo"):
        R = run_demo(25, log=log, progress=progress)
    else:
        R = run_raid_rotation(_client(creds, job, log), params["url"], params.get("fight"), top_n=n, log=log,
                              progress=progress, overrides=load_overrides(_spell_meta_path()),
                              max_age_s=_max_age_s(params), save=not SERVER["public"], pick=params.get("pick"),
                              mythic=params.get("mythic") is not False)
    I = R["info"]
    _excel_bytes(job, f"Ротация_рейда_{I['boss']}", write_raid_rotation_workbook, R)
    players = []
    for r in R["rows"]:
        row = {k: (_num(v, 4) if isinstance(v, float) else v) for k, v in r.items() if k not in ("result", "actions")}
        if "result" in r:
            row["actions"] = r["actions"]
            row["detail"] = to_json(r["result"], job["id"])
        players.append(row)
    job["result"] = {"mode": "raidrot", "info": I, "brief": R["brief"], "players": players,
                     "skipped": [p.get("name") for p in R["skipped"]], "not_picked": R.get("not_picked", []),
                     "pick": R.get("pick"),
                     "excel": f"/api/report/{job['id']}", "source_url": params.get("url")}
    log("Готово.")


def _gear_json(g: dict) -> dict | None:
    if not g.get("rows"):
        return None
    return {"my_ilvl": _num(g.get("my_ilvl"), 1), "ref_ilvl": _num(g.get("ref_ilvl"), 1), "ref_n": g.get("ref_n", 0),
            "has_my": g.get("has_my", False), "missing": g.get("missing_enchants", []), "missing_temp": g.get("missing_temp", []),
            "rows": [{k: _num(v, 2) if isinstance(v, float) else v for k, v in x.items()} for x in g["rows"]]}


def _page() -> bytes:
    """Страница интерфейса с русскими названиями классов и спеков из names_ru.py (один словарь на всё)."""
    from .names_ru import CLASSES, SPECS
    html = STATIC.read_text(encoding="utf-8")
    html = html.replace("/*CLASS_RU*/{}", json.dumps(CLASSES, ensure_ascii=False))
    html = html.replace("/*SPEC_RU*/{}", json.dumps({f"{c}|{s}": v for (c, s), v in SPECS.items()}, ensure_ascii=False))
    html = html.replace("<script>", f'<script nonce="{PAGE_NONCE}">')
    return html.encode("utf-8")


# Политика содержимого страницы: выполняется только код самой программы (файлы /static и встроенный скрипт
# с одноразовой меткой). Даже если в данные (имя игрока, гильдии, файл резервной копии) попадёт чужой HTML
# со скриптом, браузер его не запустит.
import secrets as _secrets  # noqa: E402
PAGE_NONCE = _secrets.token_urlsafe(16)
CSP = ("default-src 'self'; script-src 'self' 'nonce-{n}'; style-src 'self' 'unsafe-inline' https://fonts.googleapis.com; "
       "font-src 'self' data: https://fonts.gstatic.com; img-src 'self' data: blob: https:; media-src 'self' blob: https:; "
       "connect-src 'self'; object-src 'none'; base-uri 'none'; form-action 'self'; frame-ancestors 'none'")


def _friendly(e: Exception) -> str:
    msg = str(e)
    if "ConnectionError" in type(e).__name__ or "Max retries" in msg:
        import re as _re
        host = (_re.search(r"host='([^']+)'", msg) or [None, ""])[1]
        if host and "warcraftlogs" not in host:
            return f"Нет связи с {host}. Проверьте интернет и попробуйте ещё раз."
        return "Нет связи с Warcraft Logs. Проверьте интернет и попробуйте ещё раз."
    if isinstance(e, (KeyError, IndexError, TypeError)):
        # Не голое «'id'», а понятное сообщение: в данных Warcraft Logs нет ожидаемого поля
        import traceback
        tb = traceback.extract_tb(e.__traceback__)
        where = f"{Path(tb[-1].filename).stem}:{tb[-1].lineno}" if tb else "?"
        return (f"Неожиданный формат данных ({type(e).__name__}: {msg}, место: {where}). "
                "Пришлите это сообщение разработчику — по нему видно, что поправить.")
    return msg or type(e).__name__


# ------------------------------------------------------- сериализация
def _num(x, nd=3):
    if x is None:
        return None
    try:
        v = float(x)
    except (TypeError, ValueError):
        return x
    return None if math.isnan(v) or math.isinf(v) else round(v, nd)


def _finding(f) -> dict:
    return {"section": f.section, "title": f.title, "my": _num(f.my), "ref": _num(f.ref), "unit": f.unit,
            "diff": _num(f.diff), "impact": _num(f.impact, 4), "level": f.impact_level, "note": f.note,
            "focus": f.focus, "key": f.key, "noise": f.noise, "context": f.context}


def to_json(r: CompareResult, job_id: str) -> dict:
    from .guide import guide_sections
    me, ref, mm, agg = r.me, r.ref, r.mm, r.ref.agg
    name = ref.name_of
    k = me.duration / 60
    my_ng = mm["ngrams"][3]
    my_total = sum(my_ng.values()) or 1
    g, gr = mm["gcd"], agg["gcd"]
    return {
        "mode": "player",
        "info": {"name": me.name, "cls": me.cls, "spec": me.spec, "boss": me.encounter_name,
                 "difficulty": me.difficulty_name, "duration": _fmt_t(me.duration), "kill": me.kill,
                 "url": me.url, "ref_label": ref.label, "ref_n": ref.n, "dps": _num(me.dps, 0),
                 "ref_dps": _num(agg["dps"]["median"], 0), "demo": me.report_code.startswith("DEMO"),
                 "main_cd": name(ref.main_cd) if ref.main_cd else None,
                 # проверка сложности: сколько логов топа той же сложности, что и ваш бой
                 "ref_same_diff": sum(1 for x in ref.logs if x.difficulty == me.difficulty)},
        "excel": f"/api/report/{job_id}",
        "brief": {**r.brief, "gap": _num(r.brief.get("gap"), 4), "explained": _num(r.brief.get("explained"), 4)},
        "progress_keys": {k: {"title": v["title"], "impact": _num(v["impact"], 4)} for k, v in r.progress_keys.items()},
        "build_note": ref.build_note,
        "boss_debuffs": [{kk: _num(v, 4) if isinstance(v, float) else v for kk, v in x.items() if kk != "id"}
                         for x in r.tables.get("boss_debuffs", [])],
        "externals": [{kk: _num(v, 2) if isinstance(v, float) else v for kk, v in x.items() if kk != "id"}
                      for x in r.tables.get("externals", [])],
        "priority": [{kk: _num(v, 3) if isinstance(v, float) else v for kk, v in x.items()}
                     for x in r.tables.get("priority", [])],
        "opener": ({**r.tables["opener"], "dist": _num(r.tables["opener"]["dist"], 3),
                    "ref_dist": _num(r.tables["opener"]["ref_dist"], 3)} if r.tables.get("opener") else None),
        "reliability": r.reliability,
        "summary": [{"metric": x["metric"], "my": _num(x["my"], 4), "ref": _num(x["ref"], 4), "fmt": x["fmt"]}
                    for x in r.summary],
        "top5": [_finding(f) for f in r.top5],
        "findings": [_finding(f) for f in sorted(r.findings, key=lambda f: (f.noise or f.context, -(f.impact or 0), -f.severity))],
        "plan": r.plan,
        "gap_total": _num(r.gap[0]["total"], 0) if r.gap else 0,
        "gap": [{"factor": x["factor"], "dps": _num(x["dps"], 0), "status": x["status"]} for x in r.gap],
        "casts": [{"name": x["name"], "cat": x["cat"], "my": x["my"], "ref": _num((x["ref_cpm"] or 0) * k, 1),
                   "p25": _num((agg["abilities"][x["id"]]["cpm"]["p25"] or 0) * k, 1),
                   "p75": _num((agg["abilities"][x["id"]]["cpm"]["p75"] or 0) * k, 1),
                   "share": _num(x["share"])} for x in r.tables["casts"]],
        "ngrams": [{"seq": " → ".join(name(a) for a in n["seq"]), "share": _num(n["share"], 4),
                    "players": _num(n["player_share"]), "my": _num(my_ng.get(n["seq"], 0) / my_total, 4)}
                   for n in agg["ngrams"][3]],
        "cooldowns": [{**{kk: _num(v, 1) if isinstance(v, float) else v for kk, v in x.items()}}
                      for x in r.tables["cooldowns"]],
        "cd_timing": [{"name": x["name"], "k": x["k"], "my": _num(x["my"], 1), "ref": _num(x["ref"], 1),
                       "p25": _num(x["p25"], 1), "p75": _num(x["p75"], 1),
                       "diff": _num(x["my"] - x["ref"], 1) if x["my"] is not None and x["ref"] is not None else None,
                       "confidence": _num(x["confidence"])} for x in r.tables["cd_timing"]],
        "burst": [{"window": x["window"], "name": x["name"], "kind": x.get("kind", ""), "ref_t0": _num(x["ref_t0"], 1),
                   "my_t0": _num(x["my_t0"], 1), "ref_offset": _num(x["ref_offset"], 1),
                   "my_offset": _num(x["my_offset"], 1), "share": _num(x["share"])} for x in r.tables["burst"]],
        "gear": _gear_json(r.tables.get("gear") or {}),
        "defensives": [{kk: _num(v, 1) if isinstance(v, float) else v for kk, v in x.items()}
                       for x in r.tables["defensives"]],
        "deaths": [{"t": _fmt_t(d["t"]), "sources": [[n, _num(v, 0)] for n, v in d["sources"]],
                    "defensives": d["defensives"], "ref_share": _num(d["ref_def_share"])} for d in r.tables["deaths"]],
        "uptime": [{kk: _num(v, 4) if isinstance(v, float) else v for kk, v in x.items()} for x in r.tables["uptime"]],
        "resource_name": mm["resource"]["name"] if mm["resource"] else None,
        "resource": [{kk: _num(v, 4) if isinstance(v, float) else v for kk, v in x.items()} for x in r.tables["resource"]],
        "procs": [{kk: _num(v, 2) if isinstance(v, float) else v for kk, v in x.items()} for x in r.tables["procs"]],
        "gcd": {"my": {"gcd": _num(g["gcd"], 2), "idle": _num(g["idle_per_action"]), "share": _num(g["idle_share"], 4),
                       "total": _num(g["idle_total"], 1)},
                "ref": {"gcd": _num(gr["gcd"]["median"], 2), "idle": _num(gr["idle_per_action"]["median"]),
                        "share": _num(gr["idle_share"]["median"], 4),
                        "total": _num((gr["idle_share"]["median"] or 0) * me.duration, 1)},
                "by30_my": [_num(v, 2) for v in g["idle_by_30s"]],
                "by30_ref": [_num(v, 2) for v in gr["idle_by_30s"]],
                "worst": [{"from": _fmt_t(x["from"]), "to": _fmt_t(x["to"]), "idle": _num(x["idle"], 1),
                           "prev": x["prev_name"], "next": x["next_name"], "ref": _num(x["ref_bucket_idle"], 2)}
                          for x in r.tables["gcd_worst"]]},
        "timeline": [{"lane": e["lane"], "t": _num(e["t"], 1), "cat": e["cat"], "name": e["name"], "k": e["k"],
                      "delta": _num(e.get("delta"), 1)} for e in r.timeline],
        "duration_s": _num(max(me.duration, agg["duration"]["median"] or 0), 1),
        "dps_5s": [_num(v, 0) for v in (mm.get("dps_5s") or [])],
        "mechanics": [{"name": x["name"], "k": x["k"], "ref": _num(x["ref"], 1), "my": _num(x["my"], 1),
                       "key": x["key"]} for x in r.tables["mechanics"]],
        "guide": [{"title": s["title"], "rows": [[a, b, c, _num(d)] for a, b, c, d in s["rows"]]}
                  for s in guide_sections(ref)],
        "categories": CATEGORY_RU,
    }


# ---------------------------------------------------------------- HTTP
WEB = STATIC.parent
STATIC_FILES = {
    "/static/chart.umd.min.js": ("chart.umd.min.js", "application/javascript; charset=utf-8"),
    "/static/qrcode.js": ("qrcode.js", "application/javascript; charset=utf-8"),
    "/static/icon-192.png": ("icon-192.png", "image/png"),
    "/static/icon-512.png": ("icon-512.png", "image/png"),
    "/apple-touch-icon.png": ("icon-180.png", "image/png"),
    "/favicon.ico": ("icon-192.png", "image/png"),
}
MANIFEST = {
    "name": "Разбор лога Warcraft Logs", "short_name": "Разбор WCL", "lang": "ru",
    "start_url": "/", "display": "standalone", "background_color": "#EEF1F4", "theme_color": "#16213A",
    "icons": [{"src": "/static/icon-192.png", "sizes": "192x192", "type": "image/png"},
              {"src": "/static/icon-512.png", "sizes": "512x512", "type": "image/png", "purpose": "any"}],
}
DENIED_PAGE = """<!doctype html><html lang="ru"><meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1"><title>Нет доступа</title>
<body style="font:16px/1.5 system-ui,sans-serif;background:#EEF1F4;color:#16213A;margin:0;padding:32px 20px">
<div style="max-width:460px;margin:0 auto;background:#fff;border-radius:14px;padding:24px;border:1px solid #D5DBE4">
<h1 style="font-size:22px;margin:0 0 10px">Нужен ключ доступа</h1>
<p>Откройте программу на компьютере, нажмите «На другом устройстве» и отсканируйте QR-код камерой телефона.</p>
<p style="color:#5C6781">Если QR-код уже сканировали, а доступ пропал — ключ могли сменить или доступ с телефона выключен.</p>
</div></body></html>"""
MAX_BODY = 32 * 1024 * 1024  # резервная копия с историей разборов может весить несколько МБ


def lan_ip() -> str | None:
    """IP компьютера в локальной сети (пакеты не отправляются)."""
    try:
        with socket.socket(socket.AF_INET, socket.SOCK_DGRAM) as s:
            s.connect(("10.255.255.255", 1))
            ip = s.getsockname()[0]
        return None if ip.startswith("127.") else ip
    except OSError:
        return None


class Handler(BaseHTTPRequestHandler):
    server_version = "WCLAnalyzer/3.0"

    def log_message(self, fmt, *args):  # сервер не ведёт журнал запросов
        pass

    # ------------------------------------------------------ доступ
    def _is_local(self) -> bool:
        ip = self.client_address[0]
        return ip in ("127.0.0.1", "::1") or ip.startswith(("::ffff:127.", "127."))

    def _cookie(self, name: str) -> str | None:
        for part in (self.headers.get("Cookie") or "").split(";"):
            k, _, v = part.strip().partition("=")
            if k == name:
                return v
        return None

    def _creds(self) -> tuple[str, str] | None:
        cid = (self.headers.get("X-WCL-Id") or "").strip()
        sec = (self.headers.get("X-WCL-Secret") or "").strip()
        return (cid, sec) if cid and sec else None

    def _same_site(self) -> bool:
        """Защита локальной программы от чужих сайтов, открытых в том же браузере.
        Host — только адрес-число (127.0.0.1, адрес в Wi-Fi сети) или localhost: так не пройдёт подмена DNS
        (чужой домен, указывающий на 127.0.0.1). POST — только со страницы самой программы (Origin совпадает
        с адресом) и только JSON: простой кросс-сайтовый запрос без предварительной проверки браузера
        с типом application/json отправить нельзя."""
        host = (self.headers.get("Host") or "").strip().lower()
        name = host.rsplit(":", 1)[0].strip("[]") if host else ""
        if name not in ("localhost",) and not re.fullmatch(r"[0-9.]+|[0-9a-f:]+", name or "x"):
            return False
        if self.command == "POST":
            origin = (self.headers.get("Origin") or "").strip().lower()
            if origin and origin not in (f"http://{host}", f"https://{host}"):
                return False
            ctype = (self.headers.get("Content-Type") or "").split(";")[0].strip().lower()
            if ctype != "application/json":
                return False
        return True

    def _allowed(self) -> bool:
        """Публичный сервис открыт всем (данные у каждого в своём браузере).
        Локальная программа: этот компьютер — всегда, телефон — по ссылке-ключу из QR-кода.
        Запросы с чужих сайтов (подмена адреса, POST не со страницы программы) — отказ."""
        if not SERVER["public"] and not self._same_site():
            if urlparse(self.path).path.startswith("/api/"):
                self._json({"error": "Запрос не со страницы программы — отклонён"}, 403)
            else:
                self._file(b"Forbidden", "text/plain; charset=utf-8", status=403)
            return False
        if SERVER["public"] or self._is_local():
            return True
        st = settings.load()
        if not st["lan"]:
            self._deny()
            return False
        q = parse_qs(urlparse(self.path).query)
        if q.get("k", [None])[0] == st["token"]:
            self.send_response(302)
            self.send_header("Set-Cookie", f"wclk={st['token']}; Path=/; Max-Age=31536000; SameSite=Lax; HttpOnly")
            self.send_header("Location", urlparse(self.path).path or "/")
            self.end_headers()
            return False
        if self._cookie("wclk") == st["token"]:
            return True
        self._deny()
        return False

    def _deny(self):
        if urlparse(self.path).path.startswith("/api/"):
            return self._json({"error": "Нет доступа: отсканируйте QR-код в программе на компьютере."}, 403)
        self._file(DENIED_PAGE.encode("utf-8"), "text/html; charset=utf-8", status=403)

    # ------------------------------------------------------ ответы
    def _json(self, data, status=200):
        body = json.dumps(data, ensure_ascii=False).encode("utf-8")
        self._file(body, "application/json; charset=utf-8", status=status)

    def _file(self, data: bytes, ctype: str, cache: str = "no-store", extra: dict | None = None, status=200):
        self.send_response(status)
        self.send_header("Content-Type", ctype)
        self.send_header("Content-Length", str(len(data)))
        self.send_header("Cache-Control", cache)
        self.send_header("X-Content-Type-Options", "nosniff")
        self.send_header("Referrer-Policy", "no-referrer")
        for k, v in (extra or {}).items():
            self.send_header(k, v)
        self.end_headers()
        self.wfile.write(data)

    def _body(self) -> dict:
        n = int(self.headers.get("Content-Length") or 0)
        if n > (1024 * 1024 if SERVER["public"] else MAX_BODY):  # публичному серверу большие запросы не нужны
            raise ValueError("Слишком большой запрос")
        return json.loads(self.rfile.read(n) or b"{}") if n else {}

    # ------------------------------------------------------ GET
    def do_GET(self):  # noqa: N802
        if not self._allowed():
            return
        SERVER["seen"] = True
        path = urlparse(self.path).path
        if path in ("/", "/index.html"):
            self._file(_page(), "text/html; charset=utf-8", extra={"Content-Security-Policy": CSP.format(n=PAGE_NONCE) if not SERVER["public"]
                                         else CSP.format(n=PAGE_NONCE).replace("; frame-ancestors 'none'", "")})
        elif path in STATIC_FILES:
            name, ctype = STATIC_FILES[path]
            self._file((WEB / name).read_bytes(), ctype, "max-age=86400")
        elif path == "/manifest.webmanifest":
            self._file(json.dumps(MANIFEST, ensure_ascii=False).encode("utf-8"), "application/manifest+json")
        elif path == "/api/version":
            from . import update
            self._json({**update.current(), "updates": not SERVER["public"] and self._is_local()})
        elif path == "/api/limit":
            # Часовой лимит Warcraft Logs для ключа из заголовков: сколько очков потрачено и когда сброс.
            # Запрос rateLimitData лёгкий; страница спрашивает раз в минуту и после каждого разбора.
            try:
                client = CLIENT_FACTORY(self._creds())
                if not hasattr(client, "rate_limit"):
                    return self._json({"available": False})
                info = client.rate_limit()
                self._json({"available": True, "limit": info.get("limitPerHour"),
                            "spent": info.get("pointsSpentThisHour"), "reset_in": info.get("pointsResetIn")})
            except SystemExit:
                self._json({"available": False, "need_key": True})
            except Exception as e:  # noqa: BLE001
                self._json({"available": False, "error": _friendly(e)})
        elif path == "/api/status":
            from .platform_support import app_mode, downloads_dir
            desk = app_mode() in ("exe", "python") and not SERVER["public"] and self._is_local()
            self._json({"has_key": bool(self._creds() or env_key()), "server_key": bool(env_key()),
                        "public": SERVER["public"], "local": self._is_local(), "app": app_mode(),
                        # программа на этом компьютере: полный путь «Загрузок» — папка по умолчанию
                        "downloads": str(downloads_dir()) if desk else None,
                        # где лежат скачанные с WCL данные (кэш) — полным путём
                        "cache": str(_cache_file()) if desk else None,
                        "cache_mb": round(_cache_file().stat().st_size / 1048576, 1) if desk and _cache_file().exists() else None})
        elif path == "/api/phone":
            # Только локальная программа и только на самом компьютере: адрес для телефона в Wi-Fi сети
            if SERVER["public"] or not self._is_local():
                return self._json({"error": "Доступно только на компьютере с программой"}, 403)
            st, ip = settings.load(), lan_ip()
            base = f"http://{ip}:{SERVER['port']}/?k={st['token']}" if ip and st["lan"] and SERVER["lan"] else None
            self._json({"setting": st["lan"], "bound_lan": SERVER["lan"], "ip": ip, "base": base})
        elif path.startswith("/api/job/"):
            job = JOBS.get(path.rsplit("/", 1)[-1])
            if not job:
                return self._json({"error": "Разбор не найден: результаты хранятся на сервере не дольше часа"}, 404)
            since = int(self.headers.get("X-Log-Since") or 0)
            pv = int(self.headers.get("X-Partial-V") or 0)  # готовая часть «Разобрать бой» — один раз на версию
            v, part = job.get("partial_v", 0), job.get("partial")  # читаем один раз: версия и часть — пара
            self._json({"state": job["state"], "progress": job["progress"], "log": job["log"][since:],
                        "partial": part if v > pv else None,
                        "partial_v": v,
                        "wait_left": max(0.0, job.get("wait_until", 0) - time.time()),
                        "log_total": len(job["log"]), "error": job.get("error"), "need_key": job.get("need_key", False),
                        "result": job.get("result") if job["state"] == "done" else None})
        elif path.startswith("/api/report/"):
            parts = path[len("/api/report/"):].split("/")
            job, key = JOBS.get(parts[0]), "xlsx_alt" if parts[1:] == ["alt"] else "xlsx"
            if not job or not job.get(key):
                return self._json({"error": "Отчёт уже удалён с сервера — откройте его из истории разборов"}, 404)
            self._file(job[key], "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                       extra={"Content-Disposition": f"attachment; filename*=UTF-8''{_quote(job[key + '_name'])}"})
        else:
            self._json({"error": "Не найдено"}, 404)

    # ------------------------------------------------------ POST
    def do_POST(self):  # noqa: N802
        if not self._allowed():
            return
        path = urlparse(self.path).path
        try:
            body = self._body()
            creds = self._creds()
            if path == "/api/key":
                CHECK_KEY(str(body.get("client_id", "")).strip(), str(body.get("client_secret", "")).strip())
                return self._json({"ok": True})
            if path == "/api/zones":  # рейды текущего дополнения — для «Топ прогресса» без лога
                cl = _limited(CLIENT_FACTORY(creds), 0)
                return self._json({"zones": cl.raid_zones()})
            if path == "/api/spell":  # краткое описание способности с Wowhead (по нажатию «i», если в разборе его не было)
                from . import wowhead
                try:
                    sid = int(body.get("id"))
                except (TypeError, ValueError):
                    raise ValueError("Нет номера способности") from None
                w = wowhead.lookup([sid], limit=1).get(sid) or {}
                return self._json({"name": w.get("name"), "desc": w.get("desc"), "url": wowhead.page_url(sid)})
            if path == "/api/inspect":
                from .collect import inspect_report
                cl = _limited(CLIENT_FACTORY(creds), 0)  # поиск боя не ждёт сброса лимита: сразу объясняем, что случилось
                return self._json(inspect_report(cl, body["url"], body.get("fight")))
            if path == "/api/analyze":
                params = {k: body.get(k) for k in ("mode", "demo", "url", "fight", "actor", "ref", "against", "units", "prev", "wait",
                                                   "zone", "encounter", "difficulty", "roster", "minutes",
                                                   "refresh", "max_age_days", "pick", "mythic")}
                if params["mode"] not in (None, "fight", "progress", "raid", "raidrot", "saves", "allbosses", "rosterplan"):
                    params["mode"] = None
                return self._json({"job": start_job(creds, params)})
            if path == "/api/refs/refresh":
                params = {"mode": "refresh", "refs": body.get("refs") or [], "max_age_days": body.get("max_age_days")}
                return self._json({"job": start_job(creds, params)})
            if path == "/api/job/done":
                # Браузер забрал результат и Excel — сервер удаляет их сразу, не дожидаясь часа
                with JOBS_LOCK:
                    job = JOBS.get(str(body.get("job", "")))
                    if job and job["state"] != "running":
                        JOBS.pop(job["id"], None)
                return self._json({"ok": True})
            if path == "/api/cache/clear" and not SERVER["public"] and self._is_local():
                if any(j["state"] == "running" for j in list(JOBS.values())):
                    raise ValueError("Сейчас идёт разбор — очистите кэш, когда он закончится")
                return self._json(_shared_cache().clear())
            if path in ("/api/pickdir", "/api/opendir") and app_mode_is_desktop() and not SERVER["public"] and self._is_local():
                # Окно выбора папки и «Открыть папку» — на этом же компьютере; путь — полностью
                from .platform_support import downloads_dir, open_folder, pick_folder
                if path == "/api/opendir":
                    open_folder(str(_cache_file().parent) if body.get("cache") else (str(body.get("dir") or "") or str(downloads_dir())))
                    return self._json({"ok": True})
                try:
                    return self._json({"dir": pick_folder(str(body.get("dir") or ""))})
                except Exception as e:  # noqa: BLE001 — нет окна выбора: путь впишут вручную
                    return self._json({"dir": None, "error": str(e)})
            if path == "/api/save" and app_mode_is_desktop() and not SERVER["public"] and self._is_local():
                # Программа на этом же компьютере: файл пишется в папку, выбранную в «Куда сохранять файлы»
                import base64
                name = re.sub(r'[\\/:*?"<>|]+', "_", str(body.get("name") or "file"))[:120]
                folder = Path(str(body.get("dir") or "")).expanduser()
                if not str(body.get("dir") or "").strip() or not folder.is_absolute():
                    raise ValueError("Укажите полный путь к папке, например C:\\Users\\Имя\\Documents\\WCL")
                folder.mkdir(parents=True, exist_ok=True)
                (folder / name).write_bytes(base64.b64decode(body.get("data") or ""))
                return self._json({"saved": str(folder / name)})
            if path in ("/api/android/backups", "/api/android/backup"):
                # Android-приложение: окно выбора файла в нём не открывается — список резервных копий
                # из «Загрузок» даёт сама программа, и она же читает выбранную
                from .platform_support import android_list_backups, android_read_backup, app_mode
                if SERVER["public"] or not self._is_local() or app_mode() != "android":
                    return self._json({"error": "Доступно только в Android-приложении"}, 403)
                if path == "/api/android/backups":
                    return self._json({"files": android_list_backups()})
                text = android_read_backup(body.get("id"))
                if len(text) > MAX_BODY:
                    raise ValueError("Файл слишком большой")
                return self._json({"text": text})
            if path in ("/api/save", "/api/open"):
                # Только Android-приложение: встроенное окно не умеет скачивать файлы и открывать ссылки
                from .platform_support import android_open_url, android_save_download, app_mode
                if SERVER["public"] or not self._is_local() or app_mode() != "android":
                    return self._json({"error": "Доступно только в Android-приложении"}, 403)
                if path == "/api/open":
                    url = str(body.get("url", ""))
                    if not url.startswith(("https://", "http://")):
                        raise ValueError("Неверная ссылка")
                    android_open_url(url)
                    return self._json({"ok": True})
                import base64
                name = re.sub(r'[\\/:*?"<>|]+', "_", str(body.get("name") or "file"))[:120]
                where = android_save_download(name, base64.b64decode(body.get("data") or ""),
                                              str(body.get("mime") or "application/octet-stream"),
                                              sub=str(body.get("dir") or ""))
                return self._json({"saved": where})
            if path in ("/api/update/check", "/api/update/apply"):
                # Обновление программы — только на самом устройстве с программой, не на публичном сервере
                if SERVER["public"] or not self._is_local():
                    return self._json({"error": "Обновление доступно только в самой программе"}, 403)
                from . import update
                try:
                    if path == "/api/update/check":
                        return self._json(update.check())
                except update.UpdateError as e:
                    return self._json({"error": str(e)}, 502)
                if any(j.get("state") == "running" for j in list(JOBS.values())):  # перезапуск оборвал бы разбор
                    return self._json({"error": "Сейчас идёт разбор — обновитесь, когда он закончится "
                                                "(иначе разбор прервётся и его придётся запускать заново)."}, 409)
                if body.get("kind") == "full":
                    return self._json(update.apply_full())
                res = update.apply_code()
                if res.get("ok"):
                    request_restart()  # новый код подключается сразу, страница перезагрузится сама
                    res = {**res, "restart": "Перезапускаю программу…", "soft_restart": True}
                return self._json(res)
            if path == "/api/phone":
                if SERVER["public"] or not self._is_local():
                    return self._json({"error": "Доступно только на компьютере с программой"}, 403)
                ch = {}
                if "enabled" in body:
                    ch["lan"] = bool(body["enabled"])
                if body.get("new_token"):
                    ch["token"] = secrets.token_urlsafe(12)
                settings.save(**ch)
                return self._json({"ok": True})
            return self._json({"error": "Не найдено"}, 404)
        except SystemExit:
            self._json({"error": NEED_KEY, "need_key": True}, 400)
        except (ValueError, LookupError, KeyError, TypeError) as e:
            self._json({"error": str(e)}, 400)
        except Exception as e:  # noqa: BLE001
            self._json({"error": _friendly(e)}, 500)


def _quote(s: str) -> str:
    from urllib.parse import quote
    return quote(s)


def _free_port(preferred: int, host: str) -> int:
    for port in [preferred] + list(range(preferred + 1, preferred + 50)):
        with socket.socket() as s:
            if os.name != "nt":  # как у самого сервера: порт после недавних соединений (TIME_WAIT) свободен
                s.setsockopt(socket.SOL_SOCKET, socket.SO_REUSEADDR, 1)
            try:
                s.bind((host, port))
                return port
            except OSError:
                continue
    return 0


def serve(port: int = 8765, open_browser: bool = True, local_only: bool = False, public: bool = False,
          same_port: bool = False) -> None:
    """same_port — перезапуск после обновления: тот же порт, что был, иначе открытая страница его не найдёт."""
    public = public or os.environ.get("WCL_PUBLIC") == "1"
    from .platform_support import android_preload, app_mode
    if app_mode() == "android":
        try:  # классы Java — в главном потоке, иначе сохранение файлов и ссылки падают в потоках сервера
            android_preload()
        except Exception as e:  # noqa: BLE001
            print(f"Android: классы Java не загружены заранее: {e}")
    if public:  # публичный сервер ничего не пишет на диск — и игровые данные держит только в памяти
        from . import game_data
        game_data.SAVE["enabled"] = False
        from . import wowhead
        wowhead.SAVE["enabled"] = False
    if public:
        host, port = "0.0.0.0", int(os.environ.get("PORT", port))
        open_browser = False
    else:
        host = "127.0.0.1" if local_only else "0.0.0.0"
        if not same_port:
            port = _free_port(port, host)
    for attempt in range(30):  # после перезапуска порт может освободиться не сразу
        try:
            httpd = ThreadingHTTPServer((host, port), Handler)
            break
        except OSError:
            if not same_port or attempt == 29:
                raise
            time.sleep(0.3)
    httpd.daemon_threads = True
    if not public and not same_port:  # кэш скачанного с WCL: старое и лишнее — удалить, старые записи — сжать
        threading.Thread(target=_prune_cache, daemon=True).start()
    SERVER.update(port=httpd.server_address[1], lan=not local_only and not public, public=public,
                  httpd=httpd, restart=False, seen=False)
    if not public:  # сервер поднялся — загруженная версия кода рабочая: отката к встроенной не будет
        try:
            from . import update
            update.confirm_running_build()
        except Exception:  # noqa: BLE001
            pass
    # После перезапуска (обновление .exe целиком) браузер уже открыт — страница сама переподключится.
    # Если за 20 с никто не зашёл (вкладку закрыли), открываем браузер как обычно.
    reopen = bool(os.environ.pop("WCL_NO_BROWSER", None))
    if reopen:
        open_browser = False
    _CACHE["obj"] = None
    if public:
        print(f"WCL Analyzer работает как сервис на порту {SERVER['port']}. "
              "Данные пользователей хранятся только в их браузерах.")
    else:
        url = f"http://127.0.0.1:{SERVER['port']}/"
        print(f"WCL Analyzer открыт в браузере: {url}")
        if SERVER["lan"] and settings.load()["lan"] and lan_ip():
            print("С телефона в той же Wi-Fi сети: нажмите «На другом устройстве» и отсканируйте QR-код.")
        print("Не закрывайте это окно, пока пользуетесь программой. Остановить — Ctrl+C.")
        if open_browser:
            threading.Timer(0.8, lambda: webbrowser.open(url)).start()
        elif reopen:
            threading.Timer(20, lambda: None if SERVER.get("seen") else webbrowser.open(url)).start()
    try:
        httpd.serve_forever()
    except KeyboardInterrupt:
        print("\nОстановлено.")
    httpd.server_close()
    if SERVER.get("restart"):
        _reload_and_serve(port=SERVER["port"], local_only=local_only, public=public)


def request_restart(delay: float = 0.7) -> None:
    """Мягкий перезапуск после обновления кода: сервер останавливается, программа подключает новый код
    и снова запускает сервер на том же адресе — уже открытая страница сама перезагрузится.
    Процесс не завершается: так работает и на Windows, и в приложении Android."""
    SERVER["restart"] = True
    httpd = SERVER.get("httpd")
    if httpd is not None:
        threading.Timer(delay, httpd.shutdown).start()  # после того, как ответ на запрос ушёл


def _reload_and_serve(port: int, local_only: bool, public: bool) -> None:
    import sys
    print("Перезапуск: подключаю обновлённый код…")
    sys.meta_path[:] = [f for f in sys.meta_path if type(f).__name__ != "_Finder"]  # прежний скачанный код
    for m in [m for m in list(sys.modules) if m == "wcl_analyzer" or m.startswith("wcl_analyzer.")]:
        del sys.modules[m]
    try:
        import wcl_boot  # загрузчик оболочки (.exe / APK): подключает новый скачанный код
        build = wcl_boot.activate()
        if build:
            print(f"Обновлённая версия кода: сборка {build}")
    except ImportError:
        pass
    from wcl_analyzer.web import serve as new_serve  # noqa: PLC0415 — уже новый модуль
    new_serve(port=port, open_browser=False, local_only=local_only, public=public, same_port=True)
