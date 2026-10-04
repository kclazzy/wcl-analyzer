"""Excel-отчёт по всему рейду: обзор, игроки, механики, смерти, пуллы, таймлайн."""
from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference, ScatterChart, Series
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.formatting.rule import CellIsRule, ColorScaleRule, FormulaRule
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from .names_ru import cls_ru, spec_ru
from .excel_report import BORDER, F_1, F_INT, F_PCT, FONT, Sheet, _mmss, _v

RED = PatternFill("solid", fgColor="FFC7CE")


def _mmss_val(t: float) -> str:
    t = max(0.0, float(t))
    return f"{int(t // 60)}:{int(t % 60):02d}"
F_RATIO = '0.0"×"'
ROLE_COLOR = {"Танк": "4F7CC0", "Лекарь": "2F8F4E", "DPS": "E39B2F"}


def write_raid_workbook(R: dict, path: str | Path, wb=None) -> Path:
    own = wb is None   # wb передан — листы дописываются в общую книгу (разбор боя целиком)
    if own:
        wb = Workbook()
        wb.remove(wb.active)
    demo = R["info"]["demo"]
    _brief(wb, R, demo)
    _damage(wb, R, demo)
    _vs_top(wb, R, demo)
    _overview(wb, R, demo)
    _players(wb, R, demo)
    _mechanics(wb, R, demo)
    _deaths(wb, R, demo)
    _pulls(wb, R, demo)
    _timeline(wb, R, demo)
    if not own:
        return wb
    path = Path(path)
    wb.save(path)
    return path


# --------------------------------------------------------------- Выжимка
def _brief(wb, R, demo):
    s = Sheet(wb, "Выжимка", demo, {"A": 4, "B": 24, "C": 90})
    I, S = R["info"], R["summary"]
    outcome = "килл" if I["kill"] else f"вайп на {I['boss_pct']:.1f}% здоровья босса" if I["boss_pct"] is not None else "вайп"
    s.title(f"{I['boss']}, {I['difficulty']}: {outcome} за {I['duration']}", f"Пулл {S['pull_n']} из {S['pulls']}.")
    s.section("Главное по бою")
    for line in R.get("brief") or ["Явных проблем не найдено."]:
        s.cell(s.row, 2, "• " + line)
        s.ws.merge_cells(start_row=s.row, start_column=2, end_row=s.row, end_column=3)
        s.row += 1
    s.row += 1
    seen: set = set()
    top = [i for i in R["issues"] if i["severity"] >= 4 and not (i["player"] in seen or seen.add(i["player"]))][:5]
    if top:
        s.section("Кому что поправить")
        import re as _re
        gi = R.get("guide_info") or {n: {"url": u, "src": "mt", "video": (R.get("guide_videos") or {}).get(n)}
                                     for n, u in (R.get("guide_links") or {}).items()}

        def g_of(t):  # первая способность из текста, о которой есть гайд или страница Wowhead
            return next((gi[n] for n in _re.findall(r"«([^»]+)»", t) if n in gi), {})

        def brief(g):
            parts = [g.get("name_ru") or "", f"[{g['type']}]" if g.get("type") else "", g.get("desc") or "",
                     f"Что делать: {g['todo']}" if g.get("todo") else ""]
            return " ".join(p for p in parts if p)
        f, _l = s.table(["", "Игрок", "Что", "Гайд по способности", "Ролик", "Кратко о способности"],
                        [[k + 1, f"{i['player']} ({spec_ru(i.get('cls', ''), i['spec'])})", i["text"],
                          g_of(i["text"]).get("url", ""), g_of(i["text"]).get("video", ""), brief(g_of(i["text"]))]
                         for k, i in enumerate(top)])
        for k, i in enumerate(top):  # кликабельные ссылки в ячейках
            g = g_of(i["text"])
            label = "Wowhead" if g.get("src") == "wowhead" else "Mythic Trap"
            for col, url, text in ((4, g.get("url"), label), (5, g.get("video"), "▶ Смотреть")):
                if url:
                    c = s.ws.cell(f + k, col)
                    c.hyperlink, c.value, c.style = url, text, "Hyperlink"
            from openpyxl.styles import Alignment
            s.ws.cell(f + k, 6).alignment = Alignment(wrap_text=True, vertical="top")
        s.ws.column_dimensions["F"].width = max(s.ws.column_dimensions["F"].width or 0, 60)
    X = R.get("extras") or {}
    if X.get("heaviest"):
        s.section("Самые тяжёлые моменты", "Подробно и с графиком — на листе «Урон по рейду».")
        s.table(["", "Когда", "Механики и рейдовые кулдауны"],
                [[k + 1, x["time"], "«" + "», «".join(x["abilities"]) + "» — "
                  + (", ".join(x["cds"]) if x["cds"] else "без рейдовых кулдаунов")]
                 for k, x in enumerate(X["heaviest"][:3])])
    s.note("Подробности — на следующих листах.")


# ---------------------------------------------------------- Урон по рейду
def _damage(wb, R, demo):
    X = R.get("extras") or {}
    T = X.get("damage_timeline") or []
    if not T:
        return
    s = Sheet(wb, "Урон по рейду", demo, {"A": 10, "B": 18, "C": 30, "D": 4, "E": 10, "F": 16, "G": 12,
                                          "H": 30, "I": 40})
    s.title("Когда рейд получает больше всего урона", "Урон по рейду без танков за каждые 5 секунд; "
            "рейдовые кулдауны и самые тяжёлые моменты — справа.")
    top = s.row
    first, last = s.table(["Время", "Урон за 5 с", "Больше всего от"],
                          [[_mmss_val(x["t"]), x["damage"], x["ability"] or "—"] for x in T], [None, F_INT, None])
    ch = LineChart()
    ch.title = "Урон по рейду за 5 секунд"
    ch.add_data(Reference(s.ws, min_col=2, min_row=first - 1, max_row=last), titles_from_data=True)
    ch.set_categories(Reference(s.ws, min_col=1, min_row=first, max_row=last))
    ch.y_axis.title = "урон"
    ch.y_axis.numFmt = "#,##0"
    ch.x_axis.tickLblSkip = 6
    ch.legend = None
    ch.series[0].graphicalProperties.line.solidFill = "1F8A7E"
    ch.series[0].graphicalProperties.line.width = 22000
    s.chart(ch, f"E{top}", w=24, h=9)
    s.row = top + 20
    end_row = s.row
    s.row = top + 20
    s.ws.cell(row=s.row, column=5, value="Самые тяжёлые моменты").font = Font(name=FONT, bold=True, size=12, color="2E75B6")
    s.row += 1
    s.table(["Когда", "Урон за 5 с", "Доля", "Механики", "Рейдовые кулдауны"],
            [[h["time"], h["damage"], h["share"], ", ".join(h["abilities"]),
              ", ".join(h["cds"]) or "нет"] for h in sorted(X.get("heaviest", []), key=lambda h: h["t"])],
            [None, F_INT, F_PCT, None, None], col=5)
    s.ws.cell(row=s.row, column=5, value="Рейдовые кулдауны").font = Font(name=FONT, bold=True, size=12, color="2E75B6")
    s.row += 1
    cds = X.get("raid_cds") or []
    f2, l2 = s.table(["Когда", "Кулдаун", "Игрок", "На пик урона", "Урон за 10 с после"],
                     [[c["time"], c["name"], f"{c['player']} ({c['role']})", c["on_peak"] or "вне пиков", c["taken_10s"]]
                      for c in cds] or [["—", "не найдено", "", "", None]], [None, None, None, None, F_INT], col=5)
    if cds:
        s.ws.conditional_formatting.add(f"H{f2}:H{l2}", CellIsRule(operator="equal", formula=['"вне пиков"'], fill=RED))
    s.row = max(s.row, end_row)


# ------------------------------------------------------------------ Рейд
def _overview(wb, R, demo):
    s = Sheet(wb, "Рейд", demo, {"A": 30, "B": 22, "C": 16, "D": 70})
    I, S = R["info"], R["summary"]
    outcome = "килл" if I["kill"] else f"вайп на {I['boss_pct']:.1f}% здоровья босса" if I["boss_pct"] is not None else "вайп"
    s.title(f"Разбор рейда: {I['boss']}, {I['difficulty']}",
            f"Пулл {S['pull_n']} из {S['pulls']}: {outcome}, {I['duration']}. "
            f"Состав: {I['tanks']} танка, {I['healers']} лекаря, {I['dps']} DPS.")
    fd = S["first_death"]
    rows = [
        ("Отчёт WCL", I["url"]),
        ("Урон рейда, DPS", S["raid_dps"]),
        ("Лечение рейда, HPS", S["raid_hps"]),
        ("Медиана активного времени", S["median_active"]),
        ("Смертей (без конца вайпа)", f"{S['deaths_before_tail']} из {S['deaths']}"),
        ("Первая смерть", f"{fd['time']} — {fd['player']}, «{fd['ability']}»" if fd else "нет"),
        ("Жажда крови", f"{S['lust']['caster']}, {_fmt(S['lust']['t'])}" if S["lust"] else "не найден"),
        ("Урон от выборочных механик", S["selective_share"]),
        ("Зелий за бой", f"{S['potions']} ({'по данным WCL, с пре-потом' if S['potion_source'] == 'playerDetails' else 'по кастам, без пре-пота'})"),
        ("Камней здоровья", S["healthstones"]),
        ("Прерываний / диспелов", f"{S['interrupts']} / {S['dispels']}"),
        ("Пуллов на боссе / киллов", f"{S['pulls']} / {S['kills']}"),
    ]
    for k, v in rows:
        s.cell(s.row, 1, k, bold=True)
        c = s.cell(s.row, 2, v)
        if k == "Отчёт WCL":
            c.hyperlink = v
            c.font = Font(name=FONT, color="0563C1", underline="single")
        elif k.startswith(("Урон рейда", "Лечение рейда")):
            c.number_format = F_INT
        elif k.startswith(("Медиана активного", "Урон от выборочных")):
            c.number_format = F_PCT
        s.row += 1
    s.row += 1
    s.section("Что проверить", "Факты из лога, отсортированы по важности. Пороги — в конце листа.")
    sev = {5: "Высокая", 4: "Высокая", 3: "Средняя", 2: "Низкая", 1: "Низкая"}
    f, l = s.table(["Игрок", "Роль", "Важность", "Что произошло"],
                   [[i["player"], i["role"], sev[i["severity"]], i["text"]] for i in R["issues"]])
    if R["issues"]:
        for lvl, color in (("Высокая", "FFC7CE"), ("Средняя", "FFEB9C"), ("Низкая", "E7E6E6")):
            s.ws.conditional_formatting.add(f"C{f}:C{l}", CellIsRule(operator="equal", formula=[f'"{lvl}"'],
                                            fill=PatternFill("solid", fgColor=color)))
    th = R["thresholds"]
    s.note(f"Пороги: активное время < {th['active']:.0%} и ниже медианы рейда; выборочные механики ≥ {th['select']:.1f}× "
           f"среднего; общая механика ≥ {th['raidwide']:.2f}× медианы; парс < {th['parse']}%; "
           f"смерти в последние {th['wipe_tail']} с вайпа считаются концом вайпа.")


def _fmt(t):
    if t is None:
        return "—"
    return f"{int(t // 60)}:{t % 60:04.1f}"


# ---------------------------------------------------------------- Игроки
def _players(wb, R, demo):
    s = Sheet(wb, "Игроки", demo, {"A": 16, "B": 13, "C": 15, "D": 9, "E": 8, "F": 14, "G": 11, "H": 14, "I": 11,
                                   "J": 11, "K": 9, "L": 9, "M": 11, "N": 10, "O": 8, "P": 8, "Q": 11, "R": 9,
                                   "S": 14, "T": 14, "U": 11})
    I = R["info"]
    s.title("Игроки рейда", "DPS и HPS считаются формулами от урона, лечения и длительности боя.")
    dur_row = s.row
    s.cell(dur_row, 1, "Длительность боя, с", bold=True)
    s.cell(dur_row, 2, I["duration_s"], fmt=F_1)
    s.row += 2
    players = R["players"]
    first = s.row + 1
    last = first + len(players) - 1
    D = f"$B${dur_row}"
    rows = []
    for p in players:
        rows.append([p["name"], cls_ru(p["cls"]), spec_ru(p["cls"], p["spec"]).split(" (")[0], p["role_ru"], _v(p["ilvl"], 1), p["damage"],
                     f"=F{{r}}/{D}", p["healing"], f"=H{{r}}/{D}", p["active"],
                     None if p["parse"] is None else p["parse"] / 100, p["deaths"], p["first_death"],
                     _mmss("M").replace("—", ""), p["potions"], p["healthstones"], p["interrupts"],
                     p["dispels"], p["taken"], p["selective"],
                     f'=IF(D{{r}}="Танк","",IFERROR(T{{r}}/AVERAGEIFS($T${first}:$T${last},$D${first}:$D${last},"<>Танк"),""))'])
    f, l = s.table(["Игрок", "Класс", "Спек", "Роль", "ilvl", "Урон", "DPS", "Лечение", "HPS", "Активное время",
                    "Парс", "Смертей", "Первая смерть, с", "Первая смерть", "Зелья", "Камни", "Прерывания",
                    "Диспелы", "Получено урона", "От выборочных механик", "× среднего"], rows,
                   [None, None, None, None, F_1, F_INT, F_INT, F_INT, F_INT, F_PCT, F_PCT, F_INT, F_1, None,
                    F_INT, F_INT, F_INT, F_INT, F_INT, F_INT, F_RATIO])
    s.ws.conditional_formatting.add(f"J{f}:J{l}", FormulaRule(
        formula=[f'AND(ISNUMBER(J{f}),D{f}<>"Танк",J{f}<{R["thresholds"]["active"]})'], fill=RED))
    s.ws.conditional_formatting.add(f"U{f}:U{l}", FormulaRule(
        formula=[f'AND(ISNUMBER(U{f}),U{f}>={R["thresholds"]["select"]})'], fill=RED))
    s.ws.conditional_formatting.add(f"L{f}:L{l}", CellIsRule(operator="greaterThan", formula=["0"], fill=RED))
    s.ws.conditional_formatting.add(f"K{f}:K{l}", ColorScaleRule(
        start_type="num", start_value=0, start_color="F8696B", mid_type="num", mid_value=0.5,
        mid_color="FFEB84", end_type="num", end_value=1, end_color="63BE7B"))
    s.note("Активное время — доля времени, пока игрок жив, в которую он действовал (по данным WCL). "
           "× среднего — урон от выборочных механик относительно среднего по нетанкам.")
    if R["summary"]["potion_source"] != "playerDetails":
        s.note("Зелья посчитаны по кастам: пре-пот в них не виден.")
    s.ws.freeze_panes = f"B{f}"

    # Графики: DPS (танки и DPS) и HPS (лекари) — строки уже отсортированы так
    n_dmg = sum(1 for p in players if p["role_ru"] != "Лекарь")
    ch = BarChart()
    ch.type = "bar"
    ch.title = "DPS игроков"
    ch.add_data(Reference(s.ws, min_col=7, min_row=f, max_row=f + n_dmg - 1), titles_from_data=False)
    ch.set_categories(Reference(s.ws, min_col=1, min_row=f, max_row=f + n_dmg - 1))
    ch.legend = None
    ch.y_axis.numFmt = "#,##0"
    ch.x_axis.scaling.orientation = "maxMin"
    s.chart(ch, f"A{s.row + 1}", w=16, h=max(7, 0.55 * n_dmg + 2))
    ch2 = BarChart()
    ch2.type = "bar"
    ch2.title = "Активное время"
    ch2.add_data(Reference(s.ws, min_col=10, min_row=f, max_row=l), titles_from_data=False)
    ch2.set_categories(Reference(s.ws, min_col=1, min_row=f, max_row=l))
    ch2.legend = None
    ch2.y_axis.numFmt = "0%"
    ch2.y_axis.scaling.min = 0
    ch2.y_axis.scaling.max = 1
    ch2.x_axis.scaling.orientation = "maxMin"
    s.chart(ch2, f"H{s.row + 1}", w=14, h=max(7, 0.5 * len(players) + 2))
    if f + n_dmg <= l:
        ch3 = BarChart()
        ch3.type = "bar"
        ch3.title = "HPS лекарей"
        ch3.add_data(Reference(s.ws, min_col=9, min_row=f + n_dmg, max_row=l), titles_from_data=False)
        ch3.set_categories(Reference(s.ws, min_col=1, min_row=f + n_dmg, max_row=l))
        ch3.legend = None
        ch3.y_axis.numFmt = "#,##0"
        ch3.x_axis.scaling.orientation = "maxMin"
        s.chart(ch3, f"O{s.row + 1}", w=12, h=6)


# -------------------------------------------------------------- Механики
def _mechanics(wb, R, demo):
    s = Sheet(wb, "Механики", demo, {"A": 22, "B": 15, "C": 14, "D": 11, "E": 11, "F": 12, "G": 12, "H": 12,
                                     "I": 46})
    s.title("Урон от механик босса",
            "По всему рейду — за одно применение задевает больше 60% рейда; По танкам — 70%+ урона по танкам; "
            "Выборочно — остальное. Выборочный урон бывает и избегаемым, и назначенным: проверьте механику.")
    rows = []
    for a in R["abilities"]:
        top = ", ".join(f"{t['name']} ({t['ratio']:.1f}×)" if t["ratio"] else t["name"] for t in a["top"])
        rows.append([a["name"], a["category"], a["total"], a["share_total"], a["hits"], a["players_hit"],
                     a["coverage"], a["tank_share"], top if a["category"] != "По танкам" else "—"])
    s.table(["Способность", "Тип", "Урон всего", "Доля урона", "Попаданий", "Игроков задето",
             "Охват за применение", "Доля по танкам", "Больше всех получили (× среднего)"], rows,
            [None, None, F_INT, F_PCT, F_INT, F_INT, F_PCT, F_PCT, None])

    H = R["heatmap"]
    if H["abilities"]:
        s.section("Кто сколько получил от выборочных механик",
                  "Значение — во сколько раз больше среднего по нетанкам. 1× — как все; красный — от 3×.")
        hr = s.row
        s.cell(hr, 1, "Игрок", bold=True, color="FFFFFF", fill="1F3864").border = BORDER
        for j, name in enumerate(H["abilities"]):
            c = s.cell(hr, 2 + j, name, bold=True, color="FFFFFF", fill="1F3864", wrap=True)
            c.border = BORDER
        for i, player in enumerate(H["players"]):
            s.cell(hr + 1 + i, 1, player).border = BORDER
            for j in range(len(H["abilities"])):
                c = s.cell(hr + 1 + i, 2 + j, H["ratio"][i][j], fmt=F_RATIO)
                c.border = BORDER
        last_col = get_column_letter(1 + len(H["abilities"]))
        rng = f"B{hr + 1}:{last_col}{hr + len(H['players'])}"
        s.ws.conditional_formatting.add(rng, ColorScaleRule(
            start_type="num", start_value=0, start_color="FFFFFF", mid_type="num", mid_value=1,
            mid_color="FFFFFF", end_type="num", end_value=3, end_color="F8696B"))
        s.row = hr + len(H["players"]) + 2
        s.section("Попадания (сколько раз)")
        s.table(["Игрок"] + H["abilities"], [[p] + H["hits"][i] for i, p in enumerate(H["players"])],
                [None] + [F_INT] * len(H["abilities"]))

    s.section("Общие механики: кто получил заметно больше остальных",
              f"Урон ≥ {R['thresholds']['raidwide']:.2f}× медианы нетанков. Частая причина — нет защитной способности "
              "в этот момент, но это стоит проверить в логе.")
    s.table(["Механика", "Игрок", "Урон", "Медиана рейда", "× медианы"],
            [[o["ability"], o["player"], o["damage"], o["median"], o["ratio"]] for o in R["raidwide"]],
            [None, None, F_INT, F_INT, F_RATIO])
    if not R["raidwide"]:
        s.note("Таких игроков нет.")


# ---------------------------------------------------------------- Смерти
def _deaths(wb, R, demo):
    s = Sheet(wb, "Смерти", demo, {"A": 5, "B": 9, "C": 14, "D": 9, "E": 20, "F": 44, "G": 56, "H": 18})
    s.title("Смерти в бою", "Урон за 5 с до смерти и последние действия игрока за 10 с.")
    rows = [[d["n"], d["time"], d["player"], d["role"], d["ability"],
             ", ".join(f"{x['name']}: {x['damage']:,}".replace(",", " ") for x in d["sources"]),
             " → ".join(d["last_casts"]) or "—",
             "первая смерть" if d["first"] else "конец вайпа" if d["wipe_tail"] else ""] for d in R["deaths"]]
    f, l = s.table(["№", "Время", "Игрок", "Роль", "Причина", "Урон за 5 с", "Последние действия", "Пометка"], rows)
    if rows:
        s.ws.conditional_formatting.add(f"H{f}:H{l}", CellIsRule(operator="equal", formula=['"первая смерть"'],
                                                                fill=RED))
    else:
        s.note("В этом бою никто не умер.")


# ----------------------------------------------------------------- Пуллы
def _pulls(wb, R, demo):
    s = Sheet(wb, "Пуллы", demo, {"A": 7, "B": 7, "C": 9, "D": 12, "E": 14, "F": 10, "G": 16, "H": 12, "I": 18,
                                  "J": 20})
    s.title("Все пуллы на этом боссе", "Здоровье босса в конце: чем ниже, тем ближе к киллу.")
    rows = [[p["n"], p["fight_id"], "килл" if p["kill"] else "вайп", p["duration"], _mmss("D"),
             None if p["boss_pct"] is None else p["boss_pct"] / 100, p["deaths_before_tail"], p["deaths"],
             (p["first_death"] or {}).get("player", "—"),
             (f"{p['first_death']['time']}, {p['first_death']['ability']}" if p["first_death"] else "—")]
            for p in R["pulls"]]
    f, l = s.table(["Пулл", "Бой", "Итог", "Длительность, с", "Длительность", "Здоровье босса", "Смертей до конца вайпа",
                    "Смертей всего", "Первая смерть", "Когда и от чего"], rows,
                   [F_INT, F_INT, None, F_1, None, F_PCT, F_INT, F_INT, None, None])
    sel = [i for i, p in enumerate(R["pulls"]) if p["selected"]]
    for i in sel:
        for c in range(1, 11):
            s.ws.cell(f + i, c).font = Font(name=FONT, bold=True)
    s.note("Жирным — пулл, разобранный в этом отчёте.")
    if len(rows) >= 2:
        ch = LineChart()
        ch.title = "Здоровье босса в конце пулла"
        ch.add_data(Reference(s.ws, min_col=6, min_row=f, max_row=l), titles_from_data=False)
        ch.set_categories(Reference(s.ws, min_col=1, min_row=f, max_row=l))
        ch.legend = None
        ch.y_axis.numFmt = "0%"
        ch.y_axis.scaling.min = 0
        ch.y_axis.scaling.max = 1
        ch.x_axis.title = "пулл"
        for ser in ch.series:
            ser.smooth = False
            ser.marker.symbol = "circle"
        s.chart(ch, f"A{s.row + 1}", w=16, h=7)
        ch2 = BarChart()
        ch2.title = "Смерти до конца вайпа"
        ch2.add_data(Reference(s.ws, min_col=7, min_row=f, max_row=l), titles_from_data=False)
        ch2.set_categories(Reference(s.ws, min_col=1, min_row=f, max_row=l))
        ch2.legend = None
        ch2.x_axis.title = "пулл"
        s.chart(ch2, f"J{s.row + 1}", w=12, h=7)


# -------------------------------------------------------------- Таймлайн
def _timeline(wb, R, demo):
    s = Sheet(wb, "Таймлайн", demo, {"A": 26, "B": 6, "C": 10, "D": 10, "E": 40})
    lanes = R["lanes"]
    ys = {lane: len(lanes) - i for i, lane in enumerate(lanes)}
    s.title("Таймлайн боя", "Y — номер дорожки: " + ", ".join(f"{ys[l]} {l}" for l in lanes) + ".")
    ev = sorted(R["timeline"], key=lambda e: (lanes.index(e["lane"]) if e["lane"] in lanes else 99, e["t"]))
    rows = [[e["lane"], ys.get(e["lane"], 0), e["t"], _mmss("C"), e["label"]] for e in ev]
    f, l = s.table(["Дорожка", "Y", "Время, с", "м:сс", "Событие"], rows, [None, F_INT, F_1, None, None])
    if not rows:
        return
    ch = ScatterChart()
    ch.title = "Смерти, Жажда крови и механики по времени"
    ch.x_axis.title = "Время боя, с"
    ch.y_axis.title = "дорожка"
    ch.y_axis.scaling.min = 0
    ch.y_axis.scaling.max = len(lanes) + 1
    ch.y_axis.majorUnit = 1
    ch.x_axis.scaling.min = 0
    ch.x_axis.scaling.max = R["info"]["duration_s"] + 5
    colors = ["C00000", "ED7D31", "2E75B6", "70AD47", "7030A0", "595959"]
    i = 0
    k = 0
    while i < len(ev):
        lane = ev[i]["lane"]
        j = i
        while j < len(ev) and ev[j]["lane"] == lane:
            j += 1
        ser = Series(Reference(s.ws, min_col=2, min_row=f + i, max_row=f + j - 1),
                     Reference(s.ws, min_col=3, min_row=f + i, max_row=f + j - 1), title=lane)
        ser.marker.symbol = "x" if lane == "Смерти" else "diamond" if lane == "Жажда крови" else "circle"
        ser.marker.size = 9
        col = colors[k % len(colors)]
        ser.marker.graphicalProperties = GraphicalProperties(solidFill=col)
        ser.marker.graphicalProperties.line.solidFill = col
        ser.graphicalProperties.line.noFill = True
        ch.series.append(ser)
        i, k = j, k + 1
    s.chart(ch, "G3", w=24, h=10)


# ------------------------------------------------------------- Против топа
def _plan_text(r: dict) -> str:
    """«Божественный гимн — Элария (откат 3:00, снова готов в 4:05); …»"""
    picks = r.get("picks")
    if picks is None:  # разборы старых версий
        return (r["cd"] + " — " + (r["player"] or "") + (" (как у топа)" if r["like_top"] else "")) if r["cd"] else "нет свободного кулдауна"
    text = "; ".join(f"{p['cd']} — {p['player']}{' (как у топа)' if p['like_top'] else ''}, откат {p['cooldown']}, "
                     f"снова готов в {p['ready']}" for p in picks) or "нет свободного кулдауна"
    if r.get("heal_picks"):  # кулдауны лекарей на этот пик — своё время нажатия
        text += ". Лекари: " + "; ".join(f"{p['cd']} — {p['player']} в {p['at']}" for p in r["heal_picks"])
    if r.get("spare"):
        text += ". Запасные: " + "; ".join(f"{x['cd']} — {x['player']}" for x in r["spare"])
    if r.get("after_end"):
        text = "[По топу: ваш бой сюда не дошёл, время — у лучших киллов, в заметку MRT не входит] " + text
    return text


def _vs_top(wb, R, demo):
    X = R.get("extras") or {}
    V, T = X.get("vs_top"), X.get("pull_trend") or {}
    if not V and not T.get("rows") and not X.get("roster_cds"):
        return
    s = Sheet(wb, "Против топа", demo, {"A": 26, "B": 30, "C": 36, "D": 22, "E": 30, "F": 10, "G": 10, "H": 10})
    s.title("Рейд против лучших киллов этого босса",
            "Пики урона сопоставлены по механике и её номеру в бою; кулдауны — по номеру способности.")
    if V:
        s.section("План рейдовых кулдаунов на следующий пулл",
                  "Кулдауны состава (классы, спеки, взятые таланты) расставлены на пики с учётом перезарядки.")
        f, l = s.table(["Нажать в", "Пик", "Кулдауны и кто", "", "У топа здесь"],
                       [[r["time"] + (f" ({r['phase_name']} +{r['phase_time']})" if (r.get("phase") or 0) > 1
                                      and r.get("phase_time") else ""),
                         r["mechanic"], _plan_text(r), "", r["top"] or "—"] for r in V["plan"]])
        s.ws.column_dimensions["C"].width = 70
        if V["plan"]:
            s.ws.conditional_formatting.add(f"C{f}:C{l}", CellIsRule(operator="equal", formula=['"нет свободного кулдауна"'], fill=RED))
        from .raid_top import mrt_note
        boss = (R.get("info") or {}).get("boss", "")
        for title, note in (("Заметка для MRT — сейвы рейда и кулдауны лекарей", mrt_note(V["plan"], boss)),):
            if note:
                s.section(title, "Скопируйте строки ниже целиком: в игре /mrt → Заметки → вставить → Отправить.")
                for line in note.split("\n"):
                    s.cell(s.row, 1, line)
                    s.row += 1
                s.row += 1
        s.section("Пики урона: ваш бой и лучшие киллы",
                  f"Закрыто кулдауном: у топа {V['top_cover']:.0%}, у вас {V['my_cover']:.0%}." if V["top_cover"] is not None
                  and V["my_cover"] is not None else None)
        def mine(r, k):
            if r["my"] is None:
                return "не дошли"
            lst = r.get(k)
            if lst is None:  # разбор старой версии — без деления
                lst = r["my"] if k == "my_raid" else []
            return ", ".join(lst) or ("нет" if k == "my_raid" else "—")

        f, l = s.table(["Когда", "Пик", "Ваши: на рейд", "Ваши: на себя (усиление лекаря)", "У топа закрыт",
                        "Топ: на рейд", "Топ: на себя"],
                       [[r["time"], r["mechanic"], mine(r, "my_raid"), mine(r, "my_self"), _v(r["top_share"], 3),
                         ", ".join(r.get("top_raid", r["top_cds"])) or "—", ", ".join(r.get("top_self") or []) or "—"]
                        for r in V["rows"]],
                       [None, None, None, None, F_PCT, None, None])
        s.note("На рейд — действует сразу на всех (гимн, тотем, барьер). На себя — бафф лекаря, "
               "усиливающий его исцеление (Апофеоз, Перерождение, Древо Жизни).")
        if V["rows"]:
            s.ws.conditional_formatting.add(f"C{f}:C{l}", CellIsRule(operator="equal", formula=['"нет"'], fill=RED))
        s.section("Лучшие киллы")
        s.table(["Время боя", "Гильдия", "Ссылка"], [[k["duration"], k["guild"], k["url"]] for k in V["kills"]],
                links={2: lambda i: V["kills"][i]["url"]})
    if X.get("roster_cds"):
        s.section("Рейдовые кулдауны состава", "У кого какой рейдовый кулдаун есть: откат, сколько раз нажат, сколько можно за бой.")
        s.table(["Игрок", "Кулдаун", "Откуда известно", "Откат, с", "Нажато / можно за бой"],
                [[c["player"], c["name"], c["source"], _v(c["cd"], 0), f"{c['used']} / {c['max_uses'] or '—'}"]
                 for c in X["roster_cds"]])
    if T.get("rows"):
        s.section("От чего умирает рейд по пуллам", "Смерти до конца вайпа; последние 15 с вайпа не считаются.")
        s.table(["Механика"] + [f"Пулл {n}" for n in T["pulls"]] + ["Всего"],
                [[r["name"]] + r["counts"] + [r["total"]] for r in T["rows"]],
                [None] + [F_INT] * (len(T["pulls"]) + 1))


# ============================================================ сейвы на всех боссов
def write_saves_workbook(R: dict, path: str | Path, wb=None) -> Path:
    """План сейвов на каждого босса: лист на босса (план + заметка MRT) и общий лист с заметками."""
    from .compare import _fmt_t
    own = wb is None   # wb передан — листы дописываются в общую книгу (разбор боя целиком)
    if own:
        wb = Workbook()
        wb.remove(wb.active)
    demo = R["info"]["demo"]
    s = Sheet(wb, "Все боссы", demo, {"A": 34, "B": 14, "C": 16, "D": 12, "E": 60})
    s.title(f"Рейдовые сейвы: {R['info']['title'] or R['info']['zone']}",
            "На каждого босса — последний килл, а без киллов — лучший пулл. Подробный план — на листе босса.")
    s.table(["Босс", "Сложность", "Бой", "Пиков в плане", "Заметка MRT — скопировать целиком (ячейки ниже)"],
            [[b["boss"], b["difficulty"], ("килл " if b["kill"] else "вайп ") + b["duration"], len(b["plan"]),
              "есть" if b["mrt"] else "нет назначенных кулдаунов"] for b in R["bosses"]], [None, None, None, F_INT, None])
    for b in R["bosses"]:
        for label, note in (("MRT", b.get("mrt")),):
            if not note:
                continue
            s.section(f"{label}: {b['boss']} ({b['difficulty']})")
            for line in note.split("\n"):
                s.cell(s.row, 1, line)
                s.row += 1
    if R.get("skipped"):
        s.section("Пропущены")
        for x in R["skipped"]:
            s.cell(s.row, 1, f"{x['boss']} ({x['difficulty']}): {x['reason']}")
            s.row += 1
    used = set()
    for b in R["bosses"]:
        name = (b["boss"][:24] or "Босс").replace("/", "-").replace(":", " ")
        while name in used:
            name += " "
        used.add(name)
        t = Sheet(wb, name, demo, {"A": 26, "B": 30, "C": 80, "D": 30})
        t.title(f"{b['boss']} — {b['difficulty']}",
                ("Килл" if b["kill"] else "Лучший пулл") + f" {b['duration']}, пуллов за вечер: {b['pulls']}. " +
                ("Фазы: " + " · ".join(f"{p['name']} с {_fmt_t(p['t'])}" for p in b["phases"]) if len(b["phases"]) > 1 else ""))
        t.table(["Нажать в", "Пик", "Кулдауны и кто", "У топа здесь"],
                [[r["time"] + (f" ({r['phase_name']} +{r['phase_time']})" if (r.get("phase") or 0) > 1 and r.get("phase_time") else ""),
                  r["mechanic"], _plan_text(r), r.get("top") or "—"] for r in b["plan"]])
        for title, note in (("Заметка для MRT — сейвы рейда и кулдауны лекарей", b.get("mrt")),):
            if note:
                t.section(title, "В игре: /mrt → Заметки → вставить → Отправить.")
                for line in note.split("\n"):
                    t.cell(t.row, 1, line)
                    t.row += 1
    if not own:
        return wb
    path = Path(path)
    wb.save(path)
    return path
