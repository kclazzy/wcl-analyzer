"""Сквозной тест без сети: фейковый клиент отвечает событиями в формате WCL API.

Запуск: python tests/test_pipeline.py
"""
from __future__ import annotations

import json
import random
import sys
import tempfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from wcl_analyzer.collect import collect_reference, load_my_log  # noqa: E402
from wcl_analyzer.compare import compare  # noqa: E402
from wcl_analyzer.demo import my_policy, simulate, top_policy  # noqa: E402
from wcl_analyzer.excel_report import write_compare_workbook  # noqa: E402
from wcl_analyzer.logs import parse_report_url  # noqa: E402
from wcl_analyzer.reference import build_reference  # noqa: E402


class FakeClient:
    """Минимальная замена WCLClient: те же методы, данные из демо-симуляции."""

    def __init__(self):
        self.reports, self.raws, self.ranks = {}, {}, []
        self.requests_made = self.cache_hits = 0
        rng = random.Random(7)
        for i in range(30):
            dur = rng.uniform(250, 360)
            code = f"TOPREPORT{i:03d}"
            rep, fight, actor, raw = simulate(top_policy(rng, dur), dur, rng.uniform(691, 699), 100 + i,
                                              f"Top{i}", code, 1_788_000_000_000, (1, 2))
            self.reports[code], self.raws[code] = rep, raw
            total = sum(e["total"] for e in raw["dmg_table"]["entries"])
            self.ranks.append({"name": f"Top{i}", "amount": total / dur, "duration": dur * 1000,
                               "report": {"code": code, "fightID": 1, "startTime": 0}})
        self.ranks.sort(key=lambda r: -r["amount"])
        rep, fight, actor, raw = simulate(my_policy(rng), 305.0, 690, 999, "Me", "MYREPORT0001",
                                          1_788_000_000_000, (1, 3), with_damage_events=True)
        self.reports["MYREPORT0001"], self.raws["MYREPORT0001"] = rep, raw

    def report(self, code):
        return self.reports[code]

    def player_details(self, code, fight_id):
        a = self.reports[code]["masterData"]["actors"][0]
        return {"dps": [{"name": a["name"], "id": a["id"], "type": "Mage", "specs": [{"spec": "Frost"}]}]}

    def events(self, code, fight_id, start, end, data_type, source_id=None, target_id=None,
               hostility=None, include_resources=False):
        raw = self.raws[code]
        key = {"Casts": "boss_casts" if hostility == "Enemies" else "casts", "Buffs": "buffs",
               "Debuffs": "debuffs", "DamageTaken": "dmg_taken", "Deaths": "deaths",
               "CombatantInfo": "combatant", "DamageDone": "dmg_done"}[data_type]
        return raw.get(key, [])

    def damage_table(self, code, fight_id, source_id):
        return self.raws[code]["dmg_table"]

    def rankings(self, encounter_id, class_name, spec_name, difficulty, page=1, force=False, max_age_s=None):
        import time
        if force or not getattr(self, "_ranks_at", None):
            self._ranks_at = time.time()
        return {"name": "Демо-босс", "_fetched_at": self._ranks_at,
                "characterRankings": {"rankings": self.ranks[(page - 1) * 100: page * 100], "hasMorePages": False}}


def test_refresh():
    """Эталон сохраняется в список и обновляется принудительно."""
    import os
    import time
    from wcl_analyzer import settings
    from wcl_analyzer.api import Cache
    from wcl_analyzer.collect import refresh_refs
    tmp = Path(tempfile.mkdtemp())
    old = os.getcwd()
    os.chdir(tmp)
    try:
        client = FakeClient()
        client.cache = Cache(tmp / "wcl_cache.sqlite")
        collect_reference(client, 9999, "Mage", "Frost", 5, top_n=10, duration=300, log=lambda *_: None)
        refs = settings.list_refs(tmp / "wcl_cache.sqlite")
        assert len(refs) == 1 and refs[0]["n_logs"] == 10 and not refs[0]["stale"]
        first = refs[0]["collected_at"]
        time.sleep(0.05)
        assert refresh_refs(client, log=lambda *_: None) == 1
        assert settings.list_refs(tmp / "wcl_cache.sqlite")[0]["collected_at"] > first
        settings.save_user_prefs(tmp / "wcl_cache.sqlite", "local", ref_max_age_days=1)
        assert settings.user_prefs(tmp / "wcl_cache.sqlite", "local")["ref_max_age_days"] == 1
        assert settings.list_refs(tmp / "wcl_cache.sqlite", "someone-else") == []
        assert settings.load()["token"]
        print("OK эталоны: сохранение и обновление")
    finally:
        os.chdir(old)


def main():
    assert parse_report_url("https://www.warcraftlogs.com/reports/AbCd1234EfGh#fight=12&source=5") == \
        ("AbCd1234EfGh", "12", 5)
    assert parse_report_url("https://www.warcraftlogs.com/reports/AbCd1234EfGh") == ("AbCd1234EfGh", None, None)
    assert parse_report_url("AbCd1234EfGh")[0] == "AbCd1234EfGh"

    client = FakeClient()
    me = load_my_log(client, "https://www.warcraftlogs.com/reports/MYREPORT0001#fight=last&source=7")
    assert me.name == "Me" and me.spec == "Frost" and me.cls == "Mage"
    assert me.casts and me.dmg_timeline and me.ilvl

    tops, label = collect_reference(client, me.encounter_id, me.cls, me.spec, me.difficulty,
                                    top_n=25, duration=me.duration, log=lambda *_: None)
    assert "±" in label, label
    assert all(abs(t.duration - me.duration) / me.duration <= 0.2 for t in tops)

    ref = build_reference(tops, label=label)
    r = compare(me, ref)
    sections = {f.section for f in r.findings}
    for expected in ("Простой", "Кулдауны", "Бурст", "Защита", "Проки", "Зелья"):
        assert expected in sections, (expected, sections)
    assert len(r.top5) == 5 and r.plan
    out = Path(tempfile.gettempdir()) / "wcl_test_compare.xlsx"
    write_compare_workbook(r, out)
    assert out.exists() and out.stat().st_size > 20_000
    print(f"OK: {len(tops)} эталонных логов ({label}), {len(r.findings)} отличий, "
          f"надёжность {r.reliability['overall']}, отчёт {out}")


def test_gear_and_trinkets():
    from wcl_analyzer.demo import demo_logs
    from wcl_analyzer.logs import gear_slots
    tops, me = demo_logs()
    ref = build_reference(tops)
    r = compare(me, ref)
    assert ref.spells[443124].category == "trinket", "тринкет не распознан по иконке аксессуара"
    kinds = {b["name"]: b["kind"] for b in r.tables["burst"]}
    assert kinds["Застывший осколок"] == "Тринкет" and kinds["Стылая кровь"] == "Способность"
    assert kinds.get("Berserking") == "Расовая"
    g = r.tables["gear"]
    assert g["missing_enchants"] == ["Кольцо 2"] and g["missing_temp"] == ["Правая рука"], g["missing_enchants"]
    neck = next(x for x in g["rows"] if x["slot"] == 1)
    assert neck["few_gems"] and not neck["enchantable"]
    assert any(f.key == "gear:enchant" for f in r.findings)
    # Слоты рубашки и накидки и пустые слоты пропускаются; номер слота берётся из поля slot, если оно есть
    sl = gear_slots([{"id": 5, "slot": 3}, {"id": 0}, {"id": 7, "slot": 15, "permanentEnchant": 1, "temporaryEnchant": 2,
                                                       "gems": [{"id": 1}], "bonusIDs": [10, "11"]}])
    assert sl == [{"slot": 15, "id": 7, "ilvl": None, "enchant": 1, "enchant_name": None, "temp": 2, "gems": 1,
                   "name": None, "icon": None, "bonus": [10, 11]}], sl
    with tempfile.TemporaryDirectory() as d:
        from openpyxl import load_workbook
        wb = load_workbook(write_compare_workbook(r, Path(d) / "x.xlsx"))
        assert "Экипировка" in wb.sheetnames
    print("OK экипировка: зачарования и камни против топа, тип действий в окне бурста (тринкет/способность)")


def test_battle_analysis():
    from wcl_analyzer.battle import build_battle
    from wcl_analyzer.battle import positions as P
    from wcl_analyzer.demo import demo_logs
    # координаты: шум убирается, движения склеиваются, мелкие отбрасываются
    pts = [(t, 0.2 * (t % 2), 0.0) for t in range(0, 10)]            # стоит на месте с дрожанием
    pts += [(10.0, 4.0, 0.0), (11.0, 9.0, 0.0), (12.0, 12.0, 0.0)]    # пробежал 12 ярдов
    pts += [(t, 12.0, 0.1 * (t % 3)) for t in range(13, 20)]
    pts += [(20.0, 13.5, 0.0), (21.0, 12.0, 0.0)]                     # шаг на месте — не движение
    segs = P.segments(P.denoise(pts))
    assert len(segs) == 1 and abs(segs[0]["dist"] - 12) < 1 and 9 <= segs[0]["start"] <= 10, segs
    assert P.rdp([(0, 0, 0), (1, 1, 0.01), (2, 2, 0)], 0.5) == [(0, 0, 0), (2, 2, 0)]
    rel, origin = P.relative([(0, 105, 210)], [(0, 100, 200)])
    assert origin == "boss" and rel == [(0, 5, 10)]

    tops, me = demo_logs()
    r = compare(me, build_reference(tops))
    B = build_battle(r)
    ev = B["events"]
    assert 10 <= len(ev) <= 20, len(ev)
    types = {t for e in ev for t in e["types"]}
    assert {"opener", "phase", "end", "burst", "defensive", "movement", "switch"} <= types, types
    assert ev[0]["t"] == 0 and ev[-1]["type"] == "end" and all(a["t"] <= b["t"] for a, b in zip(ev, ev[1:]))
    assert all(1 <= len(e["lines"]) <= 4 for e in ev if e["type"] not in ("phase",)), [e for e in ev if len(e["lines"]) > 4]
    M = B["movement"]
    assert M["available"] and M["top_n"] >= 20 and M["total"] > M["top_total"], (M["total"], M["top_total"])
    reasons = {m["reason"] for m in M["moves"]}
    assert "mechanic" in reasons and "unknown" in reasons
    extra = [m for m in M["moves"] if 118 <= m["start"] <= 123][0]   # лишнее перемещение: топ здесь стоит
    assert extra["reason"] == "unknown" and extra["top_share"] < 0.3, extra
    sw = [e for e in ev if e["type"] == "switch" or "switch" in e["types"]][0]
    assert sw["diff"]["t_s"] > 1 and "Ледяной элементаль" in sw["title"], sw
    assert 3 <= len(B["differences"]) <= 5 and all(d["level"] in ("high", "medium") for d in B["differences"])
    assert 1 <= len(B["summary"]) <= 5 and B["dps"]["gap"] < 0
    json.dumps(B)
    # лог без координат: лента есть, карта честно недоступна
    me.positions, me.boss_positions = [], []
    B2 = build_battle(compare(me, build_reference(tops)))
    assert not B2["movement"]["available"] and len(B2["events"]) >= 10
    assert any("Координат" in x for x in B2["summary"])
    with tempfile.TemporaryDirectory() as d:
        from openpyxl import load_workbook
        assert "Бой по шагам" in load_workbook(write_compare_workbook(r, Path(d) / "b.xlsx")).sheetnames
    print("OK бой по шагам: 10–20 событий, движение с причиной только по данным, смена цели и бурст против топа")


def test_plan_no_duplicate_ability():
    from wcl_analyzer.raid_top import make_plan
    X = {"roster_cds": [{"pid": 1, "player": "Торвин", "name": "Ободряющий клич", "id": 97462, "cd": 180, "cls": "Warrior"},
                        {"pid": 2, "player": "Брам", "name": "Ободряющий клич", "id": 97462, "cd": 180, "cls": "Warrior"},
                        {"pid": 3, "player": "Элария", "name": "Божественный гимн", "id": 64843, "cd": 180, "cls": "Priest"}],
         "spikes": [{"t": 60, "peak_t": 62, "ability": "Волна", "ability_id": 5, "k": 1, "damage": 100, "deaths": 2},
                    {"t": 120, "peak_t": 122, "ability": "Волна", "ability_id": 5, "k": 2, "damage": 30}]}
    plan = make_plan(X, {}, [], {})
    assert [p["cd"] for p in plan[0]["picks"]] == ["Божественный гимн", "Ободряющий клич"], plan[0]["picks"]
    # второй клич — на следующий пик (от другого воина), а не вдогонку первому
    first_warrior = next(p["player"] for p in plan[0]["picks"] if p["cd"] == "Ободряющий клич")
    assert plan[1]["picks"][0]["cd"] == "Ободряющий клич" and plan[1]["picks"][0]["player"] != first_warrior
    print("OK план сейвов: одна способность на пик не дублируется (два воина — клич на разные пики)")


def test_plan_long_fight():
    """Бой длиннее, чем у топа: сильные пики в конце, которых у топа нет, тоже получают сейвы."""
    from collections import Counter
    from wcl_analyzer.raid_top import make_plan
    X = {"roster_cds": [{"pid": 1, "player": "Торвин", "name": "Ободряющий клич", "id": 97462, "cd": 180, "cls": "Warrior"},
                        {"pid": 3, "player": "Элария", "name": "Божественный гимн", "id": 64843, "cd": 180, "cls": "Priest"},
                        {"pid": 4, "player": "Таргун", "name": "Тотем целительного потока", "id": 108280, "cd": 180, "cls": "Shaman"}],
         "spikes": [{"t": t, "peak_t": t + 2, "ability": "Волна", "ability_id": 5, "k": i + 1, "damage": 100}
                    for i, t in enumerate((60, 120, 200))]
                   + [{"t": t, "peak_t": t + 2, "ability": "Ярость", "ability_id": 6, "k": i + 1, "damage": 300}
                      for i, t in enumerate((400, 410))]}
    ref = {(5, i): {"cds": Counter({c: 5}), "n_cds": [1]} for i, c in ((1, 64843), (2, 108280), (3, 97462))}
    plan = make_plan(X, ref, [], {})
    late = [r for r in plan if r["mechanic"].startswith("«Ярость»")]
    assert all(r["picks"] and r["beyond_top"] for r in late), late
    assert all(r["picks"] for r in plan), [r["mechanic"] for r in plan if not r["picks"]]
    print("OK план сейвов: сильные пики в конце длинного боя (у топа их нет) тоже закрыты")


def test_player_all_bosses():
    """Игрок на всех боссах: эталон топ-1; не хватило лимита — оставшиеся боссы в «pending», потом «Продолжить»."""
    import copy
    from wcl_analyzer import web
    client = FakeClient()
    rep = client.reports["MYREPORT0001"]
    f1 = rep["fights"][0]
    f2 = {**copy.deepcopy(f1), "id": 2, "encounterID": int(f1.get("encounterID") or 1) + 1, "name": "Второй босс"}
    f3 = {**copy.deepcopy(f1), "id": 3, "encounterID": int(f1.get("encounterID") or 1) + 2, "name": "Третий босс"}
    rep["fights"] = [f1, f2, f3]
    points = iter([1000, 960, 960, 30, 30, 30, 30])
    client.points_left = lambda: next(points)
    client.rate_limit = lambda: {"pointsResetIn": 1500}
    web.CLIENT_FACTORY = lambda creds: client
    # в тестовом клиенте логи топа есть только для первого босса — для остальных берём тот же эталон
    orig_ref, cache = web._player_ref, {}

    def ref_any_boss(client_, me, difficulty, params_, log, meta):
        assert params_["ref"] == "top1" and params_["mythic"] is False
        if "r" not in cache:
            cache["r"] = orig_ref(client_, me, difficulty, params_, log, meta)
        return cache["r"]
    web._player_ref = ref_any_boss
    job = {"id": "t" * 32, "progress": 0.0}
    params = {"url": "https://www.warcraftlogs.com/reports/MYREPORT0001", "mode": "allbosses", "actor": "7"}
    web._run_player_all(job, params, None, lambda *_: None)
    R = job["result"]
    assert R["mode"] == "playerall" and len(R["bosses"]) == 1 and R["bosses"][0]["detail"]["info"]["ref_n"] == 1, R["bosses"]
    assert [(p["fight_id"], p["actor"]) for p in R["pending"]] == [(2, 7), (3, 7)] and R["reset_in"] == 1500, R["pending"]
    assert job["xlsx"][:2] == b"PK"
    # «Продолжить»: только оставшиеся боссы, краткие строки уже разобранных — для общего Excel
    client.points_left = lambda: 1000
    job2 = {"id": "u" * 32, "progress": 0.0}
    prev = [{k: R["bosses"][0][k] for k in ("boss", "difficulty", "kill", "dps", "ref_dps", "actions")}]
    web._run_player_all(job2, {**params, "units": [[p["fight_id"], p["actor"]] for p in R["pending"]], "prev": prev},
                        None, lambda *_: None)
    R2 = job2["result"]
    assert [b["fight_id"] for b in R2["bosses"]] == [2, 3] and not R2["pending"], (R2["bosses"], R2["pending"])
    # «Все игроки»: DPS каждого боя на каждом боссе
    job3 = {"id": "v" * 32, "progress": 0.0}
    web._run_player_all(job3, {**params, "actor": "all"}, None, lambda *_: None)
    R3 = job3["result"]
    assert R3["everyone"] and len(R3["bosses"]) == 3 and {b["player"] for b in R3["bosses"]} == {"Me"}, R3["bosses"]
    assert "all_events" not in R3["bosses"][0]["detail"]["battle"]
    web._player_ref = orig_ref
    print("OK игрок на всех боссах: топ-1, остановка по лимиту WCL и продолжение с оставшихся")


def test_saves_all_bosses():
    from wcl_analyzer.excel_raid import write_saves_workbook
    from wcl_analyzer.raid_demo import DEMO_URL, FakeRaidClient
    from wcl_analyzer.raid_saves import pick_fights, run_raid_saves
    f = lambda i, enc, diff, kill, pct, t: {"id": i, "encounterID": enc, "difficulty": diff, "kill": kill,  # noqa: E731
                                          "fightPercentage": pct, "startTime": t, "endTime": t + 300_000}
    rep = {"fights": [f(1, 10, 4, False, 40.0, 0), f(2, 10, 4, True, 0, 1e6), f(3, 10, 4, True, 0, 2e6),  # килл ×2
                      f(4, 11, 5, False, 30.0, 3e6), f(5, 11, 5, False, 12.5, 4e6),                    # лучший вайп
                      f(6, 12, 3, True, 0, 5e6), f(7, 13, 1, True, 0, 6e6), f(8, 0, 0, False, None, 7e6)]}
    got = [(c["fight"]["id"], c["pulls"], c["kills"]) for c in pick_fights(rep)]
    assert got == [(3, 3, 2), (5, 2, 0)], got   # обычная сложность, ЛФР и треш пропущены
    R = run_raid_saves(FakeRaidClient(), DEMO_URL, log=lambda *_: None, talent_data=[])
    b = R["bosses"][0]
    assert R["mode"] == "saves" and b["plan"] and b["mrt"].startswith("Сейвы: Демо-босс") and len(b["phases"]) == 3
    with tempfile.TemporaryDirectory() as d:
        from openpyxl import load_workbook
        wb = load_workbook(write_saves_workbook(R, Path(d) / "s.xlsx"))
        assert wb.sheetnames[0] == "Все боссы" and len(wb.sheetnames) == 2
    print("OK сейвы на всех боссов: героическая и эпохальная, килл или лучший пулл, заметка MRT на каждого")


if __name__ == "__main__":
    test_plan_no_duplicate_ability()
    test_player_all_bosses()
    test_plan_long_fight()
    test_saves_all_bosses()
    test_battle_analysis()
    test_gear_and_trinkets()
    main()


def test_raid():
    from wcl_analyzer.excel_raid import write_raid_workbook
    from wcl_analyzer.raid import run_raid
    from wcl_analyzer.raid_demo import CODE, DEMO_URL, FakeRaidClient

    R = run_raid(FakeRaidClient(), DEMO_URL, log=lambda *_: None, talent_data=[])
    Xs = R["extras"]
    assert Xs["plan"] and all(r["picks"] for r in Xs["plan"]), "во вкладке «Полученный урон и сейвы» нет плана"
    assert Xs["roster_cds"] and Xs["saves_brief"], "нет кулдаунов состава или выжимки сейвов"
    # пик со смертью — опасный: несколько сейвов разных игроков; у каждого пика — запасные, уже откатанные
    danger = [r for r in Xs["plan"] if r["deaths"]]
    assert danger and all(len(r["picks"]) >= 2 and len({p["player"] for p in r["picks"]}) == len(r["picks"]) for r in danger)
    assert all(r["spare"] for r in Xs["plan"]), "нет запасных вариантов"
    for r in Xs["plan"]:  # одна и та же способность на пик не повторяется — ни в назначенных, ни в запасных
        cds_on_peak = [p["cd"] for p in r["picks"]] + [x["cd"] for x in r["spare"]]
        assert len(cds_on_peak) == len(set(cds_on_peak)), r
    # Заметка для MRT: таймер от пулла, имя в цвете класса, иконка способности
    from wcl_analyzer.raid_top import mrt_note
    note = mrt_note(Xs["plan"], "Демо")
    lines = note.split("\n")
    assert lines[0] == "Сейвы: Демо" and len(lines) == 1 + len(Xs["plan"]), note
    assert all(ln.startswith("{time:") and "{spell:" in ln and "|cff" in ln and "|r" in ln for ln in lines[1:]), note
    assert "«" not in note and "Ледяная волна" not in note and "{spell:900002}" in note  # механика — иконкой
    cdmap = {(c["player"], c["name"]): c["cd"] for c in Xs["roster_cds"]}
    uses = {}
    for r in Xs["plan"]:
        for p in r["picks"]:
            uses.setdefault((p["player"], p["cd"]), []).append(r["t"])
    for k, ts in uses.items():
        assert all(b - a >= cdmap[k] - 0.01 for a, b in zip(sorted(ts), sorted(ts)[1:])), (k, ts)
    assert R["info"]["size"] == 20 and R["info"]["kill"]
    cats = {a["name"]: a["category"] for a in R["abilities"]}
    assert cats["Ледяная волна"] == "По всему рейду" and cats["Лужа холода"] == "Выборочно"
    assert cats["Сокрушение"] == "По танкам" and cats["Осколки льда"] == "Выборочно"
    texts = {(i["player"], i["kind"]) for i in R["issues"]}
    for expected in (("Лиана", "Смерть"), ("Мирель", "Смерть"), ("Торвин", "Механики"), ("Мирель", "Механики"),
                     ("Брам", "Активность"), ("Квелл", "Механики")):
        assert expected in texts, expected
    assert R["extras"]["consumables"]["no_potion"] == ["Ровен"], R["extras"]["consumables"]  # зелье — в «Расходниках»
    assert R["deaths"][0]["player"] == "Лиана" and R["deaths"][0]["first"]
    assert len(R["pulls"]) == 6 and R["pulls"][-1]["kill"] and R["pulls"][0]["boss_pct"] == 78.4
    out = Path(tempfile.gettempdir()) / "wcl_test_raid.xlsx"
    write_raid_workbook(R, out)
    assert out.exists()
    print(f"OK рейд: {len(R['players'])} игроков, {len(R['issues'])} замечаний, {len(R['pulls'])} пуллов")


if __name__ == "__main__":
    test_raid()
    test_refresh()


def test_analysis_quality():
    """Смерть не порождает ложных находок; топ против остальных почти без находок;
    дебаффы на боссе и выжимка на месте."""
    import copy
    from wcl_analyzer.compare import compare
    from wcl_analyzer.demo import demo_logs
    from wcl_analyzer.raid import run_raid
    from wcl_analyzer.raid_demo import DEMO_AVOIDABLE, DEMO_URL, FakeRaidClient
    from wcl_analyzer.reference import build_reference

    tops, me = demo_logs()
    # 1. Топ, умерший на середине боя
    v = copy.deepcopy(tops[0])
    td = v.duration / 2
    v.casts = [c for c in v.casts if c.t <= td]
    v.deaths = [td]
    v.dmg_taken = [h for h in v.dmg_taken if h[0] <= td]
    v.buff_events = [e for e in v.buff_events if e[0] <= td]
    v.buffs = {k: [(a, min(b, td)) for a, b in iv if a <= td] for k, iv in v.buffs.items()}
    v.debuffs = {k: [(a, min(b, td)) for a, b in iv if a <= td] for k, iv in v.debuffs.items()}
    v.dps /= 2
    r = compare(v, build_reference(tops[1:]))
    high = [f for f in r.findings if f.impact_level == "HIGH" and not f.noise]
    assert [f.section for f in high] == ["Смерти"], [f.title for f in high]

    # 2. Шумовой порог: топ против остальных — суммарно меньше 2% «потерь»
    for i in (0, 7, 14):
        rr = compare(tops[i], build_reference([t for j, t in enumerate(tops) if j != i]))
        tot = sum(f.impact or 0 for f in rr.findings if not f.noise and not f.context)
        assert tot < 0.02, (i, tot)

    # 3. Мой лог: дебаффы на боссе, внешние баффы, выжимка
    ref = build_reference(tops, me=me)
    r = compare(me, ref)
    keys = {f.key for f in r.findings}
    assert "raiddebuff:1490" in keys and "ext:10060" in keys, keys
    assert any(k.startswith("vuln:") for k in keys), keys
    assert any(k.startswith("prio:") for k in keys) and "res:waste" in keys, keys
    assert 1 <= len(r.brief["actions"]) <= 3 and r.brief["context"], r.brief
    assert all(not f.context for f in r.top5)
    assert ref.build_note, "фильтр по билду не сработал"

    # 4. Рейд: выжимка и дополнительные проверки
    R = run_raid(FakeRaidClient(), DEMO_URL, None, log=lambda m: None, avoidable={str(x) for x in DEMO_AVOIDABLE})
    X = R["extras"]
    assert 2 <= len(R["brief"]) <= 7, R["brief"]
    assert X["missed_kicks"] and X["consumables"]["no_flask"] == ["Брам"], X["consumables"]
    assert isinstance(X["consumables"]["no_potion"], list), X["consumables"]
    assert not any("зелье" in i["text"].lower() for i in R["issues"]), "боевое зелье — в «Расходниках», не в «Что проверить»"
    assert [a["name"] for a in X["adds"]["low"]] == ["Норра"], X["adds"]
    assert any(not s["covered_by"] for s in X["spikes"]) and X["avoidable"]
    assert X["heaviest"][0]["abilities"][0] == "Ледяная волна", X["heaviest"][0]
    names = {c["name"]: c for c in X["raid_cds"]}
    assert {"Божественный гимн", "Тотем целительного потока", "Ободряющий клич"} <= set(names), names
    assert names["Ободряющий клич"]["on_peak"] is None and names["Божественный гимн"]["on_peak"]
    assert len(X["damage_timeline"]) >= 60
    V = X["vs_top"]
    assert V and V["n"] == 5 and V["top_cover"] == 1.0 and abs(V["my_cover"] - 2 / 3) < 0.01, V and V["my_cover"]
    first = V["rows"][0]  # на рейд и на себя (бафф лекаря) — раздельно
    assert first["my_self"] == ["Апофеоз (Элария)"] and first["my_raid"] == ["Божественный гимн (Элария)"], first
    second = V["rows"][1]
    assert second["top_self"] == ["Перерождение"] and "Перерождение" not in second["top_raid"], second
    third = [r for r in V["plan"] if r["mechanic"].endswith("№3")][0]
    # Гимн нажат на волну №1 (фаза 1); волна №3 — уже в фазе 2, ровно через 180 с отката. Если фаза 2 в
    # следующем пулле начнётся раньше, гимн не успеет откатиться: план берёт другой кулдаун (запас на смещение фазы)
    assert third["picks"] and third["cd"] != "Божественный гимн" and third["top"].startswith("Божественный гимн"), third
    assert third["phase"] == 3 and third["phase_name"] == "Фаза 2" and third["phase_time"] == "0:45", third
    second = [r for r in V["plan"] if r["mechanic"].endswith("№2")][0]
    assert second["intermission"] and second["phase_name"] == "Ледяной шторм", second
    assert "{time:0:45,p3}" in third["mrt"] and "{time:1:05}" in V["plan"][0]["mrt"], (third["mrt"], V["plan"][0]["mrt"])
    assert [p["n"] for p in R["info"]["phases"]] == [1, 2, 3]
    assert X["pull_trend"]["rows"][0]["name"] == "Ледяная волна", X["pull_trend"]
    # Вайп: план покрывает и пики, до которых рейд не дошёл
    Rw = run_raid(FakeRaidClient(), DEMO_URL, 1, log=lambda m: None)
    assert any(r["my"] is None for r in Rw["extras"]["vs_top"]["rows"]), Rw["extras"]["vs_top"]["rows"]
    print("OK анализ: смерть нормирована, шум топа отсечён, дебаффы босса и выжимка на месте")


if __name__ == "__main__":
    test_analysis_quality()


def test_raid_rotation():
    """Ротация всего рейда: все DPS разобраны, общие проблемы найдены, Excel собирается."""
    import json
    import tempfile
    from wcl_analyzer import web
    from wcl_analyzer.raid_rotation import run_demo
    R = run_demo(log=lambda m: None)
    ok = [r for r in R["rows"] if "error" not in r]
    assert len(ok) == 5, R["rows"]
    worst = min(ok, key=lambda r: r["gap"])
    assert worst["name"] == "Ильвен" and worst["actions"][0]["title"].startswith("Смерть"), worst["actions"]
    best = max(ok, key=lambda r: r["gap"])
    assert best["name"] == "Кассия", best["name"]
    assert any("у 3 из 5" in line for line in R["brief"]), R["brief"]
    job = {"id": "t", "progress": 0.0, "log": []}
    web._run_raid_rotation(job, {"mode": "raidrot", "demo": True}, None, lambda m: None)
    assert len(json.dumps(job["result"])) < 2_000_000 and job["xlsx"]
    from openpyxl import load_workbook
    with tempfile.NamedTemporaryFile(suffix=".xlsx") as f:
        f.write(job["xlsx"]); f.flush()
        wb = load_workbook(f.name)
        assert wb.sheetnames[:2] == ["Ротация рейда", "Что исправить"] and "Ильвен" in wb.sheetnames, wb.sheetnames
    print("OK ротация рейда: 5 DPS, общие проблемы, переход к игроку, Excel")


if __name__ == "__main__":
    test_raid_rotation()


def test_batched_fetch():
    """Пакетный запрос (все выборки игрока одним GraphQL) даёт те же данные, что и десять отдельных."""
    import re
    from wcl_analyzer.api import Cache, WCLClient
    from wcl_analyzer.logs import _fetch_raw_batched, fetch_raw

    fake = FakeClient()
    code = "TOPREPORT001"
    report = fake.report(code)
    fight = report["fights"][0]
    actor = report["masterData"]["actors"][0]

    class Batched(WCLClient):
        calls = 0

        def query(self, q, variables=None, use_cache=True, max_age_s=None):
            Batched.calls += 1
            rep = {}
            for alias, args in re.findall(r"(e\d+): events\(([^)]*)\)", q):
                kv = dict(re.findall(r"(\w+): ([^,]+)", args))
                data = fake.events(code, 1, float(kv["startTime"]), float(kv["endTime"]), kv["dataType"],
                                   source_id=int(kv["sourceID"]) if "sourceID" in kv else None,
                                   target_id=int(kv["targetID"]) if "targetID" in kv else None,
                                   hostility=kv.get("hostilityType"))
                rep[alias] = {"data": data, "nextPageTimestamp": None}
            for alias in re.findall(r"(t\d+): table\(", q):
                rep[alias] = {"data": fake.damage_table(code, 1, actor["id"])}
            return {"reportData": {"report": rep}}

    class FakeRes(FakeClient):
        def events(self, *a, **k):
            return [] if a[4] == "Resources" else super().events(*a, **k)

    fake = FakeRes()
    from wcl_analyzer.api import MemoryCache
    b = Batched("id", "secret", MemoryCache(), verbose=False)
    got = _fetch_raw_batched(b, report, fight, actor["id"], False)
    want = fetch_raw(fake, report, fight, actor["id"])
    for k in ("casts", "buffs", "debuffs", "dmg_taken", "deaths", "boss_casts", "combatant", "boss_debuffs"):
        assert len(got[k]) == len(want[k]), (k, len(got[k]), len(want[k]))
    assert got["dmg_table"] == want["dmg_table"] and Batched.calls == 1, Batched.calls
    print("OK пакетный запрос: те же данные одним запросом вместо", 10)


def test_limit_no_hang():
    """Исчерпан часовой лимит: поиск боя сразу объясняет, а не висит; запрос лимита не зацикливается."""
    import time as _t

    from wcl_analyzer.api import MemoryCache, WCLClient, WCLError

    class Resp:
        status_code = 429
        text = "Too many requests"

        def json(self):
            return {}

    class Sess:
        posts = 0

        def post(self, url, **k):
            Sess.posts += 1
            if "token" in url:
                r = Resp(); r.status_code = 200; r.json = lambda: {"access_token": "t", "expires_in": 3600}
                return r
            return Resp()

    c = WCLClient("id", "secret", MemoryCache(), verbose=False)
    c.session = Sess()
    c.max_wait_s = 0
    t0 = _t.time()
    try:
        c.query("query { x }")
        raise AssertionError("должна быть ошибка")
    except WCLError as e:
        assert "лимит" in str(e) and "мин" in str(e), e
    assert _t.time() - t0 < 2 and Sess.posts < 6, Sess.posts
    print("OK лимит: поиск боя сразу сообщает о лимите, без зависания")


def test_pick_by_percentile():
    """Отбор игроков для ротации рейда по процентилю и ilvl% из рейтингов WCL."""
    from wcl_analyzer.collect import inspect_report
    from wcl_analyzer.raid_demo import CODE, DEMO_URL, FakeRaidClient
    from wcl_analyzer.raid_rotation import parse_pick, pick_players
    c = FakeRaidClient()
    kill = next(f for f in c.report(CODE)["fights"] if f.get("kill"))
    wipe = next(f for f in c.report(CODE)["fights"] if not f.get("kill"))
    insp = inspect_report(c, DEMO_URL, kill["id"])
    dps = [p for p in insp["players"] if p["role"] == "DPS"]
    assert all(p["rank"] is not None and p["ilvl"] is not None for p in insp["players"])
    assert parse_pick("rank:50") == ("rank", 50.0) and parse_pick("x:50") is None and parse_pick("rank:100") is None
    keep, out = pick_players(dps, "rank:50", log=lambda m: None)
    assert keep and all(p["rank"] <= 50 for p in keep) and all(p["rank"] > 50 for p in out)
    assert len(keep) + len(out) == len(dps)
    k2, _ = pick_players(dps, "ilvl:75", log=lambda m: None)
    assert all(p["ilvl"] <= 75 for p in k2)
    all_, none_ = pick_players(dps, None)
    assert all_ == dps and not none_
    # у вайпа рейтингов нет — разбираются все
    wd = [p for p in inspect_report(c, DEMO_URL, wipe["id"])["players"] if p["role"] == "DPS"]
    msgs = []
    k3, _ = pick_players(wd, "rank:25", log=msgs.append)
    assert k3 == wd and "нет рейтингов" in msgs[0]
    try:
        pick_players([{**p, "rank": 80.0} for p in dps], "rank:50", log=lambda m: None)
        raise AssertionError("должна быть ошибка: никто не прошёл отбор")
    except LookupError as e:
        assert "порог выше" in str(e)
    # вкладка «Рейдовые кулдауны в пики урона» в разборе игрока
    from wcl_analyzer import web
    job = {"id": "t", "progress": 0.0, "log": []}
    web._run_player(job, {"demo": True}, None, lambda m: None)
    import json
    assert "raid_peaks" not in job["result"], "полученный урон и сейвы — вкладка разбора рейда, не разбора игрока"
    # сравнение талантов с топом
    T = job["result"]["talents"]
    kinds = {r["kind"] for r in T["rows"]}
    assert {"missing", "choice"} <= kinds, kinds
    branches = {r["branch"] for r in T["rows"]}
    assert {"Класс", "Специализация", "Героическая: Вестник льда"} <= branches, branches
    assert T["hero"]["my"] == "Вестник льда" and T["hero_same_n"] >= 3
    # таланты, у которых в событии нет номера узла, находятся по справочнику
    from wcl_analyzer.talents import _picks, demo_tree_data, spec_tree
    nodes = spec_tree(demo_tree_data(), "Mage", "Frost")["nodes"]
    by_entry = {e: nid for nid, nd in nodes.items() for e in nd["entries"]}
    class _L:
        talent_tree = [(0, 1003, 1), (5004, 1004, 1)]
    assert _picks(_L(), by_entry) == {5003: (1003, 1), 5004: (1004, 1)}
    assert T["has_names"] and all(not r["name"].startswith("узел") for r in T["rows"])
    json.dumps(T)
    import openpyxl, io
    wb = openpyxl.load_workbook(io.BytesIO(job["xlsx"]))
    if wb is not None:
        assert "Таланты" in wb.sheetnames
    print(f"OK отбор по процентилю: ≤50% — {len(keep)} из {len(dps)} DPS, вайп — все")


def test_real_talent_data():
    """Если есть интернет (на GitHub — есть): настоящий справочник Raidbots читается правильно."""
    from wcl_analyzer import talents
    talents._MEM.clear()
    data = talents.load_tree_data(log=lambda m: None, save=False)
    talents._MEM.clear()
    if not data:
        print("ПРОПУЩЕН справочник талантов: нет доступа к Raidbots")
        print("::notice title=Справочник талантов::пропущен — нет доступа к Raidbots")
        return
    # каждый спек справочника разбирается без ошибок, а у лекарей находятся их рейдовые кулдауны
    from wcl_analyzer.raid_cds import candidates
    n_specs, found_cd = 0, 0
    for spec_data in data:
        cls, spec = spec_data.get("className"), spec_data.get("specName")
        t = talents.spec_tree(data, cls, spec)
        assert t and len(t["nodes"]) > 30, (cls, spec)
        n_specs += 1
        spells = {e[1] for n in t["nodes"].values() for e in n["entries"].values() if e[1]}
        for sid, name, _cd, core in candidates(cls, spec):
            if sid in spells:
                found_cd += 1
    for cls, spec in (("Mage", "Frost"), ("Priest", "Holy"), ("Evoker", "Preservation")):
        t = talents.spec_tree(data, cls, spec)
        assert any(n["choice"] for n in t["nodes"].values()), "нет узлов выбора"
        assert any(n["part"] == "hero" and n["sub"] for n in t["nodes"].values()), "нет героических узлов"
        assert t["subs"], "нет названий героических веток"
    print(f"::notice title=Рейдовые кулдауны::в справочнике талантов найдено {found_cd} кулдаунов из таблицы сейвов, спеков {n_specs}")
    print(f"OK справочник талантов Raidbots: {len(data)} спеков, героические ветки, узлы выбора")
    print(f"::notice title=Справочник талантов::прочитан, {len(data)} спеков")


def test_same_difficulty():
    """Эталон собирается только из боёв той же сложности: бой другой сложности в рейтинге пропускается."""
    import copy
    from wcl_analyzer.collect import collect_reference
    c = FakeClient()
    my = c.reports["MYREPORT0001"]["fights"][0]
    diff = int(my.get("difficulty") or 5)
    # лучший лог рейтинга подменяем на бой другой сложности
    best = c.ranks[0]["report"]["code"]
    c.reports[best] = copy.deepcopy(c.reports[best])
    for f in c.reports[best]["fights"]:
        f["difficulty"] = 4 if diff != 4 else 5
    msgs = []
    logs, _ = collect_reference(c, int(my["encounterID"]), "Mage", "Frost", diff, top_n=5, log=msgs.append, save=False)
    assert all(lg.report_code != best for lg in logs), "бой другой сложности попал в эталон"
    assert any("другая сложность" in m for m in msgs), msgs
    assert len(logs) == 5 and all(lg.difficulty == diff for lg in logs)
    try:
        collect_reference(c, int(my["encounterID"]), "Mage", "Frost", 0, top_n=5, log=lambda m: None, save=False)
        raise AssertionError("без сложности эталон собираться не должен")
    except LookupError as e:
        assert "сложность" in str(e)
    print("OK сложность: эталон только из боёв той же сложности, бой другой сложности пропущен")


def test_two_references():
    """Героический бой: эталон 1 — героический топ, эталон 2 — эпохальный; у каждого свой Excel."""
    import copy
    from wcl_analyzer import web

    class TwoDiff(FakeClient):
        def __init__(self):
            super().__init__()
            self.heroic = []
            for rk in self.ranks:
                code = rk["report"]["code"]
                hcode = "H" + code[1:]
                rep = copy.deepcopy(self.reports[code])
                rep["code"] = hcode
                for f in rep["fights"]:
                    f["difficulty"] = 4
                self.reports[hcode], self.raws[hcode] = rep, self.raws[code]
                self.heroic.append({**rk, "report": {**rk["report"], "code": hcode}})
            for f in self.reports["MYREPORT0001"]["fights"]:
                f["difficulty"] = 4

        def rankings(self, encounter_id, class_name, spec_name, difficulty, page=1, force=False, max_age_s=None):
            out = super().rankings(encounter_id, class_name, spec_name, difficulty, page, force, max_age_s)
            if difficulty == 4:
                out = {**out, "characterRankings": {"rankings": self.heroic, "hasMorePages": False}}
            return out

    c = TwoDiff()
    old = web.CLIENT_FACTORY
    web.CLIENT_FACTORY = lambda creds: c
    try:
        job = {"id": "two", "progress": 0.0, "log": []}
        web._run_player(job, {"url": "https://www.warcraftlogs.com/reports/MYREPORT0001#fight=1&source=7",
                              "fight": "1", "actor": "7", "ref": "top10"}, ("a", "b"), lambda m: None)
    finally:
        web.CLIENT_FACTORY = old
    R = job["result"]
    assert R["info"]["ref_difficulty"] == "героический" and R["info"]["ref_same_diff"] == R["info"]["ref_n"]
    A = R.get("alt")
    assert A and A["info"]["ref_difficulty"] == "эпохальный" and A["info"]["ref_same_diff"] == 0, "нет эпохального эталона"
    assert A["excel"].endswith("/alt") and job.get("xlsx") and job.get("xlsx_alt")
    crit = {x["criterion"]: x for x in A["reliability"]["criteria"]}
    assert crit["Босс, сложность, спек"]["level"] == "MEDIUM", crit["Босс, сложность, спек"]
    # без галочки — только героический
    web.CLIENT_FACTORY = lambda creds: c
    try:
        job2 = {"id": "one", "progress": 0.0, "log": []}
        web._run_player(job2, {"url": "https://www.warcraftlogs.com/reports/MYREPORT0001#fight=1&source=7",
                               "fight": "1", "actor": "7", "ref": "top10", "mythic": False}, ("a", "b"), lambda m: None)
    finally:
        web.CLIENT_FACTORY = old
    assert "alt" not in job2["result"]
    print("OK два эталона: героический топ и эпохальный топ, у каждого свой Excel")


def test_talent_data_without_ids():
    """Справочник талантов, где у части узлов и талантов нет номера (как в реальном Raidbots), — без ошибки 'id'."""
    from wcl_analyzer import talents
    from wcl_analyzer.raid_cds import roster_cds
    data = [{"className": "Priest", "specName": "Holy",
             "classNodes": [{"name": "без номера", "entries": [{"name": "x"}]}],
             "specNodes": [{"id": 1, "name": "Божественный гимн", "type": "single",
                            "entries": [{"id": 11, "name": "Божественный гимн", "spellId": 64843}, {"name": "без номера"}]},
                           {"id": 2, "name": "Апофеоз", "type": "single", "entries": [{"id": 21, "spellId": 200183}]}],
             "heroNodes": [{"id": 3, "subTreeId": 7, "entries": [{"id": 31, "name": "Герой"}]}],
             "subTreeNodes": [{"id": 4, "entries": [{"traitSubTreeId": 7, "name": "Ветка"}, {"name": "без номера"}]}, {}]},
            "мусор"]
    t = talents.spec_tree(data, "Priest", "Holy")
    assert set(t["nodes"]) == {1, 2, 3} and t["subs"] == {7: "Ветка"}
    raw = {"details": {"healers": [{"id": 5, "name": "Элария", "type": "Priest", "specs": [{"spec": "Holy"}]}, {"name": "без id"}]},
           "combatant": [{"sourceID": 5, "talentTree": [{"id": 11, "nodeID": 1, "rank": 1}, {"nodeID": 99}]}],
           "report": {"masterData": {"abilities": [{"gameID": 64843, "name": "Божественный гимн"}, {"name": "без номера"}]}}}
    R = {"info": {"duration_s": 300}, "extras": {"raid_cds": []}}
    roster = roster_cds(raw, R, data)
    names = {c["name"] for c in roster}
    assert "Божественный гимн" in names and "Апофеоз" not in names, names  # гимн взят, апофеоз — нет
    # способность не из дерева талантов (базовая) с данными о талантах не выкидывается
    raw["details"]["dps"] = [{"id": 6, "name": "Торвин", "type": "Warrior", "specs": [{"spec": "Fury"}]}]
    raw["combatant"].append({"sourceID": 6, "talentTree": [{"id": 11, "nodeID": 1}]})
    data.append({"className": "Warrior", "specName": "Fury", "classNodes": [], "heroNodes": [], "subTreeNodes": [],
                 "specNodes": [{"id": 1, "entries": [{"id": 11, "spellId": 1}]}]})
    assert any(c["name"] == "Ободряющий клич" for c in roster_cds(raw, R, data)), "базовый кулдаун пропал"
    print("OK справочник талантов без номеров у части узлов: без ошибки 'id', таланты состава учтены")


def test_single_sources():
    """Одна таблица игровых данных и один словарь названий — без копий в разных местах."""
    import json as _json
    from wcl_analyzer import game_data, web
    from wcl_analyzer.names_ru import CLASSES
    d = _json.loads(game_data.BUNDLED.read_text(encoding="utf-8"))
    assert game_data._valid(d)
    seen = set()
    for c in d["raid_cds"]:
        assert c["class"] in CLASSES and c["cd"] > 0 and 1 <= c["power"] <= 3 and c["name"] and c["en"], c
        key = (c["class"], c.get("spec"), c["id"])
        assert key not in seen, key
        seen.add(key)
        if c.get("choice_with"):
            assert any(x["id"] == c["choice_with"] for x in d["raid_cds"]), c
    assert 64843 in game_data.cd_ids() and game_data.cooldown(64843) == 180 and 2825 in game_data.lust_ids()
    assert game_data.cd_name_re().search("Healing Tide Totem") and not game_data.cd_name_re().search("Shadow Word: Pain")
    page = web._page().decode("utf-8")
    assert "/*CLASS_RU*/" not in page and "рыцарь смерти" in page and '"Mage|Frost": "Лёд"' in page
    print("OK один источник: таблица игровых данных и словарь названий")


def test_code_update():
    """Обновление без переустановки: скачанный код подключается при запуске; сломанный — откат на встроенный."""
    import hashlib, io, json as _json, subprocess, tempfile, zipfile
    from pathlib import Path as _P
    from wcl_analyzer import update
    root = _P(__file__).resolve().parents[1]
    tmp = _P(tempfile.mkdtemp())
    shell = "testshell01"
    # архив кода «сборки 77» — как его собирает GitHub
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        for f in (root / "wcl_analyzer").rglob("*"):
            if f.is_file() and "__pycache__" not in f.parts:
                rel = f.relative_to(root).as_posix()
                data = f.read_bytes()
                if rel == "wcl_analyzer/_build.py":
                    data = f'BUILD = "77"\nSHELL_ID = "{shell}"\n'.encode()
                z.writestr(rel, data)
    code = buf.getvalue()
    info = {"build": 77, "shell": shell, "sha256": hashlib.sha256(code).hexdigest()}
    old = (update._get, update.shell_id, update.code_root)
    update._get = lambda url, timeout=30: _json.dumps(info).encode() if url.endswith("update.json") else code
    update.shell_id = lambda: shell
    update.code_root = lambda: tmp / "code"
    try:
        res = update.apply_code()
        assert res["ok"] and res["build"] == 77 and (tmp / "code" / "current.json").exists()
        # испорченный архив не принимается
        update._get = lambda url, timeout=30: (_json.dumps({**info, "sha256": "0" * 64}).encode()
                                               if url.endswith("update.json") else code)
        try:
            update.apply_code()
            raise AssertionError("испорченный архив принят")
        except ValueError as e:
            assert "контрольная сумма" in str(e)
        # GitHub недоступен — номер сборки и код берутся через jsDelivr (по файлам)
        import requests as _rq
        files = {"/" + f.relative_to(root).as_posix(): f.read_bytes() for f in (root / "wcl_analyzer").rglob("*")
                 if f.is_file() and "__pycache__" not in f.parts}
        shell_j = update.shell_fingerprint([(root / f).read_bytes() for f in update.SHELL_FILES])

        def fake_get(url, timeout=30):
            if "github.com" in url:
                raise _rq.ConnectionError("blocked")
            if url.endswith("/resolved?specifier=latest"):
                return b'{"version": "1.1.78"}'
            if "structure=flat" in url:
                return _json.dumps({"files": [{"name": n} for n in files]}).encode()
            path = url.split("@1.1.78", 1)[1]
            return files[path] if path in files else (root / path.lstrip("/")).read_bytes()
        update._get, update.shell_id = fake_get, (lambda: shell_j)
        chk = update.check()
        assert chk["source"] == "jsdelivr" and chk["latest"] == 78 and chk["kind"] == "code", chk
        res = update.apply_code()
        assert res["build"] == 78
        built = (tmp / "code" / "build-78" / "wcl_analyzer" / "_build.py").read_text()
        assert 'BUILD = "78"' in built and shell_j in built
        # всё недоступно — понятная причина, а не «нет связи с Warcraft Logs»
        update._get = lambda url, timeout=30: (_ for _ in ()).throw(_rq.ConnectionError("down"))
        try:
            update.check()
            raise AssertionError("ошибка не выдана")
        except update.UpdateError as e:
            assert "github.com — нет соединения" in str(e) and "jsdelivr.net" in str(e), e
        # вернуть «сборку 77» для проверок загрузчика ниже
        update._get = lambda url, timeout=30: _json.dumps(info).encode() if url.endswith("update.json") else code
        update.shell_id = lambda: shell
        update.apply_code()
    finally:
        update._get, update.shell_id, update.code_root = old
    probe = (f"import sys; sys.path.insert(0, {str(root)!r}); import wcl_boot; wcl_boot.SHELL_ID = {shell!r}; "
             "b = wcl_boot.activate(log=lambda m: print('LOG', m)); import wcl_analyzer, wcl_analyzer._build as v; "
             "print(b, v.BUILD, 'code' in wcl_analyzer.__file__)")
    env = {**__import__("os").environ, "WCL_DATA_DIR": str(tmp)}
    out = subprocess.run([sys.executable, "-c", probe], capture_output=True, text=True, env=env, cwd=str(tmp)).stdout.split()
    assert out[-3:] == ["77", "77", "True"], out
    # другая оболочка — скачанный код не берётся
    out = subprocess.run([sys.executable, "-c", probe.replace(repr(shell), "'other'", 1)], capture_output=True, text=True,
                         env=env, cwd=str(tmp)).stdout.split()
    assert out[-3:] == ["None", "dev", "False"], out
    # сломанный код — откат на встроенный, обновление помечено сломанным
    cur = _json.loads((tmp / "code" / "current.json").read_text())
    (_P(cur["path"]) / "wcl_analyzer" / "web.py").write_text("def broken(:\n")
    out = subprocess.run([sys.executable, "-c", probe], capture_output=True, text=True, env=env, cwd=str(tmp)).stdout
    assert out.split()[-3:] == ["None", "dev", "False"] and "не запустилось" in out, out
    assert (tmp / "code" / "failed.json").exists()
    print("OK обновление кода: подключается при запуске, чужая оболочка и сломанный код — встроенная версия")


def test_android_java_threads():
    """Android: классы приложения из потока сервера не находятся — загружаем их заранее в главном потоке."""
    import threading
    import types
    calls = []

    class _Cls:
        def __init__(self, name):
            self.name = name
            self.mActivity = self
            self.SDK_INT = 33
            self.ACTION_VIEW, self.FLAG_ACTIVITY_NEW_TASK = "view", 1

        def __call__(self, *a):
            return self

        def parse(self, u):
            return u

        def addFlags(self, f):
            pass

        def startActivity(self, intent):
            calls.append("open")

    def autoclass(name):
        if name.startswith("org.kivy") and threading.current_thread() is not threading.main_thread():
            raise Exception("JVM exception occurred: java.lang.ClassNotFoundException")
        return _Cls(name)

    sys.modules["jnius"] = types.SimpleNamespace(autoclass=autoclass)
    try:
        from wcl_analyzer import platform_support as ps
        ps._J.clear()
        errs = []
        # без загрузки заранее — в потоке ошибка (как была у пользователя)
        def no_preload():
            try:
                ps.android_open_url("https://x")
            except Exception as e:  # noqa: BLE001
                errs.append(str(e))
        t = threading.Thread(target=no_preload); t.start(); t.join()
        assert errs and "PythonActivity" in errs[0], errs
        ps._J.clear()
        ps.android_preload()  # главный поток — как в serve()
        t = threading.Thread(target=lambda: ps.android_open_url("https://x")); t.start(); t.join()
        assert calls == ["open"], calls
    finally:
        del sys.modules["jnius"]
        ps._J.clear()
    print("OK Android: классы Java загружаются в главном потоке — сохранение файлов и ссылки работают из потоков")


def test_talents_fallback():
    """Если в событиях боя талантов нет — они берутся из playerDetails(includeCombatantInfo)."""
    from wcl_analyzer.collect import ensure_talents
    from wcl_analyzer.logs import talents_from_details

    class _PL:
        talent_tree, report_code, fight_id, actor_id = [], "R", 1, 7

    class _C:
        def player_details(self, code, fid, combatant=False):
            assert combatant
            return {"dps": [{"id": 7, "name": "Me", "combatantInfo": {"talentTree": [
                {"id": 1003, "nodeID": 5003, "rank": 1}, {"id": 1004, "nodeID": 5004, "rank": 2}]}}]}
    pl = ensure_talents(_C(), _PL())
    assert pl.talent_tree == [(5003, 1003, 1), (5004, 1004, 2)], pl.talent_tree
    assert talents_from_details({"dps": [{"id": 8}]}, 7) == []
    print("OK таланты: запасной источник — сведения об игроках боя")


if __name__ == "__main__":
    test_batched_fetch()
    test_talents_fallback()
    test_android_java_threads()
    test_code_update()
    test_single_sources()
    test_talent_data_without_ids()
    test_two_references()
    test_same_difficulty()
    test_real_talent_data()
    test_pick_by_percentile()
    test_limit_no_hang()
