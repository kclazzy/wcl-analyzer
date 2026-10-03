"""Excel-отчёт: сводка, находки, план, разделы сравнения и эталон с графиками."""
from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference, ScatterChart, Series
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .compare import CompareResult, _fmt_t
from .names_ru import cls_ru, spec_ru
from .metrics import CATEGORY_RU
from .reference import Reference as Ref

FONT = "Arial"
C_HEADER = "1F3864"
C_ACCENT = "2E75B6"
LEVEL_FILL = {"HIGH": "C6EFCE", "MEDIUM": "FFEB9C", "LOW": "FFC7CE", "Unknown": "E7E6E6", "—": "E7E6E6",
              # надёжность (женский род) и влияние на DPS (средний род) — по-русски в ячейках
              "высокая": "C6EFCE", "средняя": "FFEB9C", "низкая": "FFC7CE",
              "высокое": "FFC7CE", "среднее": "FFEB9C", "низкое": "E7E6E6", "не оценено": "E7E6E6"}
REL_RU = {"HIGH": "высокая", "MEDIUM": "средняя", "LOW": "низкая"}
IMPACT_RU = {"HIGH": "высокое", "MEDIUM": "среднее", "LOW": "низкое", "Unknown": "не оценено"}
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

F_INT, F_1, F_2, F_PCT = "#,##0", "#,##0.0", "0.00", "0.0%"
F_SIGNED = "+#,##0;-#,##0;0"
F_SIGNED1 = "+#,##0.0;-#,##0.0;0.0"
F_SIGNED_PCT = "+0.0%;-0.0%;0.0%"
# Секунды -> m:ss.0 (для колонок «время»)
MMSS = 'IF({c}="","—",INT({c}/60)&":"&TEXT(MOD({c},60),"00.0"))'


class Sheet:
    def __init__(self, wb: Workbook, title: str, demo: bool, widths: dict | None = None):
        self.ws = wb.create_sheet(title)
        self.ws.sheet_view.showGridLines = False
        self.row = 1
        for col, w in (widths or {}).items():
            self.ws.column_dimensions[col].width = w
        if demo:
            c = self.ws.cell(row=1, column=1,
                             value="ДЕМО: синтетические данные для проверки программы, не реальный лог")
            c.font = Font(name=FONT, bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor="C00000")
            self.row = 2

    def cell(self, r, c, v=None, fmt=None, bold=False, color=None, fill=None, wrap=False, size=10,
             italic=False):
        cell = self.ws.cell(row=r, column=c, value=v)
        cell.font = Font(name=FONT, bold=bold, color=color, size=size, italic=italic)
        if fmt:
            cell.number_format = fmt
        if fill:
            cell.fill = PatternFill("solid", fgColor=fill)
        cell.alignment = Alignment(wrap_text=wrap, vertical="top")
        return cell

    def title(self, text: str, sub: str | None = None):
        self.cell(self.row, 1, text, bold=True, size=14, color=C_HEADER)
        self.row += 1
        if sub:
            self.cell(self.row, 1, sub, color="595959", italic=True)
            self.row += 1
        self.row += 1

    def section(self, text: str, sub: str | None = None):
        self.cell(self.row, 1, text, bold=True, size=12, color=C_ACCENT)
        self.row += 1
        if sub:
            self.cell(self.row, 1, sub, color="595959", italic=True)
            self.row += 1

    def note(self, text: str):
        self.cell(self.row, 1, text, color="595959", italic=True)
        self.row += 1

    def table(self, headers: list[str], rows: list[list], fmts: list | None = None,
              col: int = 1, links: dict | None = None) -> tuple[int, int]:
        """Пишет таблицу; строка-формула может содержать {r} — номер текущей строки."""
        hr = self.row
        for j, h in enumerate(headers):
            c = self.cell(hr, col + j, h, bold=True, color="FFFFFF", fill=C_HEADER, wrap=True)
            c.border = BORDER
        first = hr + 1
        for i, row in enumerate(rows):
            r = first + i
            for j, v in enumerate(row):
                if isinstance(v, str) and v.startswith("=") and "{r}" in v:
                    v = v.replace("{r}", str(r))
                fmt = fmts[j] if fmts and j < len(fmts) else None
                c = self.cell(r, col + j, v, fmt=fmt, wrap=isinstance(v, str) and len(v) > 40)
                c.border = BORDER
                if links and j in links and links[j](i):
                    c.hyperlink = links[j](i)
                    c.font = Font(name=FONT, color="0563C1", underline="single")
        last = first + len(rows) - 1
        self.row = max(last, hr) + 2
        return first, last

    def level_colors(self, rng: str):
        for lvl, color in LEVEL_FILL.items():
            self.ws.conditional_formatting.add(
                rng, CellIsRule(operator="equal", formula=[f'"{lvl}"'],
                                fill=PatternFill("solid", fgColor=color)))

    def chart(self, chart, anchor: str, w: float = 18, h: float = 8):
        chart.width, chart.height = w, h
        for ax in (getattr(chart, "x_axis", None), getattr(chart, "y_axis", None)):
            if ax is not None:
                ax.delete = False
        self.ws.add_chart(chart, anchor)


def _mmss(col_letter: str) -> str:
    return "=" + MMSS.format(c=f"{col_letter}{{r}}")


def _v(x, nd=None):
    if x is None:
        return None
    return round(float(x), nd) if nd is not None else x


# ============================================================ COMPARE
def write_compare_workbook(r: CompareResult, path: str | Path) -> Path:
    wb = Workbook()
    wb.remove(wb.active)
    demo = r.me.report_code.startswith("DEMO")
    _brief_sheet(wb, r, demo)
    _battle_sheet(wb, r, demo)
    _summary_sheet(wb, r, demo)
    _top5_sheet(wb, r, demo)
    _rotation_sheet(wb, r, demo)
    _cooldown_sheet(wb, r, demo)
    _burst_sheet(wb, r, demo)
    _talents_sheet(wb, r, demo)
    _gear_sheet(wb, r, demo)
    _defensive_sheet(wb, r, demo)
    _uptime_sheet(wb, r, demo)
    _resource_sheet(wb, r, demo)
    _gcd_sheet(wb, r, demo)
    _timeline_sheet(wb, r, demo)
    _mechanics_sheet(wb, r, demo)
    _guide_sheet(wb, r.ref, demo)
    _reference_sheet(wb, r.ref, demo)
    _players_sheet(wb, r.ref, demo)
    _raw_sheet(wb, r, demo)
    path = Path(path)
    wb.save(path)
    return path


TALENT_KIND = {"missing": "У топа почти у всех, у вас нет", "choice": "Другой выбор в узле",
               "extra": "У вас есть, у топа редко", "rank": "Меньше рангов, чем у топа"}


def _talents_sheet(wb, r: CompareResult, demo: bool) -> None:
    T = getattr(r, "talents", None)
    if not T:
        return
    if T.get("error"):
        s = Sheet(wb, "Таланты", demo, {"A": 100})
        s.title("Таланты против топа", T["error"])
        return
    s = Sheet(wb, "Таланты", demo, {"A": 26, "B": 34, "C": 30, "D": 26, "E": 26, "F": 14, "G": 16, "H": 16})
    s.title("Таланты против топа", f"Все скачанные логи топа ({T['n']}), без отбора по билду. "
            + ("Названия — из справочника Raidbots." if T["has_names"] else "Справочник талантов недоступен: показаны номера узлов."))
    for line in T["summary"]:
        s.cell(s.row, 1, "• " + line)
        s.row += 1
    s.row += 1
    if T.get("hero"):
        s.section("Героическая ветка")
        s.table(["Ветка", "Доля топа", "Медиана DPS топа"],
                [[h["name"] + (" (у вас)" if h["name"] == T["hero"]["my"] else ""), h["share"], _v(h["dps"], 0)] for h in T["hero"]["top"]],
                [None, F_PCT, F_INT])
    s.section("Отличия от большинства топа")
    s.table(["Ветка", "Талант", "Отличие", "У вас", "У топа", "Доля топа", "DPS топа: с ним", "DPS топа: без"],
            [[x.get("branch", ""), x["name"], TALENT_KIND[x["kind"]], x["my"], x["top"], x["share"], _v(x["dps_with"], 0),
              _v(x["dps_without"], 0)] for x in T["rows"]], [None, None, None, None, None, F_PCT, F_INT, F_INT])


def _battle_sheet(wb, r: CompareResult, demo: bool) -> None:
    """«Бой по шагам»: ключевые события, главные отличия, движение. Карта — только в приложении."""
    try:
        from .battle import build_battle
        B = build_battle(r)
    except Exception:  # noqa: BLE001 — лист необязательный
        return
    sev = {"high": "сильное", "medium": "заметное"}
    s = Sheet(wb, "Бой по шагам", demo, {"A": 9, "B": 14, "C": 34, "D": 90, "E": 14})
    s.title("Бой по шагам", "Ключевые события боя и сравнение с медианой топа. Карта движения и повтор — в приложении.")
    for line in B["summary"]:
        s.cell(s.row, 1, "• " + line)
        s.row += 1
    s.row += 1
    s.section("Лента боя")
    s.table(["Время", "Событие", "Что", "Подробно", "Отличие"],
            [[_fmt_t(e["t"]), e["type_ru"], e["title"], " | ".join(f"{ln['icon']} {ln['text']}".strip() for ln in e["lines"]),
              sev.get(e["severity"], "")] for e in B["events"]])
    s.section("Главные отличия от топа")
    s.table(["", "Раздел", "Отличие", "Что сделать", "Значение"],
            [[sev.get(d["level"], ""), d["section"], d["text"], d["do"], d["value"] or d["gain"]] for d in B["differences"]])
    M = B["movement"]
    if M.get("available"):
        s.section("Движение", f"За бой: {M['total']} с (топ {M['top_total']} с). Причина — только подтверждённая данными.")
        why = {"mechanic": "механика", "target": "к цели", "unknown": "неизвестна"}
        s.table(["Начало", "Длительность, с", "Причина", "Ярдов", "Топ в этом окне, с"],
                [[_fmt_t(m["start"]), m["moving"], why[m["reason"]] + (f" «{m['mechanic']}»" if m["mechanic"] else ""),
                  m["dist"], m["top_window"]] for m in M["moves"]], [None, F_1, None, F_1, F_1])


def _gear_sheet(wb, r: CompareResult, demo: bool) -> None:
    G = r.tables.get("gear") or {}
    if not G.get("rows"):
        return
    s = Sheet(wb, "Экипировка", demo, {"A": 16, "B": 30, "C": 10, "D": 12, "E": 26, "F": 18, "G": 20, "H": 10, "I": 12})
    miss = G["missing_enchants"] + [f"{x} (временное усиление)" for x in G["missing_temp"]]
    s.title("Экипировка против топа",
            f"Медиана топа — по {G['ref_n']} игрокам. Слот считается зачаровываемым, если зачарование в нём есть "
            "у большинства топа. " + (f"Нет зачарования: {', '.join(miss)}." if miss else "Все нужные слоты зачарованы."))
    rows = [[x["slot_name"], x["name"] or (f"Предмет {x['id']}" if x["id"] else "—"), x["ilvl"], _v(x["ref_ilvl"], 0),
             ("есть" + (f" — {x['enchant_name']}" if x["enchant_name"] else "")) if x["enchant"]
             else ("НЕТ" if x["enchantable"] else "не нужно"),
             "" if x["temp"] is None else ("есть" if x["temp"] else "нет"),
             x["ref_enchant_share"], x["gems"], _v(x["ref_gems"], 0)] for x in G["rows"]]
    first, last = s.table(["Слот", "Предмет", "Ур.", "Ур. топа", "Зачарование", "Временное усиление",
                           "Зачарование у топа", "Камни", "Камни у топа"],
                          rows, [None, None, F_INT, F_INT, None, None, F_PCT, F_INT, F_INT])
    if rows:
        s.ws.conditional_formatting.add(f"E{first}:E{last}", FormulaRule(
            formula=[f'E{first}="НЕТ"'], fill=PatternFill("solid", fgColor="FFC7CE")))


def write_reference_workbook(ref: Ref, path: str | Path) -> Path:
    wb = Workbook()
    wb.remove(wb.active)
    demo = ref.logs[0].report_code.startswith("DEMO")
    _guide_sheet(wb, ref, demo)
    _reference_sheet(wb, ref, demo)
    _players_sheet(wb, ref, demo)
    path = Path(path)
    wb.save(path)
    return path


# ------------------------------------------------------------ Выжимка
IMPACT_FILL = {"HIGH": "FFC7CE", "MEDIUM": "FFEB9C", "LOW": "E7E6E6", "Unknown": "E7E6E6"}


def _brief_sheet(wb, r: CompareResult, demo, title: str = "Выжимка"):
    """Одна страница: вердикт, три действия, что уже хорошо, что не зависит от игрока."""
    s = Sheet(wb, title, demo, {"A": 5, "B": 62, "C": 22, "D": 44})
    b, me = r.brief, r.me
    s.title(f"{me.name}, {spec_ru(me.cls, me.spec)} — {me.encounter_name} ({me.difficulty_name})", b["verdict"])
    s.cell(s.row, 1, f"Надёжность сравнения: {REL_RU.get(b['reliability'], b['reliability'])} — {b['reliability_note']}"
           + (f". {r.ref.build_note}" if r.ref.build_note else ""), color="595959")
    s.row += 2
    s.section("Что сделать в следующем бою")
    if b["actions"]:
        first, last = s.table(["№", "Что делать", "Ожидаемый эффект", "Почему"],
                              [[i + 1, a["do"], a["gain"], (a["title"] + (" — " + a["why"] if a["why"] else ""))]
                               for i, a in enumerate(b["actions"])])
        for i, a in enumerate(b["actions"]):
            s.ws.cell(row=first + i, column=3).fill = PatternFill("solid", fgColor=IMPACT_FILL.get(a["level"], "E7E6E6"))
            for c in (2, 4):
                s.ws.cell(row=first + i, column=c).alignment = Alignment(wrap_text=True, vertical="top")
    else:
        s.note("Существенных отличий от топа нет: отклонения в пределах разброса самих лучших игроков.")
        s.row += 1
    if b["good"]:
        s.section("Уже на уровне топа")
        for g in b["good"]:
            s.cell(s.row, 2, "• " + g)
            s.row += 1
        s.row += 1
    if b["context"]:
        s.section("Не зависит от вас")
        for g in b["context"]:
            s.cell(s.row, 2, "• " + g, wrap=True)
            s.row += 1
        s.row += 1
    if b.get("noise_hidden"):
        s.note(f"Скрыто отличий, которые так же часто бывают у самих игроков топа: {b['noise_hidden']} "
               "(лист «Сводка» → все отличия).")
    s.note("Подробности — на следующих листах. Оценка эффекта — доля вашего DPS; пересечения факторов исключены.")


# ------------------------------------------------------------ Сводка
def _summary_sheet(wb, r: CompareResult, demo):
    s = Sheet(wb, "Сводка", demo, {"A": 34, "B": 16, "C": 16, "D": 14, "E": 14, "F": 50})
    me, ref = r.me, r.ref
    s.title(f"Анализ лога: {me.name} — {spec_ru(me.cls, me.spec)}",
            f"Сравнение с эталоном {ref.label} ({ref.n} логов). Эталон — медиана игроков, а не один лог.")
    info = [
        ("Персонаж", me.name), ("Класс", cls_ru(me.cls)), ("Специализация", spec_ru(me.cls, me.spec)),
        ("Босс", me.encounter_name), ("Сложность", me.difficulty_name),
        ("Длительность боя", _fmt_t(me.duration)), ("Исход", "Килл" if me.kill else "Вайп"),
        ("Отчёт WCL", me.url), ("Эталон", f"{ref.label}, {ref.n} логов"),
        ("Надёжность сравнения", REL_RU.get(r.reliability["overall"], r.reliability["overall"])),
    ]
    for k, v in info:
        s.cell(s.row, 1, k, bold=True)
        c = s.cell(s.row, 2, v)
        if k == "Отчёт WCL":
            c.hyperlink = v
            c.font = Font(name=FONT, color="0563C1", underline="single")
        if k == "Надёжность сравнения":
            c.fill = PatternFill("solid", fgColor=LEVEL_FILL.get(v, "E7E6E6"))
            c.font = Font(name=FONT, bold=True)
        s.row += 1
    s.row += 1

    s.section("Мой лог против эталона", "Разница считается формулами: поменяйте значение — пересчитается.")
    rows = []
    for x in r.summary:
        fmt = {"num": F_INT, "num1": F_1, "pct": F_PCT}[x["fmt"]]
        rows.append([x["metric"], _v(x["my"], 4), _v(x["ref"], 4), "=B{r}-C{r}",
                     '=IF(C{r}=0,"",B{r}/C{r}-1)'])
    first, last = s.table(["Показатель", "Мой лог", f"{ref.label} (медиана)", "Разница", "Разница, %"], rows)
    for i, x in enumerate(r.summary):
        fmt = {"num": F_INT, "num1": F_1, "pct": F_PCT}[x["fmt"]]
        rr = first + i
        s.ws.cell(rr, 2).number_format = fmt
        s.ws.cell(rr, 3).number_format = fmt
        s.ws.cell(rr, 4).number_format = {"num": F_SIGNED, "num1": F_SIGNED1, "pct": F_SIGNED_PCT}[x["fmt"]]
        s.ws.cell(rr, 5).number_format = F_SIGNED_PCT
    s.ws.conditional_formatting.add(f"E{first}:E{first}", CellIsRule(operator="lessThan", formula=["0"],
                                    font=Font(name=FONT, color="C00000", bold=True)))
    summary_end = s.row
    r.tables["_summary_dps_row"] = first + [x["metric"] for x in r.summary].index("DPS")

    # Графики: DPS и доли
    dps = BarChart()
    dps.type = "col"
    dps.title = "DPS: мой лог и эталон"
    dps.add_data(Reference(s.ws, min_col=2, max_col=3, min_row=first - 1, max_row=first), titles_from_data=True)
    dps.set_categories(Reference(s.ws, min_col=1, min_row=first, max_row=first))
    dps.y_axis.numFmt = "#,##0"
    dps.y_axis.title = "DPS"
    dps.dataLabels = DataLabelList()
    dps.dataLabels.showVal = True
    dps.dataLabels.showSerName = False
    dps.dataLabels.showCatName = False
    dps.dataLabels.showLegendKey = False
    s.chart(dps, "H4", w=12, h=7.5)

    pct_rows = [first + i for i, x in enumerate(r.summary) if x["fmt"] == "pct"]
    if pct_rows:
        pc = BarChart()
        pc.type = "bar"
        pc.title = "Эффективность: мой лог и эталон"
        lo, hi = min(pct_rows), max(pct_rows)
        for col in (2, 3):
            ser = Series(Reference(s.ws, min_col=col, min_row=lo, max_row=hi),
                         title=s.ws.cell(first - 1, col).value)
            pc.series.append(ser)
        pc.set_categories(Reference(s.ws, min_col=1, min_row=lo, max_row=hi))
        pc.y_axis.numFmt = "0%"
        pc.y_axis.scaling.min = 0
        pc.y_axis.scaling.max = 1
        s.chart(pc, "H20", w=16, h=8)

    s.row = summary_end
    s.section("Надёжность сравнения", "Итог = минимальный уровень среди проверенных критериев.")
    f, l = s.table(["Критерий", "Уровень", "Основание"],
                   [[c["criterion"], REL_RU.get(c["level"], c["level"]), c["reason"]] for c in r.reliability["criteria"]])
    s.ws.column_dimensions["C"].width = 16
    for rr in range(f, l + 1):
        s.ws.merge_cells(start_row=rr, start_column=3, end_row=rr, end_column=6)
    s.level_colors(f"B{f}:B{l}")
    s.note("Высокий DPS эталона не доказывает, что каждое отличие — ошибка: смотрите контекст в листах разделов.")


# ------------------------------------------------------ Top 5 и план
def _top5_sheet(wb, r: CompareResult, demo):
    s = Sheet(wb, "Топ-5 и план", demo, {"A": 7, "B": 14, "C": 58, "D": 12, "E": 12, "F": 16,
                                         "G": 11, "H": 11, "I": 14, "J": 60})
    s.title("Что исправить в первую очередь",
            "Сначала отличия с измеримым влиянием на DPS, затем неоценённые — по величине отклонения; "
            "не больше двух пунктов из одного раздела.")
    rows = [_finding_row(i, f) for i, f in enumerate(r.top5)]
    first, last = s.table(FINDING_HEADERS, rows, FINDING_FMTS)
    if rows:
        s.level_colors(f"H{first}:H{last}")

    s.section("План тренировок", "Одна сессия — один фокус.")
    s.table(["Сессия", "Раздел", "Фокус: что тренировать"],
            [[p["session"], p["section"], p["focus"]] for p in r.plan])
    s.note("Проверка: после каждой сессии вставьте ссылку на новый лог — показатель должен приблизиться к эталону.")
    s.row += 1

    s.section("Разрыв в DPS: из чего он складывается",
              "Оценки факторов пересекаются (простой уменьшает и число кастов), поэтому их сумма "
              "может не совпасть с общим разрывом. «Возможная причина» — вклад не доказан.")
    gs = s.row
    s.cell(gs, 3, "Общий разрыв DPS (эталон − мой)", bold=True)
    s.cell(gs, 4, f"=-'Сводка'!D{r.tables['_summary_dps_row']}", fmt=F_INT, bold=True)
    s.row += 2
    factors = [g for g in r.gap if g["factor"] != "Не объяснено моделью"]
    rows = [[g["factor"], _v(g["dps"], 1), f'=IF($D${gs}=0,"",D{{r}}/$D${gs})', g["status"]] for g in factors]
    rows.append(["Не объяснено моделью", None, f'=IF($D${gs}=0,"",D{{r}}/$D${gs})',
                 "Остаток: другие факторы и пересечения оценок"])
    first, ur = s.table(["Фактор", "DPS", "Доля разрыва", "Статус"], rows, [None, F_INT, F_PCT, None], col=3)
    s.ws.cell(ur, 4).value = f"=D{gs}-SUM(D{first}:D{ur - 1})" if ur > first else f"=D{gs}"
    s.ws.cell(ur, 3).font = Font(name=FONT, bold=True)
    ch = BarChart()
    ch.type = "bar"
    ch.title = "Оценка вклада факторов в разрыв DPS"
    ch.add_data(Reference(s.ws, min_col=4, min_row=first, max_row=ur), titles_from_data=False)
    ch.set_categories(Reference(s.ws, min_col=3, min_row=first, max_row=ur))
    ch.legend = None
    ch.y_axis.numFmt = "#,##0"
    s.chart(ch, f"H{gs}", w=16, h=7)
    s.row = max(s.row, gs + 16)

    s.section("Все найденные отличия", "Каждое — факт расхождения с эталоном, а не вердикт.")
    rows = [_finding_row(i, f) for i, f in
            enumerate(sorted(r.findings, key=lambda f: (-(f.impact or 0), -f.severity)))]
    first, last = s.table(FINDING_HEADERS, rows, FINDING_FMTS)
    if rows:
        s.level_colors(f"H{first}:H{last}")


FINDING_HEADERS = ["№", "Раздел", "Отличие от эталона", "Моё", "Эталон", "Ед.", "Разница", "Влияние",
                   "Оценка, % DPS", "Контекст"]
FINDING_FMTS = [None, None, None, F_2, F_2, None, "+#,##0.00;-#,##0.00;0", None, F_PCT, None]


def _finding_row(i, f):
    my, ref, unit = f.my, f.ref, f.unit
    if unit == "доля" or unit == "доля кастов":
        unit = "доля (0–1)"
    return [i + 1, f.section, f.title, _v(my, 3), _v(ref, 3), unit,
            '=IF(OR(D{r}="",E{r}=""),"",D{r}-E{r})', IMPACT_RU[f.impact_level], _v(f.impact, 4), f.note]


# ------------------------------------------------------------ Ротация
def _rotation_sheet(wb, r: CompareResult, demo):
    s = Sheet(wb, "Ротация", demo, {"A": 26, "B": 14, "C": 10, "D": 10, "E": 12, "F": 9, "G": 9,
                                    "H": 11, "I": 11, "J": 10, "K": 12})
    s.title("Касты: мой лог против эталона",
            f"Эталон — медиана {r.ref.label}, пересчитанная на длительность вашего боя.")
    k = r.me.duration / 60
    rows = [[x["name"], x["cat"], x["my"], _v(x["my_cpm"], 2), _v((x["ref_cpm"] or 0) * k, 1),
             _v((r.ref.agg["abilities"][x["id"]]["cpm"]["p25"] or 0) * k, 1),
             _v((r.ref.agg["abilities"][x["id"]]["cpm"]["p75"] or 0) * k, 1),
             _v(x["share"], 3), _v(x["ref_cpm"], 2), "=C{r}-E{r}",
             '=IF(C{r}<F{r},"ниже эталона",IF(C{r}>G{r},"выше эталона","в норме"))']
            for x in r.tables["casts"]]
    first, last = s.table(["Способность", "Категория", "Мои касты", "Мои касты в минуту", "Эталон (медиана)",
                           "Эталон P25", "Эталон P75", "Доля игроков", "Касты в минуту у эталона", "Разница", "Статус"],
                          rows, [None, None, F_INT, F_2, F_1, F_1, F_1, F_PCT, F_2, F_SIGNED1, None])
    if rows:
        s.ws.conditional_formatting.add(f"K{first}:K{last}", CellIsRule(
            operator="equal", formula=['"ниже эталона"'], fill=PatternFill("solid", fgColor="FFC7CE")))
        s.ws.conditional_formatting.add(f"K{first}:K{last}", CellIsRule(
            operator="equal", formula=['"выше эталона"'], fill=PatternFill("solid", fgColor="FFEB9C")))
        ch = BarChart()
        ch.type = "bar"
        ch.title = "Касты за бой: мой лог и медиана эталона"
        ch.add_data(Reference(s.ws, min_col=3, min_row=first - 1, max_row=last), titles_from_data=True)
        ch.add_data(Reference(s.ws, min_col=5, min_row=first - 1, max_row=last), titles_from_data=True)
        ch.set_categories(Reference(s.ws, min_col=1, min_row=first, max_row=last))
        s.chart(ch, "M3", w=17, h=max(7, 0.9 * len(rows) + 3))

    s.section("Последовательности (3 способности подряд)",
              "Отличие от частой последовательности не означает ошибку: проверьте проки, ресурс и механику.")
    my_ng = r.mm["ngrams"][3]
    my_total = sum(my_ng.values()) or 1
    rows = []
    for g in r.ref.agg["ngrams"][3]:
        rows.append([" → ".join(r.ref.name_of(x) for x in g["seq"]), _v(g["share"], 4),
                     _v(g["player_share"], 3), _v(my_ng.get(g["seq"], 0) / my_total, 4), "=D{r}-B{r}"])
    s.table(["Последовательность эталона", "Доля в эталоне", "Игроков эталона (≥2 раз)",
             "Доля у меня", "Разница"], rows, [None, F_PCT, F_PCT, F_PCT, F_SIGNED_PCT])
    s.ws.column_dimensions["A"].width = 44
    mine_top = [(g, c) for g, c in my_ng.most_common(8)]
    ref_seqs = {tuple(g["seq"]) for g in r.ref.agg["ngrams"][3]}
    s.section("Мои частые последовательности")
    s.table(["Последовательность", "Доля у меня", "Есть в топ-10 эталона"],
            [[" → ".join(r.ref.name_of(x) for x in g), _v(c / my_total, 4), "да" if g in ref_seqs else "нет"]
             for g, c in mine_top], [None, F_PCT, None])


# ----------------------------------------------------------- Кулдауны
def _cooldown_sheet(wb, r: CompareResult, demo):
    s = Sheet(wb, "Кулдауны", demo, {"A": 24, "B": 14, "C": 10, "D": 11, "E": 11, "F": 12, "G": 11,
                                     "H": 13, "I": 13, "J": 13, "K": 13})
    s.title("Кулдауны: использование и тайминги",
            "Доступно = сколько раз CD можно нажать за бой при использовании сразу по готовности.")
    rows = [[x["name"], x["cat"], _v(x["cd"], 0), x["ideal"], x["used"], _v(x["ref_used"], 1),
             '=IF(B{r}="Защитный","под механики",D{r}-E{r})',
             _v(x["avg_delay"], 1), _v(x["max_delay"], 1), _v(x["ref_delay"], 1), '=IF(J{r}="","",H{r}-J{r})']
            for x in r.tables["cooldowns"]]
    s.section("Упущено")
    first, last = s.table(["Кулдаун", "Категория", "CD, с (оценка)", "Доступно", "Использовано",
                           "Эталон (медиана)", "Пропущено", "Средняя задержка, с", "Макс. задержка, с",
                           "Задержка эталона, с", "Разница задержки, с"],
                          rows, [None, None, F_INT, F_INT, F_INT, F_1, F_INT, F_1, F_1, F_1, F_SIGNED1])
    s.note("CD оценён по минимальному интервалу между кастами у эталона; точное значение задаётся в spell_meta.json.")
    s.row += 1

    s.section("Тайминги: k-е применение", "Отрицательная разница — раньше эталона, положительная — позже.")
    rows = [[x["name"], x["k"], _v(x["my"], 1), _mmss("C").replace("{r}", "{r}"), _v(x["ref"], 1),
             _mmss("E"), _v(x["p25"], 1), _v(x["p75"], 1), '=IF(C{r}="","",C{r}-E{r})',
             x["players"], _v(x["confidence"], 3)] for x in r.tables["cd_timing"]]
    first, last = s.table(["Действие", "№", "Моё время, с", "Моё (м:сс)", "Эталон, с", "Эталон (м:сс)",
                           "P25, с", "P75, с", "Разница, с", "Игроков", "Доверие"], rows,
                          [None, F_INT, F_1, None, F_1, None, F_1, F_1, F_SIGNED1, F_INT, F_PCT])
    if rows:
        s.ws.conditional_formatting.add(f"I{first}:I{last}", FormulaRule(
            formula=[f'AND(ISNUMBER(I{first}),ABS(I{first})>10)'], fill=PatternFill("solid", fgColor="FFC7CE")))
        # Подписи категорий для графика: «Действие №k»
        for rr in range(first, last + 1):
            s.cell(rr, 13, f'=A{rr}&" №"&B{rr}')
        s.cell(first - 1, 13, "Подпись", bold=True, color="FFFFFF", fill=C_HEADER)
        ch = BarChart()
        ch.type = "bar"
        ch.title = "Отклонение от медианы эталона, с (+ позже, − раньше)"
        ch.add_data(Reference(s.ws, min_col=9, min_row=first, max_row=last), titles_from_data=False)
        ch.set_categories(Reference(s.ws, min_col=13, min_row=first, max_row=last))
        ch.legend = None
        s.chart(ch, "O3", w=16, h=max(7, 0.7 * len(rows) + 3))


# -------------------------------------------------------------- Бурст
def _burst_sheet(wb, r: CompareResult, demo):
    s = Sheet(wb, "Бурст", demo, {"A": 10, "B": 26, "C": 16, "D": 16, "E": 14, "F": 14, "G": 14, "H": 14, "I": 14})
    main = r.ref.name_of(r.ref.main_cd) if r.ref.main_cd else "—"
    s.title("Окна бурста: мой лог против эталона",
            f"Окно строится от главного атакующего кулдауна ({main}); смещения — секунды относительно его нажатия.")
    rows = [[x["window"], x["name"], _v(x["ref_t0"], 1), _v(x["my_t0"], 1), _v(x["ref_offset"], 1),
             _v(x["my_offset"], 1), '=IF(F{r}="","нет в окне",F{r}-E{r})', _v(x["share"], 3), x.get("kind", "")]
            for x in r.tables["burst"]]
    first, last = s.table(["Окно", "Действие", "Старт окна у эталона, с", "Старт окна у меня, с",
                           "Смещение у эталона, с", "Моё смещение, с", "Разница, с", "Доля игроков", "Тип"],
                          rows, [F_INT, None, F_1, F_1, F_SIGNED1, F_SIGNED1, F_SIGNED1, F_PCT, None])
    if rows:
        s.ws.conditional_formatting.add(f"G{first}:G{last}", FormulaRule(
            formula=[f'OR(G{first}="нет в окне",AND(ISNUMBER(G{first}),ABS(G{first})>3))'],
            fill=PatternFill("solid", fgColor="FFC7CE")))
    s.note("Красным — действие не попало в окно ±3 с от позиции эталона.")
    s.row += 1
    s.section("Паттерн эталона словами")
    for w in r.ref.agg["burst"]:
        comps = [f"{r.ref.name_of(c['id'])} ({c['offset']:+.0f} с, {c['share']:.0%})" for c in w["components"]]
        lust = f"; Жажда крови в окне у {w['lust_share']:.0%}" if w["lust_share"] else ""
        s.cell(s.row, 1, f"Окно {w['window']}: {_fmt_t(w['t0']['median'])} — {main}"
               + (" + " + ", ".join(comps) if comps else "") + lust)
        s.row += 1
    s.row += 1
    s.section("Мои окна")
    for i, w in enumerate(r.mm["burst"]):
        comps = [f"{r.ref.name_of(ab)} ({off:+.0f} с)" for ab, off in w["components"]]
        s.cell(s.row, 1, f"Окно {i + 1}: {_fmt_t(w['t0'])} — {main}" + (" + " + ", ".join(comps) if comps else ""))
        s.row += 1


# ---------------------------------------------------------- Дефенсивы
def _defensive_sheet(wb, r: CompareResult, demo):
    s = Sheet(wb, "Защита", demo, {"A": 22, "B": 22, "C": 6, "D": 12, "E": 12, "F": 16, "G": 16,
                                      "H": 12, "I": 12, "J": 14})
    s.title("Сейвы и механики",
            "Упреждение — за сколько секунд до удара нажата защита (отрицательное — после удара).")
    rows = [[x["name"], x["mechanic"], x["k"], _v(x["t_hit"], 1), _v(x["my_def_t"], 1), _v(x["my_lead"], 1),
             _v(x["ref_lead"], 1), '=IF(F{r}="","нет защиты",F{r}-G{r})', _v(x["hp_after"], 0), _v(x["hit"], 0)]
            for x in r.tables["defensives"]]
    first, last = s.table(["Защитная способность", "Механика", "№", "Удар, с", "Моя защита, с", "Моё упреждение, с",
                           "Упреждение эталона, с", "Разница, с", "Здоровье после удара, %", "Урон удара"],
                          rows, [None, None, F_INT, F_1, F_1, F_SIGNED1, F_1, F_SIGNED1, F_INT, F_INT])
    if rows:
        s.ws.conditional_formatting.add(f"H{first}:H{last}", FormulaRule(
            formula=[f'OR(H{first}="нет защиты",AND(ISNUMBER(F{first}),F{first}<0))'],
            fill=PatternFill("solid", fgColor="FFC7CE")))
    s.section("Паттерн эталона")
    rows = []
    for ab, d in r.ref.agg["defensives"].items():
        rows.append([r.ref.name_of(ab), _v(d["uses"]["median"], 1), d["users"],
                     r.ref.name_of(d["mechanic"]) if d["mechanic"] else "—", _v(d["offset"], 1),
                     _v(d["link_share"], 3), _v(d["hp"], 0), _v(d["confidence"], 3)])
    s.table(["Защитная способность", "Применений (медиана)", "Игроков", "Чаще всего перед", "Упреждение, с",
             "Доля привязанных", "Здоровье при нажатии, %", "Доверие"], rows,
            [None, F_1, F_INT, None, F_1, F_PCT, F_INT, F_PCT])
    s.section("Смерти")
    if r.tables["deaths"]:
        s.table(["Время", "Урон за 5 с до смерти", "Защита за 15 с", "Эталон нажимал защиту"],
                [[_fmt_t(d["t"]), ", ".join(f"{n}: {v:,.0f}" for n, v in d["sources"]),
                  ", ".join(d["defensives"]) or "нет", _v(d["ref_def_share"], 3)] for d in r.tables["deaths"]],
                [None, None, None, F_PCT])
    else:
        s.note("Смертей в этом бою нет.")


# ------------------------------------------------------------- Uptime
def _uptime_sheet(wb, r: CompareResult, demo):
    s = Sheet(wb, "Эффекты", demo, {"A": 26, "B": 10, "C": 10, "D": 11, "E": 11, "F": 11, "G": 12})
    s.title("Время действия баффов и дебаффов", "Дебаффы — на основной цели (боссе).")
    rows = [[x["name"], x["kind"], _v(x["my"], 4), _v(x["ref"], 4), _v(x["p25"], 4), _v(x["p75"], 4), "=C{r}-D{r}"]
            for x in r.tables["uptime"]]
    first, last = s.table(["Эффект", "Тип", "У меня", "Эталон", "P25", "P75", "Разница"], rows,
                          [None, None, F_PCT, F_PCT, F_PCT, F_PCT, F_SIGNED_PCT])
    if rows:
        s.ws.conditional_formatting.add(f"C{first}:C{last}", FormulaRule(
            formula=[f"C{first}<E{first}-0.02"], fill=PatternFill("solid", fgColor="FFC7CE")))
        ch = BarChart()
        ch.type = "bar"
        ch.title = "Время действия: мой лог и эталон"
        ch.add_data(Reference(s.ws, min_col=3, max_col=4, min_row=first - 1, max_row=last), titles_from_data=True)
        ch.set_categories(Reference(s.ws, min_col=1, min_row=first, max_row=last))
        ch.y_axis.numFmt = "0%"
        ch.y_axis.scaling.min = 0
        ch.y_axis.scaling.max = 1
        s.chart(ch, "I3", w=16, h=max(7, 0.8 * len(rows) + 3))
    bd = r.tables.get("boss_debuffs") or []
    if bd:
        s.section("Дебаффы на боссе", "Свои, рейдовые и окна уязвимости. «Бурст внутри» — доля времени главного CD "
                                      "под этим дебаффом.")
        s.table(["Дебафф", "Тип", "У меня", "Эталон", "Бурст внутри: я", "Бурст внутри: топ"],
                [[x["name"], x["kind"], _v(x["my"], 4), _v(x["ref"], 4), _v(x["my_overlap"], 4), _v(x["ref_overlap"], 4)]
                 for x in bd], [None, None, F_PCT, F_PCT, F_PCT, F_PCT])
    ext = r.tables.get("externals") or []
    if ext:
        s.section("Внешние баффы", "От других игроков рейда; на них вы не влияете.")
        s.table(["Бафф", "У меня", "Эталон (медиана)", "Доля топа"],
                [[x["name"], x["my"], _v(x["ref"], 1), _v(x["share"], 3)] for x in ext], [None, F_INT, F_1, F_PCT])


# --------------------------------------------------- Ресурсы и проки
def _resource_sheet(wb, r: CompareResult, demo):
    s = Sheet(wb, "Ресурсы и проки", demo, {"A": 32, "B": 12, "C": 12, "D": 10, "E": 10, "F": 12,
                                            "G": 11, "H": 11, "I": 11, "J": 12, "K": 12, "L": 12, "M": 12})
    name = r.mm["resource"]["name"] if r.mm["resource"] else "Ресурс"
    s.title(f"Управление ресурсом: {name}",
            "Переполнение: ресурс у предела и после следующего каста всё ещё у предела — прирост потерян.")
    if r.tables["resource"]:
        rows = [[x["metric"], _v(x["my"], 4), _v(x["ref"], 4), _v(x["p25"], 4), _v(x["p75"], 4), "=B{r}-C{r}"]
                for x in r.tables["resource"]]
        s.table(["Показатель", "Мой лог", "Эталон", "P25", "P75", "Разница"], rows,
                [None, F_PCT, F_PCT, F_PCT, F_PCT, F_SIGNED_PCT])
    else:
        s.note("В событиях кастов нет данных ресурса.")
    s.section("Проки", "Потеряно = истёк без использования + перезаписан новым проком.")
    rows = [[x["name"], x["consumer"], x["gained"], x["used"], x["expired"], x["refreshed"], "=E{r}+F{r}",
             _v(x["ref_gained"], 1), _v(x["ref_used"], 1), _v(x["ref_expired"], 1), _v(x["ref_refreshed"], 1),
             "=J{r}+K{r}", _v(x["reaction"], 2), _v(x["ref_reaction"], 2)] for x in r.tables["procs"]]
    first, last = s.table(["Прок", "Тратится через", "Получено", "Использовано", "Истекло", "Перезаписано",
                           "Потеряно", "Эталон: получено", "Эталон: использовано", "Эталон: истекло",
                           "Эталон: перезаписано", "Эталон: потеряно", "Моя реакция, с", "Реакция эталона, с"],
                          rows, [None, None] + [F_INT] * 5 + [F_1] * 5 + [F_2, F_2])
    if rows:
        ch = BarChart()
        ch.type = "col"
        ch.title = "Потерянные проки: мой лог и эталон"
        ch.add_data(Reference(s.ws, min_col=7, min_row=first - 1, max_row=last), titles_from_data=True)
        ch.add_data(Reference(s.ws, min_col=12, min_row=first - 1, max_row=last), titles_from_data=True)
        ch.set_categories(Reference(s.ws, min_col=1, min_row=first, max_row=last))
        s.chart(ch, f"A{s.row}", w=14, h=7)


# --------------------------------------------------------------- GCD
def _gcd_sheet(wb, r: CompareResult, demo):
    s = Sheet(wb, "Простой", demo, {"A": 34, "B": 14, "C": 14, "D": 14, "E": 20, "F": 20, "G": 18})
    g, gr = r.mm["gcd"], r.ref.agg["gcd"]
    s.title("Простой между действиями", "Простой — время от конца глобального кулдауна или каста до начала следующего действия.")
    rows = [
        ["Глобальный кулдаун, с", _v(g["gcd"], 2), _v(gr["gcd"]["median"], 2), "=B{r}-C{r}"],
        ["Средний простой на действие, с", _v(g["idle_per_action"], 3), _v(gr["idle_per_action"]["median"], 3), "=B{r}-C{r}"],
        ["Доля времени в простое", _v(g["idle_share"], 4), _v(gr["idle_share"]["median"], 4), "=B{r}-C{r}"],
        ["Простой за бой, с", _v(g["idle_total"], 1), _v((gr["idle_share"]["median"] or 0) * r.me.duration, 1), "=B{r}-C{r}"],
    ]
    first, _ = s.table(["Показатель", "Мой лог", "Эталон", "Разница"], rows)
    for i, fmt in enumerate([(F_2, "+0.00;-0.00;0.00"), ("0.000", "+0.000;-0.000;0"),
                             (F_PCT, F_SIGNED_PCT), (F_1, F_SIGNED1)]):
        for c in (2, 3):
            s.ws.cell(first + i, c).number_format = fmt[0]
        s.ws.cell(first + i, 4).number_format = fmt[1]
    s.note("Движение и фазы без цели тоже дают простой: сравнивайте с эталоном на том же отрезке боя.")
    s.row += 1

    s.section("Простой по отрезкам боя (30 с)")
    n = max(len(g["idle_by_30s"]), len(gr["idle_by_30s"]))
    rows = [[f"{_fmt_t(i * 30)}–{_fmt_t(i * 30 + 30)}",
             _v(g["idle_by_30s"][i] if i < len(g["idle_by_30s"]) else None, 2),
             _v(gr["idle_by_30s"][i] if i < len(gr["idle_by_30s"]) else None, 2),
             '=IF(OR(B{r}="",C{r}=""),"",B{r}-C{r})'] for i in range(n)]
    first, last = s.table(["Отрезок", "Мой простой, с", "Эталон (медиана), с", "Разница, с"], rows,
                          [None, F_2, F_2, F_SIGNED1])
    if rows:
        ch = LineChart()
        ch.title = "Простой по 30-секундным отрезкам"
        ch.add_data(Reference(s.ws, min_col=2, max_col=3, min_row=first - 1, max_row=last), titles_from_data=True)
        ch.set_categories(Reference(s.ws, min_col=1, min_row=first, max_row=last))
        ch.y_axis.title = "секунд простоя"
        for ser in ch.series:
            ser.smooth = False
        s.chart(ch, "F3", w=17, h=8)

    s.section("Самые длинные паузы", "Сравните с простоем эталона в том же 30-секундном отрезке.")
    s.table(["С (м:сс)", "До (м:сс)", "Пауза, с", "После действия", "Перед действием",
             "Простой эталона в отрезке, с"],
            [[_fmt_t(x["from"]), _fmt_t(x["to"]), _v(x["idle"], 1), x["prev_name"], x["next_name"],
              _v(x["ref_bucket_idle"], 2)] for x in r.tables["gcd_worst"]],
            [None, None, F_1, None, None, F_2])


# ----------------------------------------------------------- Таймлайн
LANE_Y = {"Медиана топа": 3, "Мой лог": 2, "Механики": 1}


def _timeline_sheet(wb, r: CompareResult, demo):
    s = Sheet(wb, "Таймлайн", demo, {"A": 13, "B": 10, "C": 11, "D": 16, "E": 26, "F": 6, "G": 14, "H": 8})
    s.title("Таймлайн: медиана эталона против моего лога",
            "Y: 3 — медиана эталона, 2 — мой лог, 1 — механики вашего боя. Красный ромб — отличие > 10 с.")
    ev = sorted(r.timeline, key=lambda e: (e["cat"], e["lane"], e["t"]))
    rows = [[e["lane"], _v(e["t"], 1), _mmss("B"), e["cat"], e["name"], e["k"], _v(e.get("delta"), 1),
             LANE_Y[e["lane"]]] for e in ev]
    first, last = s.table(["Дорожка", "Время, с", "м:сс", "Категория", "Действие", "№", "Отличие, с", "Y"],
                          rows, [None, F_1, None, None, None, F_INT, F_SIGNED1, F_INT])
    table_end = s.row
    if not rows:
        return
    ch = ScatterChart()
    ch.title = "Кулдауны, зелья и механики по времени"
    ch.x_axis.title = "Время боя, с"
    ch.y_axis.title = "3 эталон · 2 мой лог · 1 механики"
    ch.y_axis.scaling.min = 0
    ch.y_axis.scaling.max = 4
    ch.y_axis.majorUnit = 1
    ch.y_axis.crossBetween = "midCat"
    ch.x_axis.crossBetween = "midCat"
    ch.x_axis.scaling.min = 0
    ch.x_axis.scaling.max = max(r.me.duration, r.ref.agg["duration"]["median"] or 0) + 5
    symbols = {"Атакующий кулдаун": ("circle", "2E75B6"), "Защитный": ("square", "70AD47"),
               "Зелье": ("triangle", "7030A0"), "Жажда крови": ("diamond", "ED7D31"),
               "Механика": ("x", "595959"), "Расовая": ("circle", "FFC000"), "Тринкет": ("square", "00B0F0")}
    i = 0
    while i < len(ev):
        cat = ev[i]["cat"]
        j = i
        while j < len(ev) and ev[j]["cat"] == cat:
            j += 1
        ser = Series(Reference(s.ws, min_col=8, min_row=first + i, max_row=first + j - 1),
                     Reference(s.ws, min_col=2, min_row=first + i, max_row=first + j - 1), title=cat)
        sym, color = symbols.get(cat, ("circle", "A5A5A5"))
        ser.marker.symbol = sym
        ser.marker.size = 8
        ser.marker.graphicalProperties = GraphicalProperties(solidFill=color)
        ser.marker.graphicalProperties.line.solidFill = color
        ser.graphicalProperties.line.noFill = True
        ch.series.append(ser)
        i = j
    # Отличия > 10 с — отдельная серия
    dev = [e for e in ev if e.get("delta") is not None and abs(e["delta"]) > 10]
    if dev:
        s.cell(first - 1, 10, "Отличие > 10 с: время", bold=True, color="FFFFFF", fill=C_HEADER)
        s.cell(first - 1, 11, "Y", bold=True, color="FFFFFF", fill=C_HEADER)
        s.cell(first - 1, 12, "Действие", bold=True, color="FFFFFF", fill=C_HEADER)
        for k, e in enumerate(dev):
            s.cell(first + k, 10, _v(e["t"], 1), fmt=F_1)
            s.cell(first + k, 11, 2)
            s.cell(first + k, 12, f"{e['name']} №{e['k']}: {e['delta']:+.0f} с")
        s.ws.column_dimensions["J"].width = 14
        s.ws.column_dimensions["L"].width = 30
        ser = Series(Reference(s.ws, min_col=11, min_row=first, max_row=first + len(dev) - 1),
                     Reference(s.ws, min_col=10, min_row=first, max_row=first + len(dev) - 1),
                     title="Отличие > 10 с")
        ser.marker.symbol = "diamond"
        ser.marker.size = 13
        ser.marker.graphicalProperties = GraphicalProperties(solidFill="FF0000")
        ser.marker.graphicalProperties.line.solidFill = "C00000"
        ser.graphicalProperties.line.noFill = True
        ch.series.append(ser)
    s.chart(ch, "N3", w=26, h=10)

    if r.mm.get("dps_5s"):
        s.row = max(table_end, 26)
        s.section("Мой DPS по 5-секундным отрезкам")
        d = r.mm["dps_5s"]
        rows = [[i * 5, _mmss("A"), _v(v, 0), _v(r.ref.agg["dps"]["median"], 0)] for i, v in enumerate(d)]
        f2, l2 = s.table(["Время, с", "м:сс", "Мой DPS", "DPS эталона (медиана за бой)"], rows,
                         [F_INT, None, F_INT, F_INT], col=1)
        lc = LineChart()
        lc.title = "Мой DPS по времени и средний DPS эталона"
        lc.add_data(Reference(s.ws, min_col=3, max_col=4, min_row=f2 - 1, max_row=l2), titles_from_data=True)
        lc.set_categories(Reference(s.ws, min_col=2, min_row=f2, max_row=l2))
        lc.y_axis.numFmt = "#,##0"
        lc.x_axis.tickLblSkip = 6
        for ser in lc.series:
            ser.smooth = False
        s.chart(lc, "N26", w=26, h=9)


# ----------------------------------------------------------- Механики
def _mechanics_sheet(wb, r: CompareResult, demo):
    s = Sheet(wb, "Механики", demo, {"A": 26, "B": 6, "C": 14, "D": 14, "E": 12, "F": 12})
    s.title("Механики босса: тайминги вашего боя и эталона",
            "Большие расхождения означают другую стратегию или темп боя и снижают надёжность сравнения.")
    rows = [[x["name"], x["k"], _v(x["ref"], 1), _v(x["my"], 1), '=IF(D{r}="","",D{r}-C{r})',
             "ключевая" if x["key"] else ""] for x in r.tables["mechanics"]]
    s.table(["Механика", "№", "Эталон (медиана), с", "Мой бой, с", "Разница, с", "Урон игрокам"], rows,
            [None, F_INT, F_1, F_1, F_SIGNED1, None])
    s.note("Ключевая = входит в топ-8 по урону игрокам эталона и бьёт больше чем в половине боёв.")


# -------------------------------------------------- Практический гайд
def _guide_sheet(wb, ref: Ref, demo):
    from .guide import guide_sections
    s = Sheet(wb, "Гайд эталона", demo, {"A": 16, "B": 60, "C": 22, "D": 12})
    log0 = ref.logs[0]
    s.title(f"Как играть этот бой: {log0.encounter_name} — {log0.spec} {log0.cls}",
            f"Построено по {ref.n} логам ({ref.label}). Доверие — нижняя граница 95% интервала Уилсона.")
    for sec in guide_sections(ref):
        s.section(sec["title"])
        if sec["rows"]:
            s.table(["Когда", "Что делать", "Подтверждение", "Доверие"],
                    [[a, b, c, _v(d, 3)] for a, b, c, d in sec["rows"]], [None, None, None, F_PCT])
        else:
            s.note("Недостаточно данных.")


# ------------------------------------------------------------ Эталон
def _reference_sheet(wb, ref: Ref, demo):
    s = Sheet(wb, "Эталон", demo, {"A": 26, "B": 14, "C": 10, "D": 10, "E": 8, "F": 8, "G": 10, "H": 12,
                                   "I": 12, "J": 12, "K": 12, "L": 10})
    agg = ref.agg
    s.title(f"Статистика эталона ({ref.label}, {ref.n} логов)",
            "Категории способностей — эвристика по логам или spell_meta.json (колонка «Источник»).")
    rows = []
    for ab, a in sorted(agg["abilities"].items(), key=lambda kv: -(kv[1]["casts"]["median"] or 0)):
        sp = ref.spells[ab]
        c = a["casts"]
        rows.append([sp.name, CATEGORY_RU[sp.category], _v(c["mean"], 1), _v(c["median"], 1), _v(c["min"], 0),
                     _v(c["max"], 0), _v(a["share"], 3), _v(a["first_use"], 1), _v(a["interval"], 1),
                     _v(a["last_use"], 1), _v(a["uptime"], 3), sp.source])
    s.table(["Способность", "Категория", "Средн. касты", "Медиана", "Мин", "Макс", "Доля игроков",
             "Первое, с", "Интервал, с", "Последнее, с", "Время действия", "Источник"], rows,
            [None, None, F_1, F_1, F_INT, F_INT, F_PCT, F_1, F_1, F_1, F_PCT, None])

    s.section("Тайминги кулдаунов (k-е применение)",
              "Поддержка — игроки, чьё применение не дальше 10 с от медианы (или дальше, если так велик разброс топа).")
    rows = []
    for ab, rs in agg["cd_timing"].items():
        for x in rs:
            rows.append([ref.name_of(ab), x["k"], _v(x["median"], 1), _v(x["p25"], 1), _v(x["p75"], 1),
                         x["players"], x["support"], _v(x["confidence"], 3)])
    s.table(["Действие", "№", "Медиана, с", "P25, с", "P75, с", "Игроков", "Поддержка", "Доверие"], rows,
            [None, F_INT, F_1, F_1, F_1, F_INT, F_INT, F_PCT])

    s.section("Частые последовательности")
    for n in (2, 3, 4):
        rows = [[" → ".join(ref.name_of(x) for x in g["seq"]), g["count"], _v(g["share"], 4),
                 g["players"], _v(g["player_share"], 3)] for g in agg["ngrams"][n][:8]]
        s.table([f"{n} способности подряд", "Раз", "Доля", "Игроков (≥2 раз)", "Доля игроков"], rows,
                [None, F_INT, F_PCT, F_INT, F_PCT])


def _players_sheet(wb, ref: Ref, demo):
    s = Sheet(wb, "Игроки топа", demo, {"A": 6, "B": 16, "C": 12, "D": 11, "E": 8, "F": 14, "G": 14, "H": 12,
                                        "I": 12, "J": 50})
    s.title("Игроки эталона", "Паттерн, а не рейтинг: DPS зависит также от гира, состава и длительности.")
    main = ref.main_cd
    rows = []
    for log, m in zip(ref.logs, ref.metrics):
        first_main = m["cast_times"].get(main, [None])[0] if main else None
        rows.append([log.rank, log.name, _v(log.dps, 0), _v(log.duration, 1), _v(log.ilvl, 1),
                     _v(first_main, 1), len(m["potion_times"]), _v(m["gcd"]["idle_per_action"], 3),
                     _v(m["lust"], 1), log.url])
    s.table(["Ранг", "Игрок", "DPS", "Длительность, с", "Уровень предметов", f"Первый {ref.name_of(main) if main else 'CD'}, с",
             "Зелий", "Простой на действие, с", "Жажда крови, с", "Лог WCL"], rows,
            [F_INT, None, F_INT, F_1, F_1, F_1, F_INT, "0.000", F_1, None],
            links={9: lambda i: rows[i][9]})


def _raw_sheet(wb, r: CompareResult, demo):
    s = Sheet(wb, "Сырые данные", demo, {"A": 10, "B": 10, "C": 26, "D": 10, "E": 14, "F": 12, "G": 10})
    s.title("Мои касты", f"Источник: {r.me.url}")
    rows = []
    for c in r.me.casts:
        sp = r.ref.spells.get(c.id)
        res = f"{c.res[0]:.0f}/{c.res[1]:.0f}" if c.res else ""
        rows.append([_v(c.t, 2), _mmss("A"), r.me.name_of(c.id), c.id,
                     CATEGORY_RU[sp.category] if sp else "—", res, _v(c.hp, 0)])
    s.table(["Время, с", "м:сс", "Способность", "ID", "Категория", "Ресурс", "Здоровье, %"], rows,
            ["0.00", None, None, "0", None, None, F_INT])


# ===================================================== РОТАЦИЯ ВСЕГО РЕЙДА
def write_raid_rotation_workbook(R: dict, path: str | Path) -> Path:
    """Сводка по всем DPS рейда + лист «Что исправить» + выжимка на каждого игрока."""
    import re as _re
    wb = Workbook()
    wb.remove(wb.active)
    demo = R["info"]["demo"]
    I = R["info"]
    s = Sheet(wb, "Ротация рейда", demo, {"A": 18, "B": 28, "C": 12, "D": 14, "E": 12, "F": 13, "G": 60, "H": 18,
                                                      "I": 14, "J": 22, "K": 16, "L": 13, "M": 14, "N": 16})
    s.title(f"Сравнение ротации рейда: {I['boss']}, {I['difficulty']}",
            f"Каждый DPS против топ-{I['top_n']} своего спека на этом боссе. Бой {I['duration']}.")
    for line in R["brief"]:
        s.cell(s.row, 1, "• " + line)
        s.row += 1
    s.row += 1
    ok = sorted([r for r in R["rows"] if "error" not in r], key=lambda r: r["gap"])
    rows = [[r["name"], spec_ru(r["cls"], r["spec"]), _v(r["dps"], 0), _v(r["ref_dps"], 0), "=C{r}/D{r}-1",
             REL_RU.get(r["reliability"], r["reliability"]), (r["actions"][0]["do"] if r["actions"] else "существенных отличий нет"),
             (r["actions"][0]["gain"] if r["actions"] else ""),
             _v(r.get("rank"), 0), _v(r.get("ilvl"), 0), _v(r.get("mythic_dps"), 0), _v(r.get("mythic_gap"), 4),
             _v(r.get("item_level"), 1), _v(r.get("ref_item_level"), 1)] for r in ok]
    first, last = s.table(["Игрок", "Спек", "DPS", "Медиана топа", "К топу", "Надёжность", "Главное действие",
                           "Эффект", "Процентиль, %", "С учётом экипировки, %", "Эпох. топ (рейтинг)", "К эпох. топу",
                           "Ур. предметов", "Ур. предметов топа"], rows,
                          [None, None, F_INT, F_INT, F_SIGNED_PCT, None, None, None, F_INT, F_INT, F_INT, F_SIGNED_PCT,
                           F_1, F_1])
    if rows:
        for i in range(len(rows)):
            s.ws.cell(row=first + i, column=7).alignment = Alignment(wrap_text=True, vertical="top")
        s.ws.conditional_formatting.add(f"E{first}:E{last}", CellIsRule(
            operator="lessThan", formula=["-0.1"], fill=PatternFill("solid", fgColor="FFC7CE")))
        ch = BarChart()
        ch.type = "bar"
        ch.title = "Отставание от медианы топа своего спека"
        ch.add_data(Reference(s.ws, min_col=5, min_row=first - 1, max_row=last), titles_from_data=True)
        ch.set_categories(Reference(s.ws, min_col=1, min_row=first, max_row=last))
        ch.y_axis.numFmt = "0%"
        ch.legend = None
        s.chart(ch, f"A{s.row}", w=18, h=max(6, 0.7 * len(rows) + 3))
        s.row += int(max(6, 0.7 * len(rows) + 3) * 2) + 2
    errs = [r for r in R["rows"] if "error" in r]
    if errs:
        s.section("Не удалось разобрать")
        s.table(["Игрок", "Причина"], [[r["name"], r["error"]] for r in errs])

    t = Sheet(wb, "Что исправить", demo, {"A": 18, "B": 5, "C": 62, "D": 20, "E": 50})
    t.title("Что исправить каждому", "До трёх действий на игрока, по убыванию эффекта.")
    t.table(["Игрок", "№", "Что делать", "Эффект", "Почему"],
            [[r["name"], i + 1, a["do"], a["gain"], a["title"]] for r in ok for i, a in enumerate(r["actions"])])
    for col in (3, 5):
        for row in t.ws.iter_rows(min_row=1, min_col=col, max_col=col):
            for c in row:
                c.alignment = Alignment(wrap_text=True, vertical="top")

    used = set(wb.sheetnames)
    for r in ok:
        name = _re.sub(r"[\\/*?:\[\]]", "", r["name"])[:28] or "Игрок"
        base, k = name, 2
        while name in used:
            name, k = f"{base[:26]} {k}", k + 1
        used.add(name)
        _brief_sheet(wb, r["result"], demo, title=name)
    path = Path(path)
    wb.save(path)
    return path
